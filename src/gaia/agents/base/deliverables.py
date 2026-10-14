# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""Do the figures in the files this turn wrote rest on what its tools saw?

The answer checks in :mod:`gaia.agents.base.grounding` read the chat reply
only, so a price invented inside ``run_python`` and saved to ``bom.xlsx``
went unseen. This reads each deliverable the turn wrote — text formats and
the cells of .xlsx/.docx/.pptx, bounded in size — and finds its money,
percentages and unit quantities. Each must appear in a tool result this turn,
an input file, or the conversation, or follow from ones that do: a line
total, a column sum, a share, a per-unit split. A figure the file itself
marks as an estimate is left alone.

Evidence is what tools *returned*, never what the call sent: a search that
echoes its query, a script's own literals, and reading back the file this
turn wrote prove nothing about where a figure came from.

A finding asks the model to verify or mark the figures; nothing here edits
the file.
"""

from __future__ import annotations

import ast
import csv
import io
import json
import logging
import os
import re
import zipfile
from dataclasses import dataclass
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Set, Tuple

from gaia.agents.base.claims import admits_unverified
from gaia.agents.base.grounding import (
    _EXEC_WORDS,
    _TABLE_RULE,
    Finding,
    _args_text,
    _Corpus,
    _has,
    _name,
    _ran,
    _rounds_to,
    _Value,
    _value,
    evidence_text,
)

logger = logging.getLogger(__name__)

FILE = "file"

#: Formats whose figures are read; anything else is not a deliverable here.
DELIVERABLE_SUFFIXES = frozenset(
    {".md", ".markdown", ".txt", ".csv", ".tsv", ".json", ".xlsx", ".docx", ".pptx"}
)
MAX_FILE_BYTES = 5_000_000
MAX_TEXT_CHARS = 400_000
MAX_ROWS = 2_000
MAX_COLS = 40
MAX_SHEETS = 10
MAX_CLAIMS = 200
#: Values named per file in the correction and the note.
_SHOWN = 5


@dataclass
class Block:
    """A table (``header`` + rows of cells) or prose (one-cell rows)."""

    rows: List[List[str]]
    header: Optional[List[str]] = None
    table: bool = False


# ---------------------------------------------------------------------------
# Reading a file into blocks
# ---------------------------------------------------------------------------

_PURE_NUMBER = re.compile(r"^\s*-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\s*$")


def _table(rows: List[List[str]]) -> Block:
    """A table whose header is its first row of labels, if one comes early."""
    rows = [r[:MAX_COLS] for r in rows[:MAX_ROWS] if any(c.strip() for c in r)]
    for index, row in enumerate(rows[:5]):
        filled = [c for c in row if c.strip()]
        if len(filled) >= 2 and not any(_PURE_NUMBER.match(c) for c in filled):
            return Block(rows[:index] + rows[index + 1 :], row, True)
    return Block(rows, None, True)


def _md_cells(line: str) -> List[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _markdown_blocks(text: str) -> List[Block]:
    blocks: List[Block] = []
    prose: List[List[str]] = []
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        rule = index + 1 < len(lines) and _TABLE_RULE.match(lines[index + 1])
        if line.strip().startswith("|") and rule:
            header = _md_cells(line)
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_md_cells(lines[index]))
                index += 1
            blocks.append(Block(rows[:MAX_ROWS], header, True))
            continue
        if line.strip():
            prose.append([line])
        index += 1
    if prose:
        blocks.insert(0, Block(prose[:MAX_ROWS]))
    return blocks


def _scalar(value: Any) -> bool:
    return value is None or isinstance(value, (str, int, float, bool))


def _json_cell(value: Any) -> str:
    return "" if value is None else str(value)


def _json_blocks(data: Any, out: List[Block], depth: int = 0) -> None:
    if depth > 20 or len(out) > 500:
        return
    if isinstance(data, list):
        dicts = [d for d in data[:MAX_ROWS] if isinstance(d, dict)]
        keys = list(dict.fromkeys(k for d in dicts for k, v in d.items() if _scalar(v)))
        if keys:
            out.append(
                Block(
                    [[_json_cell(d.get(k)) for k in keys] for d in dicts],
                    [str(k) for k in keys],
                    True,
                )
            )
        strings = [[s] for s in data[:MAX_ROWS] if isinstance(s, str)]
        if strings:
            out.append(Block(strings))
        for item in data[:MAX_ROWS]:
            children = item.values() if isinstance(item, dict) else [item]
            for child in children:
                if not _scalar(child):
                    _json_blocks(child, out, depth + 1)
    elif isinstance(data, dict):
        keys = [k for k, v in data.items() if _scalar(v)]
        if keys:
            out.append(
                Block(
                    [[_json_cell(data[k]) for k in keys]], [str(k) for k in keys], True
                )
            )
        for value in data.values():
            if not _scalar(value):
                _json_blocks(value, out, depth + 1)


def text_blocks(text: str, suffix: str) -> List[Block]:
    """*text* as blocks, read the way its *suffix* says."""
    text = text[:MAX_TEXT_CHARS]
    if suffix in (".csv", ".tsv"):
        delimiter = "\t" if suffix == ".tsv" else ","
        try:
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        except csv.Error:
            return _markdown_blocks(text)
        return [_table(rows)] if rows else []
    if suffix == ".json":
        try:
            data = json.loads(text)
        except json.JSONDecodeError:
            return _markdown_blocks(text)
        out: List[Block] = []
        _json_blocks(data, out)
        return out
    return _markdown_blocks(text)


def _format_decimals(section: str, value: Any) -> int:
    match = re.search(r"\.(0+)", section)
    if match:
        return len(match.group(1))
    if isinstance(value, float) and not value.is_integer():
        digits = repr(value)
        if "e" in digits or "." not in digits:
            return 2
        return min(len(digits.split(".", 1)[1]), 4)
    return 0


def _xlsx_cell(cell: Any) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return str(value)
    section = (getattr(cell, "number_format", None) or "General").split(";")[0]
    if "%" in section:
        return f"{value * 100:.{_format_decimals(section, value * 100)}f}%"
    decimals = _format_decimals(section, value)
    text = f"{value:,.{decimals}f}" if "," in section else f"{value:.{decimals}f}"
    symbol = re.search(r"[$€£¥]", section)
    return symbol.group(0) + text if symbol else text


def _xlsx_blocks(path: str) -> List[Block]:
    import openpyxl  # pylint: disable=import-outside-toplevel

    # Formulas stay formulas: a computed cell is derived by construction.
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        blocks = []
        for sheet in book.worksheets[:MAX_SHEETS]:
            rows = [
                [_xlsx_cell(c) for c in row]
                for row in sheet.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS)
            ]
            if any(any(r) for r in rows):
                blocks.append(_table(rows))
        return blocks
    finally:
        book.close()


def _docx_blocks(path: str) -> List[Block]:
    from docx import Document  # pylint: disable=import-outside-toplevel

    document = Document(path)
    prose = [[p.text] for p in document.paragraphs[: MAX_ROWS * 2] if p.text.strip()]
    blocks = [Block(prose)] if prose else []
    for table in document.tables[:50]:
        rows = [[c.text for c in r.cells[:MAX_COLS]] for r in table.rows[:MAX_ROWS]]
        if rows:
            blocks.append(_table(rows))
    return blocks


def _pptx_blocks(path: str) -> List[Block]:
    from pptx import Presentation  # pylint: disable=import-outside-toplevel

    blocks = []
    for slide in list(Presentation(path).slides)[:200]:
        prose = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs)
                    if text.strip():
                        prose.append([text])
            if getattr(shape, "has_table", False):
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                if rows:
                    blocks.append(_table(rows))
        if prose:
            blocks.append(Block(prose))
    return blocks


def read_blocks(path: str) -> List[Block]:
    """The figures-bearing content of the file at *path*, bounded in size.

    ``[]`` for an unsupported, oversized or unreadable file, logged: this is
    an advisory check and must never fail the turn.
    """
    suffix = os.path.splitext(path)[1].lower()
    if suffix not in DELIVERABLE_SUFFIXES:
        return []
    try:
        if os.path.getsize(path) > MAX_FILE_BYTES:
            logger.info(
                "Not checking figures in %s: larger than %d bytes", path, MAX_FILE_BYTES
            )
            return []
        if suffix == ".xlsx":
            return _xlsx_blocks(path)
        if suffix == ".docx":
            return _docx_blocks(path)
        if suffix == ".pptx":
            return _pptx_blocks(path)
        with open(path, encoding="utf-8", errors="replace") as handle:
            return text_blocks(handle.read(MAX_TEXT_CHARS), suffix)
    except ImportError as error:
        logger.info("Not checking figures in %s: %s is not installed", path, error.name)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as error:
        logger.warning("Could not read %s to check its figures: %s", path, error)
    return []


def blocks_text(blocks: Sequence[Block]) -> str:
    """*blocks* as plain lines, cells joined by `` | `` so rows stay tabular."""
    lines = []
    for block in blocks:
        if block.header:
            lines.append("| " + " | ".join(block.header) + " |")
        lines.extend("| " + " | ".join(row) + " |" for row in block.rows)
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Which tool calls and files point at inputs
# ---------------------------------------------------------------------------

_PATH_ARGS = ("file_path", "path", "input_path", "source", "filename", "file")


def input_paths(records: Sequence[Dict[str, Any]]) -> List[str]:
    """Deliverable-format files this turn's tools were pointed at, in order.

    Path arguments, plus string literals in code a tool ran — a script that
    reads ``prices.csv`` makes that file an input even if nothing printed it.
    """
    found: List[str] = []
    for record in records:
        if not _ran(record) or record.get("failed"):
            continue
        args = record.get("args") or {}
        if not isinstance(args, dict):
            continue
        for key in _PATH_ARGS:
            if isinstance(args.get(key), str):
                found.append(args[key])
        code = args.get("code")
        if isinstance(code, str):
            try:
                tree = ast.parse(code)
            except SyntaxError:
                continue
            found.extend(
                node.value
                for node in ast.walk(tree)
                if isinstance(node, ast.Constant)
                and isinstance(node.value, str)
                and "\n" not in node.value
                and len(node.value) < 1024
            )
    return [
        p
        for p in dict.fromkeys(found)
        if "\x00" not in p and os.path.splitext(p)[1].lower() in DELIVERABLE_SUFFIXES
    ]


# ---------------------------------------------------------------------------
# Figures in the file
# ---------------------------------------------------------------------------

_NUM = r"\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?"
_CODES = r"USD|EUR|GBP|CAD|AUD|JPY|CHF|CNY|INR"
_MONEY = re.compile(
    rf"(?:US\$|C\$|A\$|\$|€|£|¥|\b(?:{_CODES})\s?)\s?(?P<a>{_NUM})"
    rf"|(?<![\w.,])(?P<b>{_NUM})\s?(?:€|£|\b(?:{_CODES})\b)"
)
_PERCENT = re.compile(rf"(?<![\w.,])({_NUM})\s?%")
_UNIT = re.compile(
    rf"(?<![\w.,])({_NUM})\s?(kW|mW|W|mV|V|mAh|Ah|mA|A|GHz|MHz|kHz|Hz|TB|GB|MB|KB|kB|"
    r"Gbps|Mbps|ms|µs|ns|°C|mm|cm|kg)(?![\w/])"
)
_ANY_NUMBER = re.compile(rf"(?<![\w.,])(?:{_NUM})(?![\w]|\.\d)")
#: A column of money: "Unit price", "Line total (USD)", "Cost $".
_MONEY_HEADER = re.compile(
    r"price|cost|total|amount|subtotal|budget|fee|msrp|spend|revenue|usd|eur|gbp|[$€£]",
    re.I,
)
#: ...unless it counts or measures something else: "Total qty", "Total time (s)".
_COUNT_HEADER = re.compile(
    r"qty|quantity|\bunits\b|\bcount\b|\bpcs\b|hours|%|percent|#|number of|time|"
    r"\bsecs?\b|seconds|\bms\b|duration|size|bytes|weight|tokens|score",
    re.I,
)
#: The file itself says the figure is not confirmed.
_MARKER = re.compile(
    r"\b(?:est\.?|estimated?|estimates?|approx\.?|approximate(?:ly)?|unverified|"
    r"not\s+verified|unconfirmed|assumed|assumption|tbd|tbc|placeholder|ballpark|"
    r"indicative)(?!\w)|~|≈",
    re.I,
)
#: "All prices are estimates", "Figures below are unverified".
_BLANKET = re.compile(
    r"\b(?:all|these|the|listed|above|below)\s+(?:\w+\s+){0,2}"
    r"(?:prices?|costs?|values?|figures?|numbers?|amounts?|totals?)\b[^.\n]{0,60}?"
    r"\b(?:estimat\w*|unverified|not\s+verified|approximat\w*|indicative|placeholders?)",
    re.I,
)


@dataclass
class _Claim:
    #: ``"money"``, ``"percent"``, or the unit symbol.
    kind: str
    value: _Value
    shown: str
    #: The column or paragraph it sits in, and its place there.
    seq: Tuple[int, int]
    pos: int
    row: int
    #: The file says it is an estimate: never reported, still an operand.
    marked: bool = False


def _number(text: str) -> float:
    return float(text.replace(",", ""))


def _scan(blocks: Sequence[Block]) -> Tuple[List[_Claim], List[List[float]], List[str]]:
    """``(claims, numbers per row, every plain number's text)``."""
    claims: List[_Claim] = []
    rows: List[List[float]] = []
    plain: List[str] = []
    for b, block in enumerate(blocks):
        header = block.header or []
        money_cols = {
            j
            for j, h in enumerate(header)
            if _MONEY_HEADER.search(h) and not _COUNT_HEADER.search(h)
        }
        marked_cols = {j for j, h in enumerate(header) if _MARKER.search(h)}
        place = 0
        for i, row in enumerate(block.rows):
            row_text = " | ".join(row)
            numbers = _ANY_NUMBER.findall(row_text)
            plain.extend(numbers)
            rows.append([_number(n) for n in numbers])
            row_marked = bool(_MARKER.search(row_text))
            for j, cell in enumerate(row):
                if cell.lstrip().startswith("="):
                    continue
                for kind, text, shown in _cell_figures(cell, j in money_cols):
                    if not shown:
                        shown = f"{text} ({header[j].strip()})"
                    pos = i if block.table else place
                    place += 1
                    seq = (b, j if block.table else -1)
                    marked = row_marked or j in marked_cols
                    claims.append(
                        _Claim(
                            kind, _value(text), shown, seq, pos, len(rows) - 1, marked
                        )
                    )
    return claims[:MAX_CLAIMS], rows, plain


def _cell_figures(cell: str, money_column: bool) -> List[Tuple[str, str, str]]:
    """``(kind, number, shown)`` per figure in *cell*; ``shown`` is ``""`` for a
    bare number in a money column, which the caller labels with the header."""
    figures = [
        ("money", m.group("a") or m.group("b"), m.group(0).strip())
        for m in _MONEY.finditer(cell)
    ]
    if not figures and money_column and _PURE_NUMBER.match(cell):
        figures.append(("money", cell.strip(), ""))
    figures += [
        ("percent", m.group(1), m.group(0).strip()) for m in _PERCENT.finditer(cell)
    ]
    figures += [
        (m.group(2), m.group(1), m.group(0).strip()) for m in _UNIT.finditer(cell)
    ]
    return figures


# ---------------------------------------------------------------------------
# Evidence
# ---------------------------------------------------------------------------

_WRITE_WORDS = (
    "write",
    "edit",
    "save",
    "create",
    "append",
    "export",
    "replace",
    "insert",
)
_MONEY_WORDS = re.compile(
    r"price|pricing|cost|msrp|subtotal|quote|\bfee\b|usd|eur|gbp", re.I
)
#: A data row: CSV (commas, never ", " as prose has), tabs, or a pipe table.
_TABULAR = re.compile(r"^(?!.*, )[^,]+,|\t|\|.*\|")


def _money_evidence(text: str) -> str:
    """The numbers in *text* that can be a price.

    Tagged with a currency, on a short line that talks about price, in a
    data row (CSV, tab, pipe table), or the value under a ``price`` key. A
    datasheet's "60 Hz" is not a source for $60.
    """
    parts = [m.group("a") or m.group("b") for m in _MONEY.finditer(text)]
    previous = ""
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if len(stripped) <= 400 and _TABULAR.search(stripped):
            parts.append(stripped)
        elif len(stripped) <= 200 and _MONEY_WORDS.search(stripped):
            parts.append(stripped)
        elif _PURE_NUMBER.match(stripped) and _MONEY_WORDS.search(previous):
            parts.append(stripped)
        previous = stripped
    return "\n".join(parts)


def _basename(path: str) -> str:
    return re.split(r"[\\/]", path.rstrip("\\/"))[-1].lower()


def _evidence_texts(records: Sequence[Dict[str, Any]], written: Set[str]) -> List[str]:
    """What this turn's tools returned that can source a figure in *written*.

    Not a write tool's reply, not a failed call, and not a read of a written
    file once it was written — a read-back shows the file, not its source.
    The script that wrote it still counts for what it computed, minus its
    own literals.
    """
    first_write: Dict[str, int] = {}
    for index, record in enumerate(records):
        if _has(_name(record), _WRITE_WORDS + _EXEC_WORDS):
            args = _args_text(record)
            for name in written:
                if name in args:
                    first_write.setdefault(name, index)
    first_exec = next(
        (i for i, r in enumerate(records) if _has(_name(r), _EXEC_WORDS)),
        len(records),
    )
    texts = []
    for index, record in enumerate(records):
        name = _name(record)
        if not _ran(record) or record.get("failed"):
            continue
        if _has(name, _WRITE_WORDS) and not _has(name, _EXEC_WORDS):
            continue
        args = _args_text(record)
        if any(
            name_ in args and index > first_write.get(name_, first_exec - 1)
            for name_ in written
        ):
            continue
        texts.append(evidence_text(record, code_literals=True))
    return texts


# ---------------------------------------------------------------------------
# Derivation
# ---------------------------------------------------------------------------


def _any_rounds(results: Iterable[float], value: _Value) -> bool:
    return any(_rounds_to(r, value) for r in results)


def _run_sums(
    index: int, claims: Sequence[_Claim], known: Sequence[bool], backward_only: bool
) -> bool:
    """*claims[index]* is the sum of the known figures just before (or after)
    it, or of the leading rows of another table column — a total in prose."""
    claim = claims[index]
    peers = [
        k
        for k, c in enumerate(claims)
        if c.seq == claim.seq and c.kind == claim.kind and k != index
    ]
    before = [k for k in peers if claims[k].pos < claim.pos][::-1]
    after = [k for k in peers if claims[k].pos > claim.pos]
    runs = [before] if backward_only else [before, after]
    columns: Dict[Tuple[int, int], List[int]] = {}
    for k, c in enumerate(claims):
        if c.seq != claim.seq and c.seq[1] >= 0 and c.kind == claim.kind:
            columns.setdefault(c.seq, []).append(k)
    runs.extend(columns.values())
    for run in runs:
        total = 0.0
        for k in run:
            if not known[k]:
                break
            total += claims[k].value.number
            if _rounds_to(total, claim.value):
                return True
    return False


def _follows(
    index: int,
    claims: Sequence[_Claim],
    known: Sequence[bool],
    rows: Sequence[Sequence[float]],
    request_numbers: Sequence[float],
    plain_known: Sequence[float],
) -> bool:
    """*claims[index]* follows from known figures by one arithmetic step."""
    claim = claims[index]
    value = claim.value
    same = list(
        dict.fromkeys(
            c.value.number
            for k, c in enumerate(claims)
            if known[k] and k != index and c.kind == claim.kind
        )
    )[:40]
    if _any_rounds(same, value):
        return True
    for i, a in enumerate(same):
        for b in same[i + 1 :]:
            if _any_rounds((a + b, abs(a - b), (a + b) / 2), value):
                return True
    multipliers = {
        m
        for m in list(rows[claim.row]) + list(request_numbers)
        if m not in (0.0, 1.0) and m <= 100_000 and m != value.number
    }
    percents = [
        c.value.number for k, c in enumerate(claims) if known[k] and c.kind == "percent"
    ]
    for a in same:
        if _any_rounds((r for m in multipliers for r in (a * m, a / m)), value):
            return True
        if _any_rounds(
            (
                r
                for p in percents
                for r in (a * p / 100, a * (1 + p / 100), a * (1 - p / 100))
            ),
            value,
        ):
            return True
    if _run_sums(index, claims, known, backward_only=False):
        return True
    if claim.kind == "percent":
        money = [
            c.value.number
            for k, c in enumerate(claims)
            if known[k] and c.kind == "money"
        ]
        pool = list(dict.fromkeys(money + list(plain_known)))[:40]
        for a in pool:
            for b in pool:
                if (
                    b
                    and a != b
                    and _any_rounds((100 * a / b, 100 * abs(a - b) / b), value)
                ):
                    return True
    return False


def _downstream(
    index: int,
    claims: Sequence[_Claim],
    rows: Sequence[Sequence[float]],
    request_numbers: Sequence[float],
) -> bool:
    """A total or line amount computed from other figures, sourced or not.

    Listing it beside the unit price it was computed from is noise: fixing
    the price fixes it. Only multiplication and running sums, so an invented
    price is never excused as the difference of two others.
    """
    claim = claims[index]
    known = [True] * len(claims)
    known[index] = False
    others = list(
        dict.fromkeys(
            c.value.number
            for k, c in enumerate(claims)
            if k != index and c.kind == claim.kind
        )
    )[:60]
    multipliers = {
        m
        for m in list(rows[claim.row]) + list(request_numbers)
        if m not in (0.0, 1.0) and m <= 100_000 and m != claim.value.number
    }
    if _any_rounds((a * m for a in others for m in multipliers), claim.value):
        return True
    return _run_sums(index, claims, known, backward_only=True)


def _unsupported(
    blocks: Sequence[Block],
    general: _Corpus,
    money: _Corpus,
    request_numbers: Sequence[float],
    data_numbers: Sequence[float] = (),
) -> List[_Claim]:
    """Figures in *blocks* with no source, minus totals computed from them.

    *data_numbers* — numbers from data rows the tools returned — may be the
    operands of a percentage the file shows without them.
    """
    if _BLANKET.search(blocks_text(blocks)):
        return []
    claims, rows, plain = _scan(blocks)
    if not claims:
        return []
    known = [
        c.marked or (money if c.kind == "money" else general).holds(c.value)
        for c in claims
    ]
    plain_known = list(
        dict.fromkeys(
            [v.number for v in (_value(t) for t in plain) if general.holds(v)]
            + list(data_numbers)
        )
    )[:40]
    for _ in range(4):
        changed = False
        for index in range(len(claims)):
            if not known[index] and _follows(
                index, claims, known, rows, request_numbers, plain_known
            ):
                known[index] = changed = True
        if not changed:
            break
    missing = [k for k, ok in enumerate(known) if not ok]
    roots = [k for k in missing if not _downstream(k, claims, rows, request_numbers)]
    return [claims[k] for k in (roots or missing)]


# ---------------------------------------------------------------------------
# The check
# ---------------------------------------------------------------------------


_GLUED_UNIT = re.compile(r"(\d)([A-Za-zµμ°])")

#: Dates and clock times in the request are not quantities to multiply by.
_TIMESTAMP = re.compile(
    r"\d{4}-\d{2}-\d{2}(?:[T ]\d{1,2}:\d{2}(?::\d{2})?(?:\.\d+)?(?:Z|[+-]\d{2}:?\d{2})?)?"
    r"|(?<!\d)\d{1,2}:\d{2}(?::\d{2})?(?!\d)"
)


def _mentioned(claim: _Claim, answer: str) -> bool:
    return claim.value.number in {_number(n) for n in _ANY_NUMBER.findall(answer)}


def file_findings(
    files: Mapping[str, Sequence[Block]],
    records: Sequence[Dict[str, Any]],
    *,
    query: str = "",
    context: str = "",
    inputs: Sequence[str] = (),
    answer: str = "",
) -> List[Finding]:
    """Figures in *files* (path → blocks) that nothing this turn supports.

    Evidence is the tool results in *records*, the text of *inputs* (files
    the tools were pointed at), the request (*query*) and earlier turns
    (*context*). An answer that says its figures are unverified and names a
    value has already told the user about it.
    """
    if not files:
        return []
    written = {_basename(path) for path in files}
    texts = _evidence_texts(records, written) + list(inputs)
    # "27W", "8GB": a unit glued to its number still shows the number.
    general = _Corpus(_GLUED_UNIT.sub(r"\1 \2", "\n".join([context, query] + texts)))
    money = _Corpus(
        "\n".join([query] + [_money_evidence(t) for t in [context] + texts])
    )
    request_numbers = [
        _number(n) for n in _ANY_NUMBER.findall(_TIMESTAMP.sub(" ", query or ""))
    ][:20]
    data_numbers = list(
        dict.fromkeys(
            _number(n)
            for t in texts
            for line in t.splitlines()
            if len(line) <= 400 and _TABULAR.search(line)
            for n in re.split(r"[,;\t|]", line)
            if _PURE_NUMBER.match(n)
        )
    )[:20]
    admitted = admits_unverified(answer or "")
    findings: List[Finding] = []
    for path, blocks in files.items():
        missing = _unsupported(blocks, general, money, request_numbers, data_numbers)
        if admitted:
            missing = [c for c in missing if not _mentioned(c, answer)]
        if not missing:
            continue
        shown = list(dict.fromkeys(c.shown for c in missing))
        listed = ", ".join(shown[:_SHOWN]) + (
            f" and {len(shown) - _SHOWN} more" if len(shown) > _SHOWN else ""
        )
        name = re.split(r"[\\/]", path.rstrip("\\/"))[-1]
        findings.append(
            Finding(
                FILE,
                f"`{name}` has figures no tool result this turn supports: {listed}. "
                "Show their source with a tool (reading the file back doesn't "
                "count) and fix the file, or mark them in it as estimates; name "
                "any still unverified in your answer.",
                f"`{name}` has figures I found no source for this turn: {listed}",
            )
        )
    return findings
