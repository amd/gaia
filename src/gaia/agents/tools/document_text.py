# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""Full text of Word, PDF and Excel files for the read tools, in budget-sized parts.

Reading a handful of documents in full is how cross-document facts get noticed;
retrieval only returns the passages a query happened to ask for. Each part fits
the model's tool-result budget and says where the next one starts, so a long
document is read completely instead of being silently cut.
"""

import math
import os
from typing import Any, Dict, Optional, Tuple

#: Extensions read_file extracts text from, and the file_type it reports.
READABLE_DOCUMENTS = {".docx": "docx", ".pdf": "pdf", ".xlsx": "xlsx"}

#: Binary document formats only index_document can read.
INDEX_ONLY_DOCUMENTS = frozenset(
    {".doc", ".pptx", ".ppt", ".xls", ".odt", ".ods", ".odp", ".epub"}
)

#: Chars one read returns when the host has no model-sized budget.
DEFAULT_READ_CHARS = 20000

#: Past this many reads, a document gets a hint that retrieval is cheaper.
FULL_READ_MAX_READS = 3

#: Extraction stops here; a longer document is for index_document.
MAX_EXTRACT_CHARS = 1_000_000

#: A character ``limit`` this small is almost always meant as a line count.
SMALL_LIMIT_CHARS = 1000

#: Extracted documents kept per agent, so paging does not re-parse the file.
_CACHE_ENTRIES = 8


def read_budget(host: Any) -> int:
    """Chars one read may return: the host's tool-result target, so it is never truncated."""
    budget = getattr(host, "_truncation_budget", None)
    return budget()[1] if callable(budget) else DEFAULT_READ_CHARS


def document_type(path: str) -> Optional[str]:
    """``"docx"`` / ``"pdf"`` / ``"xlsx"`` for a readable document, else ``None``."""
    return READABLE_DOCUMENTS.get(os.path.splitext(path)[1].lower())


def index_only_error(path: str) -> Optional[Dict[str, Any]]:
    """The error for a binary document read_file cannot extract, else ``None``."""
    ext = os.path.splitext(path)[1].lower()
    if ext not in INDEX_ONLY_DOCUMENTS:
        return None
    return {
        "status": "error",
        "error": (
            f"read_file cannot extract {ext} files. Call index_document('{path}'), "
            "then query_specific_file to retrieve its content. read_file reads "
            ".docx, .pdf and .xlsx text directly."
        ),
    }


def _too_long(path: str) -> ValueError:
    return ValueError(
        f"{os.path.basename(path)} has more than {MAX_EXTRACT_CHARS:,} characters "
        "of text, too long to read in full. Call index_document on it, then "
        "query_specific_file for what you need."
    )


def _pdf_text(path: str) -> Tuple[str, Dict[str, Any]]:
    try:
        from pypdf import PdfReader  # pylint: disable=import-outside-toplevel
        from pypdf.errors import PdfReadError  # pylint: disable=import-outside-toplevel
    except ImportError as e:
        raise ImportError(
            "Reading PDF text requires pypdf. Install it with: "
            'uv pip install -e ".[rag]"'
        ) from e

    name = os.path.basename(path)
    try:
        reader = PdfReader(path)
    except PdfReadError as e:
        raise ValueError(
            f"Could not read PDF {name}: {e}. The file looks corrupted or is not "
            "a PDF; re-download or re-export it."
        ) from e
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError(
            f"{name} is password-protected. Ask the user for an unlocked copy."
        )

    pages, size, has_text = [], 0, False
    for number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        has_text = has_text or bool(text)
        pages.append(f"[Page {number}]\n{text}")
        size += len(pages[-1]) + 2
        if size > MAX_EXTRACT_CHARS:
            raise _too_long(path)
    if not has_text:
        raise ValueError(
            f"{name} has {len(pages)} page(s) but no text layer — it is probably "
            "a scan. index_document can read it with the vision model, if one is "
            "installed."
        )
    return "\n\n".join(pages), {"pages": len(pages)}


def _xlsx_text(path: str) -> Tuple[str, Dict[str, Any]]:
    try:
        import openpyxl  # pylint: disable=import-outside-toplevel
    except ImportError as e:
        raise ImportError(
            "Reading Excel text requires openpyxl. Install it with: "
            'uv pip install -e ".[rag]"'
        ) from e

    def load(data_only: bool):
        try:
            return openpyxl.load_workbook(path, read_only=True, data_only=data_only)
        except Exception as e:
            raise ValueError(
                f"Could not read Excel file {os.path.basename(path)}: {e}. It may "
                "be corrupted or not an .xlsx workbook."
            ) from e

    values = load(data_only=True)
    try:
        formulas = load(data_only=False)
    except ValueError:
        values.close()
        raise
    try:
        parts, size, uncomputed = [], 0, 0
        for sheet, source in zip(values.worksheets, formulas.worksheets):
            parts.append(f"[Sheet: {sheet.title}]")
            rows = zip(sheet.iter_rows(values_only=True), source.iter_rows())
            for row, cells_in in rows:
                for value, cell in zip(row, cells_in):
                    if value is None and getattr(cell, "data_type", None) == "f":
                        uncomputed += 1
                cells = ["" if cell is None else str(cell) for cell in row]
                while cells and not cells[-1].strip():
                    cells.pop()
                if cells:
                    parts.append(" | ".join(cells))
                    size += len(parts[-1]) + 1
                    if size > MAX_EXTRACT_CHARS:
                        raise _too_long(path)
        meta: Dict[str, Any] = {"sheets": len(values.worksheets)}
        if uncomputed:
            # A workbook saved by a script has formulas but no computed values.
            meta["warning"] = (
                f"{uncomputed} formula cell(s) have no saved value and read as "
                "blank; open and save the file in Excel, or compute them yourself."
            )
        return "\n".join(parts), meta
    finally:
        values.close()
        formulas.close()


def extract_document_text(path: str) -> Tuple[str, Dict[str, Any]]:
    """``(text, metadata)`` for a .docx, .pdf or .xlsx file.

    PDF text carries ``[Page N]`` markers and Excel text ``[Sheet: name]``
    markers, so a statement can be cited by where it came from.

    Raises:
        ImportError: the format's parser is not installed.
        ValueError: the file is unreadable, encrypted, has no text, or is too
            long to read in full.
    """
    kind = document_type(path)
    if kind == "pdf":
        return _pdf_text(path)
    if kind == "xlsx":
        return _xlsx_text(path)
    if kind == "docx":
        from gaia.rag.docx_utils import (  # pylint: disable=import-outside-toplevel
            extract_docx_text,
        )

        text = extract_docx_text(path)
        if len(text) > MAX_EXTRACT_CHARS:
            raise _too_long(path)
        return text, {}
    raise ValueError(f"Not a .docx, .pdf or .xlsx file: {path}")


def cached_document_text(host: Any, path: str) -> Tuple[str, Dict[str, Any]]:
    """:func:`extract_document_text`, reused while the file is unchanged."""
    real = os.path.realpath(path)
    stat = os.stat(real)
    key = (real, stat.st_mtime_ns, stat.st_size)
    cache = getattr(host, "_document_text_cache", None)
    if cache is None:
        cache = {}
        host._document_text_cache = cache
    entry = cache.pop(key, None)
    if entry is None:
        entry = extract_document_text(real)
    cache[key] = entry
    while len(cache) > _CACHE_ENTRIES:
        cache.pop(next(iter(cache)))
    return entry


def read_document(
    host: Any, path: str, offset: int = 0, limit: Optional[int] = None
) -> Dict[str, Any]:
    """One part of a document's extracted text, and where the next part starts.

    The caller has already checked that *path* may be read. A part is at most
    the host's read budget; ``next_offset`` is ``None`` once the end is reached.
    """
    kind = document_type(path)
    if (
        not isinstance(offset, int)
        or isinstance(offset, bool)
        or offset < 0
        or (limit is not None and (not isinstance(limit, int) or limit < 1))
    ):
        return {
            "status": "error",
            "error": "offset must be 0 or more and limit 1 or more characters.",
        }
    try:
        text, meta = cached_document_text(host, path)
    except (ImportError, ValueError) as e:
        return {"status": "error", "file_path": path, "error": str(e)}

    total = len(text)
    if offset and offset >= total:
        return {
            "status": "error",
            "error": f"offset {offset} is past the end: the text has {total} characters.",
        }
    budget = read_budget(host)
    size = budget if limit is None else min(limit, budget)
    end = min(total, offset + size)
    if end < total:
        # End the part on a line break when one is near, so no line is split.
        cut = text.rfind("\n", offset + size * 4 // 5, end)
        if cut > offset:
            end = cut + 1

    result: Dict[str, Any] = {
        "status": "success",
        "file_path": path,
        "file_type": kind,
        "content": text[offset:end],
        "offset": offset,
        "total_chars": total,
        "next_offset": end if end < total else None,
        **meta,
    }
    if end < total:
        result["note"] = (
            f"Partial: characters {offset:,}-{end:,} of {total:,}. Call read_file "
            f"with offset={end} for the next part. Read every part before "
            "comparing documents or saying something is not in this one."
        )
    if offset == 0 and total > FULL_READ_MAX_READS * budget:
        result["hint"] = (
            f"Long document: about {math.ceil(total / size)} reads in full. To look "
            "up specific facts, index_document + query_specific_file is cheaper; "
            "read it all only when the task needs every line."
        )
    return result


def text_page_notes(
    page: Dict[str, Any], size_bytes: int, line_ranges: bool = True
) -> Dict[str, Any]:
    """Say how much of the file a character page holds, and where the rest is."""
    returned = len(page.get("content") or "")
    start = page.get("offset") or 0
    page = {**page, "chars_returned": returned, "size_bytes": size_bytes}
    if page.get("next_offset") is not None:
        by_line = (
            ", or read by line number with start_line/end_line" if line_ranges else ""
        )
        page["note"] = (
            f"Partial: characters {start:,}-{start + returned:,} of a "
            f"{size_bytes:,}-byte file. Continue with "
            f"offset={page['next_offset']}{by_line}."
        )
    return page


def with_small_limit_warning(
    result: Dict[str, Any], limit: Optional[int], line_ranges: bool = True
) -> Dict[str, Any]:
    """A tiny ``limit`` is usually a line count; say it counts characters."""
    if (
        result.get("status") != "success"
        or not isinstance(limit, int)
        or isinstance(limit, bool)
        or limit >= SMALL_LIMIT_CHARS
    ):
        return result
    returned = len(result.get("content") or "")
    by_line = "To read lines, use start_line/end_line instead. " if line_ranges else ""
    return {
        **result,
        "warning": (
            f"limit={limit} counts characters, not lines, so only {returned} "
            f"characters came back. {by_line}For more text per read, raise "
            "limit or leave it out."
        ),
    }
