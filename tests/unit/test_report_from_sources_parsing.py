# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""Each parsing stage behind the report-from-sources tools, on its own.

Period labels, cell numbers, the CSV/JSON/xlsx readers, layout detection,
label matching, file-name vintages and the bounded-memory paths on a large
generated file. All data is generated here.
"""

from __future__ import annotations

import csv
import json
import sys
import time
from datetime import date, datetime
from pathlib import Path

import pytest

pytest.importorskip("openpyxl")

_SKILL = Path(__file__).resolve().parents[2] / "hub" / "skills" / "report-from-sources"
if str(_SKILL) not in sys.path:
    sys.path.insert(0, str(_SKILL))

from report_from_sources_lib import periods as P  # noqa: E402
from report_from_sources_lib import series as S  # noqa: E402
from report_from_sources_lib import tables as T  # noqa: E402

# --------------------------------------------------------------------------- #
# Periods
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    "text,label",
    [
        ("2026Q2", "2026Q2"),
        ("2026 Q2", "2026Q2"),
        ("2026-Q2", "2026Q2"),
        ("Q2 2026", "2026Q2"),
        ("Q2-26", "2026Q2"),
        ("Q2'26", "2026Q2"),
        ("2Q26", "2026Q2"),
        ("2Q 2026", "2026Q2"),
        ("Second quarter 2026", "2026Q2"),
        ("2nd quarter of 2026", "2026Q2"),
        ("2026Q2(p)", "2026Q2"),
        ("2026Q2*", "2026Q2"),
        ("Jun-2026", "2026-06"),
        ("June 2026", "2026-06"),
        ("Jun 26", "2026-06"),
        ("Sept. 2025", "2025-09"),
        ("2026-06", "2026-06"),
        ("2026/06", "2026-06"),
        ("2026M06", "2026-06"),
        ("06/2026", "2026-06"),
        ("2026-06-30", "2026-06-30"),
        ("2026-06-30T00:00:00Z", "2026-06-30"),
        ("06/30/2026", "2026-06-30"),
        ("30/06/2026", "2026-06-30"),
        ("30 Jun 2026", "2026-06-30"),
        ("June 30, 2026", "2026-06-30"),
        ("FY2026", "FY2026"),
        ("FY26", "FY2026"),
        ("FY 2025/26", "FY2026"),
        ("FY2025-2026", "FY2026"),
        ("2025-26", "FY2026"),
        ("2025/2026", "FY2026"),
        ("Fiscal year 2026", "FY2026"),
        ("FY26 Q1", "FY2026Q1"),
        ("Q3 FY2026", "FY2026Q3"),
        ("H1 2026", "2026H1"),
        ("2026H2", "2026H2"),
        ("1H26", "2026H1"),
        ("2026", "2026"),
        ("CY2026", "2026"),
    ],
)
def test_period_labels(text, label):
    period = P.parse_period(text)
    assert period is not None, text
    assert period.label == label


@pytest.mark.parametrize(
    "text",
    ["foo", "Q5 2026", "2026-13", "13/13/2026", "2025-27", "12ab", "", "Revenue 2026"],
)
def test_non_periods(text):
    assert P.parse_period(text) is None


def test_bare_numbers_are_years_only_when_asked():
    assert P.parse_period(2026) is None
    assert P.parse_period(2026, bare_year_numbers=True).label == "2026"
    assert P.parse_period(2026.5, bare_year_numbers=True) is None
    assert P.parse_period(datetime(2026, 6, 30)).label == "2026-06-30"
    assert P.parse_period(date(2026, 6, 30)).label == "2026-06-30"


def test_positions_order_across_frequencies():
    q2 = P.parse_period("2026Q2")
    jun = P.parse_period("Jun 2026")
    year = P.parse_period("2026")
    q4 = P.parse_period("Q4 2026")
    assert q2.position == jun.position
    assert year.position == q4.position
    assert P.parse_period("2025Q4").sort_key < q2.sort_key < year.sort_key


def test_axis_day_month_order_is_decided_once():
    periods, notes = P.parse_axis(["03/04/2026", "13/04/2026", "20/04/2026"])
    assert [p.label for p in periods] == ["2026-04-03", "2026-04-13", "2026-04-20"]
    periods, notes = P.parse_axis(["03/04/2026", "04/04/2026"])
    assert any("month-first" in n for n in notes)


def test_axis_harmonizes_mixed_month_formats():
    values = [datetime(2024, 1, 1), "2024-02-28", "Mar 2024", "04/15/2024", "May 2024"]
    periods, notes = P.parse_axis(values)
    assert [p.label for p in periods] == [
        "2024-01",
        "2024-02",
        "2024-03",
        "2024-04",
        "2024-05",
    ]
    assert notes


def test_axis_reads_month_start_dates_as_months_and_quarter_ends_as_quarters():
    periods, _ = P.parse_axis(["2025-01-01", "2025-02-01", "2025-03-01"])
    assert [p.label for p in periods] == ["2025-01", "2025-02", "2025-03"]
    periods, _ = P.parse_axis(["2025-03-31", "2025-06-30", "2025-09-30"])
    assert [p.label for p in periods] == ["2025Q1", "2025Q2", "2025Q3"]
    daily, _ = P.parse_axis(["2025-01-01", "2025-01-02", "2025-01-03"])
    assert [p.label for p in daily] == ["2025-01-01", "2025-01-02", "2025-01-03"]


@pytest.mark.parametrize(
    "spec,kind",
    [
        ("last 5", "last"),
        ("5", "last"),
        (5, "last"),
        ("all", "all"),
        ("2024Q1:2026Q2", "range"),
        ("2023..2025", "range"),
        ("2025Q4, 2026Q1", "list"),
        (["2025Q4"], "list"),
        ("", "last"),
    ],
)
def test_period_specs(spec, kind):
    assert P.parse_period_spec(spec).kind == kind


@pytest.mark.parametrize("spec", ["last banana", "0", "Q9:2026"])
def test_bad_period_specs_say_what_is_accepted(spec):
    with pytest.raises(P.PeriodSpecError):
        P.parse_period_spec(spec)


def test_select_periods():
    avail = [
        P.parse_period(f"{y}Q{q}") for y in (2024, 2025, 2026) for q in (1, 2, 3, 4)
    ]
    picked, _ = P.select_periods(avail, P.parse_period_spec("last 3"))
    assert [p.label for p in picked] == ["2026Q2", "2026Q3", "2026Q4"]
    picked, _ = P.select_periods(avail, P.parse_period_spec("2025Q3:2026Q1"))
    assert [p.label for p in picked] == ["2025Q3", "2025Q4", "2026Q1"]
    picked, _ = P.select_periods(avail, P.parse_period_spec("2025..2025"))
    assert [p.label for p in picked] == ["2025Q1", "2025Q2", "2025Q3", "2025Q4"]
    picked, missing = P.select_periods(avail, P.parse_period_spec("2026Q1, 2027Q1"))
    assert [p.label for p in picked] == ["2026Q1"] and missing == ["2027Q1"]


# --------------------------------------------------------------------------- #
# Cell numbers
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    "raw,value,marker",
    [
        (123.4, 123.4, None),
        (7, 7, None),
        ("123.4", 123.4, None),
        ("123.4(a)", 123.4, "(a)"),
        ("123.4 (b)", 123.4, "(b)"),
        ("123.4*", 123.4, "*"),
        ("123.4p", 123.4, "p"),
        ("1,234.5", 1234.5, None),
        ("(12.5)", -12.5, None),
        ("-3", -3, None),
        ("−3.5", -3.5, None),
        ("$1,200", 1200, None),
        ("4.5%", 4.5, None),
        ("1.5e3", 1500.0, None),
    ],
)
def test_read_number(raw, value, marker):
    num = T.read_number(raw)
    assert num.value == value and type(num.value) is type(value)
    assert num.marker == marker


@pytest.mark.parametrize(
    "raw", [None, "", "-", "..", "n/a", "NA", "x", "—", float("nan")]
)
def test_blank_cells(raw):
    num = T.read_number(raw)
    assert num.value is None and num.blank


@pytest.mark.parametrize("raw", ["abc", "(12.5", "Q2 2026", True, "12ab"])
def test_text_is_not_a_number(raw):
    assert T.read_number(raw).value is None


def test_percent_and_scale_suffix_are_reported_not_applied():
    assert T.read_number("4.5%").percent
    num = T.read_number("12.3bn")
    assert num.value == 12.3 and num.scale_suffix == "bn"


@pytest.mark.parametrize(
    "units,scale",
    [
        ("$ millions", 1e6),
        ("USD '000", 1e3),
        ("thousands of people", 1e3),
        ("bn", 1e9),
        ("%", None),
        (None, None),
    ],
)
def test_units_scale(units, scale):
    assert T.units_scale(units) == scale


# --------------------------------------------------------------------------- #
# Readers
# --------------------------------------------------------------------------- #


def _rows(path):
    with T.open_tables(path) as tables:
        return [(n, list(v)) for n, v in tables[0].rows()]


def test_csv_with_bom_crlf_and_unicode(tmp_path):
    path = tmp_path / "données.csv"
    path.write_bytes(
        "﻿date,série,value\r\n2026-01,Café,1.5\r\n2026-02,Café,2.5\r\n".encode("utf-8")
    )
    rows = _rows(path)
    assert rows[0] == (1, ["date", "série", "value"])
    assert rows[2] == (3, ["2026-02", "Café", "2.5"])


def test_csv_semicolon_and_cp1252(tmp_path):
    path = tmp_path / "eu.csv"
    path.write_bytes(
        "Monat;Umsatz\r\nJan 2026;10\r\nFeb 2026;12\r\nMär;x\r\n".encode("cp1252")
    )
    rows = _rows(path)
    assert rows[1][1] == ["Jan 2026", "10"]
    assert rows[3][1][0] == "Mär"


def test_tsv_and_multiline_quoted_cell(tmp_path):
    path = tmp_path / "t.tsv"
    path.write_bytes(b'period\tnote\tvalue\n2026Q1\t"two\nlines"\t3\n2026Q2\tok\t4\n')
    rows = _rows(path)
    assert rows[1][1] == ["2026Q1", "two\nlines", "3"]
    assert rows[2][1][2] == "4"


def test_empty_files(tmp_path):
    (tmp_path / "empty.csv").write_text("", encoding="utf-8")
    assert _rows(tmp_path / "empty.csv") == []
    with T.open_tables(tmp_path / "empty.csv") as tables:
        assert T.analyze_layout(tables[0]).kind == "unknown"
    inv = S.inventory_folder(tmp_path)
    assert inv["data_files"][0]["sheets"][0]["layout"] == "unknown"


def test_unreadable_formats_say_what_to_do(tmp_path):
    (tmp_path / "old.xls").write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(T.DataFileError, match="save it as .xlsx"):
        with T.open_tables(tmp_path / "old.xls"):
            pass
    (tmp_path / "broken.xlsx").write_bytes(b"not a zip")
    with pytest.raises(T.DataFileError):
        with T.open_tables(tmp_path / "broken.xlsx"):
            pass
    inv = S.inventory_folder(tmp_path)
    errors = {f["file"]: f.get("error") for f in inv["data_files"]}
    assert errors["old.xls"] and errors["broken.xlsx"]


def test_json_shapes(tmp_path):
    records = tmp_path / "r.json"
    records.write_text(
        json.dumps(
            {"data": [{"month": "2026-01", "v": 1}, {"month": "2026-02", "v": 2}]}
        ),
        encoding="utf-8",
    )
    assert _rows(records)[0][1] == ["month", "v"]
    mapping = tmp_path / "m.json"
    mapping.write_text(
        json.dumps({"Sales": {"2025": 1, "2026": 2}, "Costs": {"2025": 3, "2026": 4}}),
        encoding="utf-8",
    )
    result = S.extract_series(mapping, ["Costs"], periods="all")
    assert result["series"][0]["values"] == [3, 4]
    bad = tmp_path / "scalar.json"
    bad.write_text("42", encoding="utf-8")
    with pytest.raises(T.DataFileError, match="not a table"):
        list(T._JsonTable(bad).rows())


def test_json_lines_skips_and_reports_malformed_lines(tmp_path):
    path = tmp_path / "obs.jsonl"
    lines = [json.dumps({"date": f"2026-0{m}-01", "value": m}) for m in range(1, 7)]
    lines.insert(2, "{not json")
    path.write_bytes(("\r\n".join(lines) + "\r\n").encode("utf-8"))
    result = S.extract_series(path, ["value"], periods="last 3")
    assert result["periods"] == ["2026-04", "2026-05", "2026-06"]
    assert result["series"][0]["values"] == [4, 5, 6]
    assert any("line 3" in n for n in result["notes"])


def test_xlsx_row_numbers_survive_missing_rows(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Series"
    ws["B1"] = "2025"
    ws["C1"] = "2026"
    ws["A9"] = "Output"
    ws["B9"] = 1.0
    ws["C9"] = 2.0
    wb.save(tmp_path / "gap.xlsx")
    result = S.extract_series(tmp_path / "gap.xlsx", ["Output"], periods="all")
    assert result["series"][0]["cells"] == ["B9", "C9"]


def test_merged_ranges_are_read_from_the_sheet_xml(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.merge_cells("B1:E1")
    ws.merge_cells("A10:C12")
    wb.save(tmp_path / "m.xlsx")
    with T.open_tables(tmp_path / "m.xlsx") as tables:
        assert sorted(tables[0].merged) == [(1, 1, 1, 4), (10, 0, 12, 2)]


# --------------------------------------------------------------------------- #
# Layout
# --------------------------------------------------------------------------- #


def _sheet(tmp_path, rows, merges=(), name="t.xlsx"):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in rows:
        ws.append(row)
    for m in merges:
        ws.merge_cells(m)
    wb.save(tmp_path / name)
    return tmp_path / name


def _layout(path):
    with T.open_tables(path) as tables:
        return T.analyze_layout(tables[0])


def test_wide_with_year_row_over_quarters(tmp_path):
    path = _sheet(
        tmp_path,
        [
            ["Title"],
            [],
            [None, 2025, None, None, None, 2026, None],
            ["Item", "Q1", "Q2", "Q3", "Q4", "Q1", "Q2"],
            ["A", 1, 2, 3, 4, 5, 6],
        ],
        merges=["B3:E3", "F3:G3"],
    )
    lay = _layout(path)
    assert lay.kind == "wide"
    assert lay.header_rows == [3, 4]
    assert [p.label for _, p in sorted(lay.period_cols.items())] == [
        "2025Q1",
        "2025Q2",
        "2025Q3",
        "2025Q4",
        "2026Q1",
        "2026Q2",
    ]
    assert lay.merged_header_cells == 2


def test_wide_year_written_once_without_merge(tmp_path):
    path = _sheet(
        tmp_path,
        [
            [None, "2025", None, "2026", None],
            [None, "H1", "H2", "H1", "H2"],
            ["x", 1, 2, 3, 4],
        ],
    )
    lay = _layout(path)
    assert [p.label for _, p in sorted(lay.period_cols.items())] == [
        "2025H1",
        "2025H2",
        "2026H1",
        "2026H2",
    ]


def test_wide_annual_total_columns_are_kept_apart(tmp_path):
    path = _sheet(
        tmp_path, [["Item", "2025Q3", "2025Q4", "2025", "2026Q1"], ["A", 1, 2, 3, 4]]
    )
    lay = _layout(path)
    assert lay.freq == "Q"
    assert [p.label for p in lay.other_period_cols.values()] == ["2025"]
    result = S.extract_series(path, ["A"], periods="all")
    assert result["periods"] == ["2025Q3", "2025Q4", "2026Q1"] and result["series"][0][
        "values"
    ] == [1, 2, 4]
    # Asking for the year reads the annual column the sheet already has.
    annual = S.extract_series(path, ["A"], periods="2025")
    assert annual["periods"] == ["2025"] and annual["series"][0]["values"] == [3]


def test_columns_layout(tmp_path):
    path = _sheet(
        tmp_path,
        [
            ["Monthly report"],
            ["Month", "Sales", "Costs"],
            ["Jan 2026", 1, 2],
            ["Feb 2026", 3, 4],
            ["Mar 2026", 5, 6],
        ],
    )
    lay = _layout(path)
    assert lay.kind == "columns" and lay.period_col == 0
    assert lay.title_lines == ["Monthly report"]
    assert sorted(lay.column_headers.values()) == ["Costs", "Sales"]


def test_stacked_layout_needs_filters_when_split(tmp_path):
    rows = [["date", "region", "metric", "value"]]
    for m in (1, 2, 3):
        for region in ("East", "West"):
            rows.append(
                [
                    f"2026-{m:02d}-01",
                    region,
                    "Sales",
                    m * (10 if region == "East" else 1),
                ]
            )
    path = _sheet(tmp_path, rows)
    assert _layout(path).kind == "stacked"
    with pytest.raises(S.SeriesLookupError, match="filters"):
        S.extract_series(path, ["Sales"], periods="all")
    result = S.extract_series(path, ["Sales"], periods="all", filters="region=West")
    assert result["series"][0]["values"] == [1, 2, 3]


def test_units_row_column_and_title(tmp_path):
    path = _sheet(
        tmp_path,
        [
            ["Output ($ millions)"],
            ["Item", "Units", "2025", "2026"],
            [None, None, "$m", "$m"],
            ["A", "index", 1, 2],
        ],
    )
    lay = _layout(path)
    assert lay.units_row == 3 and lay.units == "$m" and lay.units_col == 1
    result = S.extract_series(path, ["A"], periods="all")
    assert result["series"][0]["units"] == "index"


def test_notes_and_footnotes(tmp_path):
    path = _sheet(
        tmp_path,
        [
            ["Item", "2025", "2026"],
            ["A", "1.5(a)", 2],
            [],
            ["Notes:"],
            ["(a) Revised."],
            ["Source: survey"],
        ],
    )
    lay = _layout(path)
    assert lay.footnotes == {"a": "Revised."}
    result = S.extract_series(path, ["A"], periods="all")
    assert result["series"][0]["flags"][0]["footnote"] == "Revised."


# --------------------------------------------------------------------------- #
# Label matching
# --------------------------------------------------------------------------- #


def _cands(*labels, section=None):
    return [
        S.Candidate("S", lab, "row", section=section, row=i + 1, col=0)
        for i, lab in enumerate(labels)
    ]


def test_normalize_label():
    assert S.normalize_label("Net revenue ($ millions)(a)") == "net revenue"
    assert S.normalize_label("Core inflation (ex food)") == "core inflation ex food"
    assert S.normalize_label("R&D") == "r and d"


def test_exact_beats_contains():
    match, _ = S.resolve_label("Net sales", _cands("Net sales", "Net sales, adjusted"))
    assert match.candidate.label == "Net sales" and match.how == "exact"


def test_two_contains_are_ambiguous():
    match, problem = S.resolve_label("sales", _cands("Net sales", "Gross sales"))
    assert match is None and problem["status"] == "ambiguous"
    assert {c["label"] for c in problem["candidates"]} == {"Net sales", "Gross sales"}


def test_one_contains_is_taken():
    match, _ = S.resolve_label("revenue", _cands("Total revenue", "Costs"))
    assert match.candidate.label == "Total revenue" and match.how == "contains"


def test_fuzzy_needs_a_clear_winner():
    match, _ = S.resolve_label(
        "Manufacturng output", _cands("Manufacturing output", "Services output")
    )
    assert match.candidate.label == "Manufacturing output" and match.how == "fuzzy"
    match, problem = S.resolve_label("Headcount", _cands("Net sales", "Costs"))
    assert match is None and problem["status"] == "not_found" and problem["nearest"]


def test_section_path_resolves_duplicates():
    cands = _cands("Total", section="Revenue") + _cands("Total", section="Costs")
    match, problem = S.resolve_label("Total", cands)
    assert match is None and "Name the section" in problem["hint"]
    match, _ = S.resolve_label("Costs > Total", cands)
    assert match.candidate.section == "Costs"


# --------------------------------------------------------------------------- #
# Vintage
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    "stem,period,key",
    [
        ("regional-output-second-quarter-2025", "2025Q2", "regional output"),
        ("regional_output_2026Q2", "2026Q2", "regional output"),
        ("Regional Output Q2 2026", "2026Q2", "regional output"),
        ("sales_2024_q3", "2024Q3", "sales"),
        ("trade_monthly_202603", None, "trade monthly"),
        ("report-jun2026", "2026-06", "report"),
        ("budget_FY26_final", "FY2026", "budget"),
        ("prices_20260630", "2026-06-30", "prices"),
    ],
)
def test_filename_vintage(stem, period, key):
    got, _, got_key = S.filename_vintage(stem)
    assert (got.label if got else None) == period
    assert got_key == key


def test_revision_words_rank():
    assert S.filename_vintage("pack_v3")[1] == ("v3", 13)
    assert S.filename_vintage("pack_final")[1][0] == "final"
    assert (
        S.filename_vintage("pack_draft")[1][1] < S.filename_vintage("pack_final")[1][1]
    )


def test_content_as_of_beats_name_when_coverage_ties(tmp_path):
    for name, asof in (
        ("kpi_v2.xlsx", "As of March 31, 2026"),
        ("kpi_v1.xlsx", "As of June 30, 2026"),
    ):
        _sheet(
            tmp_path,
            [
                [asof],
                ["Item", "2025Q4", "2026Q1"],
                ["A", 1, 2],
                ["B", 3, 4],
                ["C", 5, 6],
            ],
            name=name,
        )
    inv = S.inventory_folder(tmp_path)
    ds = inv["datasets"][0]
    assert ds["newest"] == "kpi_v1.xlsx"
    assert "same periods" in ds["note"] and "as-of" in ds["note"]


def test_unrelated_files_are_not_merged(tmp_path):
    _sheet(
        tmp_path,
        [["Item", "2025", "2026"], ["Rainfall", 1, 2]],
        name="weather_2026.xlsx",
    )
    _sheet(
        tmp_path,
        [["Item", "2025", "2026"], ["Revenue", 1, 2]],
        name="finance_2025.xlsx",
    )
    inv = S.inventory_folder(tmp_path)
    assert len(inv["datasets"]) == 2 and all("note" not in d for d in inv["datasets"])


def test_inventory_lists_documents_and_skips_lock_files(tmp_path):
    (tmp_path / "memo.pdf").write_bytes(b"%PDF-1.4")
    (tmp_path / "~$book.xlsx").write_bytes(b"lock")
    (tmp_path / "image.png").write_bytes(b"png")
    inv = S.inventory_folder(tmp_path)
    assert inv["narrative_sources"][0]["file"] == "memo.pdf"
    assert inv["not_read"][0]["file"] == "~$book.xlsx"
    assert inv["other_files"] == ["image.png"]


def test_inventory_of_an_empty_folder(tmp_path):
    inv = S.inventory_folder(tmp_path)
    assert (
        inv["data_files"] == []
        and inv["datasets"] == []
        and "No data files" in inv["message"]
    )


# --------------------------------------------------------------------------- #
# Large input
# --------------------------------------------------------------------------- #


def test_large_long_csv_streams(tmp_path):
    """300k rows (12 series x 25k days): inventory and lookup stream it."""
    path = tmp_path / "big.csv"
    start = date(1960, 1, 1).toordinal()
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["day", "sensor", "reading"])
        for i in range(25_000):
            d = date.fromordinal(start + i).isoformat()
            for s in range(12):
                w.writerow([d, f"sensor {s:02d}", (i * 7 + s) % 1000 / 10])
    t0 = time.monotonic()
    inv = S.inventory_folder(tmp_path)
    sheet = inv["data_files"][0]["sheets"][0]
    assert sheet["layout"] == "stacked" and sheet["data_rows"] == 300_000
    assert sheet["periods"]["last"] == date.fromordinal(start + 24_999).isoformat()
    result = S.extract_series(path, ["sensor 07"], periods="last 3")
    assert result["series"][0]["values"] == [
        ((24_997 + k) * 7 + 7) % 1000 / 10 for k in range(3)
    ]
    assert time.monotonic() - t0 < 120
