# Copyright(C) 2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""verify_workbook finds the formula mistakes openpyxl never computes.

A local model built a workbook whose SUMIFS criteria walked across columns, a
SUM covered a text column and a variance pointed at the wrong cell; nothing
recalculated the file, so the agent reported it as done.
"""

import hashlib
import importlib.util
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch

import pytest

from gaia.agents.base.agent import Agent
from gaia.security import PathValidator
from gaia.skills.context import SkillContextError
from gaia.skills.manager import SkillManager

openpyxl = pytest.importorskip("openpyxl")

SKILL_DIR = Path(__file__).resolve().parents[2] / "hub" / "skills" / "office-docs"


def _load(name: str, filename: str):
    """Import one of the skill's modules from its folder, as the loader does."""
    spec = importlib.util.spec_from_file_location(name, SKILL_DIR / filename)
    module = importlib.util.module_from_spec(spec)
    sys.path.insert(0, str(SKILL_DIR))
    try:
        sys.modules[name] = module
        spec.loader.exec_module(module)
    finally:
        sys.path.remove(str(SKILL_DIR))
    return module


# Registered under the name tools.py imports it by, so both see one module.
workbook_check = _load("workbook_check", "workbook_check.py")
office_tools = _load("office_docs_tools_under_test", "tools.py")
_criterion_matches = workbook_check._criterion_matches
check_workbook = workbook_check.check_workbook
soffice_installed = workbook_check.soffice_installed
needs_libreoffice = pytest.mark.skipif(
    not soffice_installed(), reason="LibreOffice not installed"
)


def _static(path: Path) -> dict:
    return check_workbook(path)


def _problems(result: dict, bucket: str) -> dict:
    return {e["cell"]: e for e in result[bucket]}


@pytest.fixture
def data_book():
    """A data sheet plus an empty summary sheet to put formulas on."""
    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "Data"
    data.append(["Region", "Quarter", "Sales", "Cost"])
    for row in [("East", "Q1", 10, 4), ("West", "Q1", 20, 8), ("East", "Q2", 15, 5)]:
        data.append(row)
    summary = wb.create_sheet("Summary")
    summary.append(["Region", "Sales"])
    summary.append(["East"])
    return wb


def _save(wb, tmp_path, name="book.xlsx") -> Path:
    path = tmp_path / name
    wb.save(path)
    return path


def test_correct_formulas_pass(data_book, tmp_path):
    ws = data_book["Summary"]
    ws["B2"] = "=SUMIFS(Data!C:C,Data!A:A,A2)"
    ws["B3"] = "=SUM(Data!C2:C4)"
    ws["B4"] = "=B2/B3"
    result = _static(_save(data_book, tmp_path))

    assert result["error_count"] == 0, result["errors"]
    assert result["warning_count"] == 0, result["warnings"]
    assert result["formula_count"] == 3


def test_sumifs_criteria_on_the_wrong_column_is_flagged(data_book, tmp_path):
    data_book["Summary"]["B2"] = "=SUMIFS(Data!D:D,Data!B:B,A2)"
    result = _static(_save(data_book, tmp_path))

    warning = _problems(result, "warnings")["Summary!B2"]
    assert "matches no cell in Data!B:B" in warning["problem"]


def test_sum_over_a_text_column_is_an_error(data_book, tmp_path):
    data_book["Summary"]["B2"] = "=SUM(Data!A2:A4)"
    result = _static(_save(data_book, tmp_path))

    assert "finds only text" in _problems(result, "errors")["Summary!B2"]["problem"]


def test_sum_over_empty_cells_is_an_error(data_book, tmp_path):
    data_book["Summary"]["B2"] = "=SUM(Data!F2:F4)"
    result = _static(_save(data_book, tmp_path))

    assert "only empty cells" in _problems(result, "errors")["Summary!B2"]["problem"]


def test_a_header_above_the_numbers_is_not_a_text_warning(data_book, tmp_path):
    data_book["Summary"]["B2"] = "=SUM(Data!C:C)"
    result = _static(_save(data_book, tmp_path))

    assert result["warning_count"] == 0 and result["error_count"] == 0


def test_numbers_stored_as_text_among_numbers_warn(data_book, tmp_path):
    data_book["Data"]["C5"] = "1,200"
    data_book["Data"]["C6"] = 7
    data_book["Summary"]["B2"] = "=SUM(Data!C2:C6)"
    result = _static(_save(data_book, tmp_path))

    assert "skips 1 text cell" in _problems(result, "warnings")["Summary!B2"]["problem"]


def test_arithmetic_on_an_empty_cell_warns(data_book, tmp_path):
    data_book["Summary"]["B2"] = "=Data!C2-Data!C9"
    result = _static(_save(data_book, tmp_path))

    assert (
        "Data!C9, which is empty"
        in _problems(result, "warnings")["Summary!B2"]["problem"]
    )


@pytest.mark.parametrize(
    "formula,code",
    [
        ("=Data!C2/Data!F2", "#DIV/0!"),
        ("=Data!A2*2", "#VALUE!"),
        ("=Missing!A1", "#REF!"),
        ("=#REF!+1", "#REF!"),
        ("=Bogus*2", "#NAME?"),
        ("=FOOBAR(1)", "#NAME?"),
        ("=STDEV.S(Data!C2:C4)", "#NAME?"),
        ('=SUMIFS(Data!C2:C4,Data!A2:A3,"East")', "#VALUE!"),
    ],
)
def test_static_errors_name_the_excel_error(data_book, tmp_path, formula, code):
    data_book["Summary"]["B2"] = formula
    result = _static(_save(data_book, tmp_path))

    assert _problems(result, "errors")["Summary!B2"]["error"] == code


def test_prefixed_newer_functions_are_accepted(data_book, tmp_path):
    data_book["Summary"]["B2"] = "=_xlfn.STDEV.S(Data!C2:C4)"
    result = _static(_save(data_book, tmp_path))

    assert result["error_count"] == 0


def test_defined_names_resolve(data_book, tmp_path):
    from openpyxl.workbook.defined_name import DefinedName

    data_book.defined_names["Sales"] = DefinedName("Sales", attr_text="Data!$C$2:$C$4")
    data_book["Summary"]["B2"] = "=SUM(Sales)"
    result = _static(_save(data_book, tmp_path))

    assert result["error_count"] == 0, result["errors"]


def test_circular_references_are_reported_with_their_cells(data_book, tmp_path):
    ws = data_book["Summary"]
    ws["B2"] = "=B3+1"
    ws["B3"] = "=B2+1"
    ws["C2"] = "=SUM(C1:C3)"
    result = _static(_save(data_book, tmp_path))

    cycles = sorted(sorted(c) for c in result["circular_references"])
    assert cycles == [["Summary!B2", "Summary!B3"], ["Summary!C2"]]
    assert result["error_count"] == 2


@pytest.mark.parametrize(
    "criterion,value,expected",
    [
        ("East", "east", True),
        ("East", "West", False),
        (">=10", 12, True),
        (">=10", 9, False),
        ("<>East", "West", True),
        ("E*", "East", True),
        (10.0, 10, True),
        (10.0, "10", True),
    ],
)
def test_criterion_matching_follows_excel(criterion, value, expected):
    assert _criterion_matches(criterion, value) is expected


def test_a_failed_recalculation_is_reported_not_hidden(data_book, tmp_path):
    def broken(_path):
        raise RuntimeError("LibreOffice exited with code 1")

    data_book["Summary"]["B2"] = "=SUM(Data!C2:C4)"
    result = check_workbook(_save(data_book, tmp_path), recalculator=broken)

    assert result["recalculated"] is False
    assert "exited with code 1" in result["recalc_error"]


def test_cached_values_from_excel_are_used_when_not_recalculated(tmp_path, monkeypatch):
    """A file Excel saved carries computed values; they are read, not ignored."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = 2
    wb.active["A2"] = "=A1*2"
    path = _save(wb, tmp_path)
    monkeypatch.setattr(
        workbook_check,
        "_computed_values",
        lambda _p: {"Sheet": {(1, 1): 2, (2, 1): "#DIV/0!"}},
    )
    result = check_workbook(path)

    assert "cached values" in result["values_source"]
    assert _problems(result, "errors")["Sheet!A2"]["error"] == "#DIV/0!"


def _with_cached_values(path: Path, cached: dict) -> Path:
    """A copy of *path* whose formula cells carry stale cached values.

    This is what an Excel-saved file looks like after openpyxl edits a cell a
    formula depends on: the old result stays in each formula's ``<v>``.
    """
    stale = path.with_name("stale.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(stale, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                for cell, value in cached.items():
                    xml, n = re.subn(
                        rf'(<c r="{cell}"[^>]*>\s*<f>[^<]*</f>)(\s*<v>[^<]*</v>|\s*<v\s*/>)?',
                        rf"\g<1><v>{value}</v>",
                        xml,
                    )
                    assert n == 1, (cell, xml)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    return stale


@needs_libreoffice
def test_libreoffice_recomputes_values_an_earlier_save_cached(agent, folders):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 5
    ws["B1"] = "=A1*2"
    ws["A2"] = 0
    ws["B2"] = "=1/A2"
    path = _with_cached_values(_save(wb, folders[0]), {"B1": 999, "B2": 5})

    result = _verify(agent, path)

    assert result["recalculated"] is True, result
    assert result["computed_values"]["Sheet!B1"] == 10
    assert _problems(result, "errors")["Sheet!B2"]["error"] == "#DIV/0!"


# ── the tool, run by an agent ───────────────────────────────────────────────


class _HostAgent(Agent):
    def _get_system_prompt(self) -> str:
        return "test"

    def _register_tools(self) -> None:
        pass

    def _create_console(self):
        from gaia.agents.base.console import AgentConsole

        return AgentConsole()


@pytest.fixture
def folders(tmp_path):
    """``books`` is in the agent's scope, ``elsewhere`` is not."""
    books, elsewhere = tmp_path / "books", tmp_path / "elsewhere"
    books.mkdir()
    elsewhere.mkdir()
    return books, elsewhere


@pytest.fixture
def gaia_home(monkeypatch):
    """A short GAIA_HOME: LibreOffice crashes on Windows when its profile path
    is as deep as pytest's tmp_path makes it."""
    home = Path(tempfile.mkdtemp(prefix="gh-"))
    monkeypatch.setenv("GAIA_HOME", str(home))
    yield home
    shutil.rmtree(home, ignore_errors=True)


@pytest.fixture
def agent(tmp_path, folders, gaia_home):
    """An agent allowed into ``books`` only, with office-docs loaded as bundled."""
    root = tmp_path / "bundled"
    shutil.copytree(SKILL_DIR, root / "office-docs")
    with patch("gaia.agents.base.agent.AgentSDK"):
        built = _HostAgent(silent_mode=True, skip_lemonade=True)
    validator = PathValidator(allowed_paths=[str(folders[0])])
    validator.set_access_prompt(lambda path: False)  # the user says no
    built.path_validator = validator
    built._instance_tools = {}
    manager = SkillManager(
        agent_skill_dirs=[root],
        user_skills_root=tmp_path / "no-user-skills",
        include_claude_roots=False,
    )
    built.load_skill("office-docs", manager=manager)
    yield built
    for name in list(built.loaded_skills):
        built.unload_skill(name)


def _verify(agent, path) -> dict:
    return agent._execute_tool("office-docs/verify_workbook", {"file_path": str(path)})


@pytest.fixture
def no_libreoffice(monkeypatch):
    monkeypatch.setattr(workbook_check, "soffice_installed", lambda: False)


def test_tool_checks_a_workbook(agent, folders, data_book, no_libreoffice):
    data_book["Summary"]["B2"] = "=SUM(Data!A2:A4)"
    result = _verify(agent, _save(data_book, folders[0]))

    assert result["status"] == "success"
    assert result["error_count"] == 1
    assert result["recalculated"] is False
    assert "LibreOffice is not installed" in result["recalc_note"]


def test_tool_refuses_a_workbook_outside_scope(agent, folders, data_book):
    path = _save(data_book, folders[1])
    with patch.object(workbook_check, "check_workbook") as check:
        result = _verify(agent, path)

    assert result["status"] == "error"
    assert "not in allowed paths" in result["error"]
    check.assert_not_called()


def test_tool_checks_a_workbook_in_a_folder_approved_this_session(
    agent, folders, data_book, no_libreoffice
):
    agent.path_validator.add_allowed_path(str(folders[1]))
    data_book["Summary"]["B2"] = "=SUM(Data!C2:C4)"

    result = _verify(agent, _save(data_book, folders[1]))

    assert result["status"] == "success"
    assert result["error_count"] == 0


def test_tool_rejects_other_formats_and_broken_files(agent, folders, no_libreoffice):
    books = folders[0]
    (books / "notes.csv").write_text("a,b\n", encoding="utf-8")
    (books / "broken.xlsx").write_bytes(b"not a zip")

    assert ".xlsx and .xlsm" in _verify(agent, books / "notes.csv")["error"]
    assert "Could not open" in _verify(agent, books / "broken.xlsx")["error"]
    missing = _verify(agent, books / "missing.xlsx")
    assert missing["status"] == "error"
    assert "No file at" in missing["error"]


def test_recalculation_runs_only_inside_the_skill_call(data_book, tmp_path):
    """Outside a tool call there is no agent scope to check the copy against."""
    with pytest.raises(SkillContextError):
        workbook_check.recalculate(_save(data_book, tmp_path))


def test_the_skill_declares_what_the_recalculation_needs():
    from gaia.skills.audit import audit_skill

    skill = SkillManager(
        agent_skill_dirs=[SKILL_DIR.parent], include_claude_roots=False
    ).load("office-docs")
    declared = {str(p) for p in skill.parsed_permissions()}
    assert {"program:run:soffice", "scratch:write"} <= declared
    assert audit_skill(SKILL_DIR).verdict == "ALLOW"


@needs_libreoffice
def test_libreoffice_recalculates_a_copy_in_the_calls_temp_folder(
    agent, folders, data_book, gaia_home
):
    ws = data_book["Summary"]
    ws["B2"] = "=SUMIFS(Data!C:C,Data!A:A,A2)"
    ws["B3"] = "=IF(ISNUMBER(A2),A2*2,0)"
    path = _save(data_book, folders[0])
    before = hashlib.sha256(path.read_bytes()).hexdigest()

    result = _verify(agent, path)

    assert result["recalculated"] is True, result
    assert result["computed_values"]["Summary!B2"] == 25
    # The static #VALUE! guess is overruled by the computed value.
    assert "Summary!B3" not in _problems(result, "errors")
    assert "Summary!B3" in _problems(result, "warnings")
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
    temp_root = gaia_home / "skill-temp" / "office-docs"
    assert not list(temp_root.iterdir()), "the call's temporary folder was kept"


def test_the_skill_registers_verify_workbook_under_its_namespace(tmp_path):
    """The manifest matches tools.py, so the loader registers the tool."""
    from gaia.agents.base.tools import _TOOL_REGISTRY
    from gaia.skills.loader import register_skill_tools, unregister_skill_tools

    root = tmp_path / "skills"
    shutil.copytree(SKILL_DIR, root / "office-docs")
    manager = SkillManager(
        user_skills_root=root, claude_skill_dirs=[tmp_path / "claude-skills"]
    )
    before = dict(_TOOL_REGISTRY)
    try:
        registered = register_skill_tools(manager.load("office-docs"))

        assert set(registered) == {"office-docs/verify_workbook"}
        entry = registered["office-docs/verify_workbook"]
        assert entry["parameters"]["file_path"]["required"] is True
    finally:
        unregister_skill_tools("office-docs")
        _TOOL_REGISTRY.clear()
        _TOOL_REGISTRY.update(before)
