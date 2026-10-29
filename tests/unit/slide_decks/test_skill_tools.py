# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""The skill's manifest, tools.py and audit agree; the tools run inside a bound call.

This is the core-lane twin of the flagship test in
``hub/agents/gaia/python/tests/test_slide_decks_skill.py``: no agent package
needed. Calls are bound with :func:`gaia.skills.context.bind` to a stand-in
agent with no path validator (library use, no sandbox declared), and
GAIA_HOME is redirected so preference writes land under tmp_path.
"""

from __future__ import annotations

import json

import pytest

from gaia.skills import parse_skill_file, validate_skill
from gaia.skills.audit import audit_skill
from gaia.skills.context import bind
from gaia.skills.loader import register_skill_tools, unregister_skill_tools

from .conftest import SKILL_DIR

TOOL_NAMES = {
    "build_deck",
    "build_deck_from_markdown",
    "verify_deck",
    "add_slide",
    "update_slide",
    "list_blueprints",
    "outline_from_blueprint",
    "deck_outline_from_markdown",
    "update_deck_prefs",
    "export_slide_image",
}


class _Host:
    """Enough of an agent for gaia.skills.context: not private, no validator."""

    _incognito = False


@pytest.fixture
def tools(monkeypatch, tmp_path):
    from gaia.agents.base.tools import _TOOL_REGISTRY

    monkeypatch.setenv("GAIA_HOME", str(tmp_path / "home"))
    skill = parse_skill_file(SKILL_DIR)
    before = dict(_TOOL_REGISTRY)
    registered = register_skill_tools(skill)
    try:
        yield skill, {
            k.split("/", 1)[1]: v["function"] for k, v in registered.items()
        }, registered
    finally:
        unregister_skill_tools(skill.name)
        _TOOL_REGISTRY.clear()
        _TOOL_REGISTRY.update(before)


def test_manifest_validates_and_declares_every_tool(tools):
    skill, fns, registered = tools
    validate_skill(skill, source=str(SKILL_DIR))
    assert set(fns) == TOOL_NAMES
    assert {t.name for t in skill.gaia.tools} == TOOL_NAMES
    assert (
        registered["slide-decks/build_deck"]["parameters"]["outline"]["type"]
        == "object"
    )
    assert (
        registered["slide-decks/build_deck"]["parameters"]["output_path"]["required"]
        is False
    )
    perms = {str(p) for p in skill.parsed_permissions()}
    assert {"program:run:soffice", "scratch:write"} <= perms


def test_the_skill_code_passes_the_bundled_audit():
    report = audit_skill(SKILL_DIR)
    assert report.verdict == "ALLOW", [
        (f.severity, f.location, f.message) for f in report.findings
    ]
    assert not [
        f for f in report.findings if f.severity in ("medium", "high", "critical")
    ]


def test_tools_build_verify_edit_and_save_prefs(tools, tmp_path):
    skill, fns, _r = tools
    target = tmp_path / "deck.pptx"
    outline = {
        "slides": [
            {"type": "title", "title": "Pine Ridge school district update"},
            {
                "type": "bullets",
                "title": "Enrolment rose 3% for a second year",
                "bullets": ["Up 3%", "Two schools full"],
            },
        ]
    }
    with bind(_Host(), skill):
        built = fns["build_deck"](outline=outline, output_path=str(target))
        assert built["status"] == "ok", built
        report = fns["verify_deck"](path=str(target), expected_slides=2, render=False)
        assert report["status"] == "pass", report["issues"]
        added = fns["add_slide"](
            deck_path=str(target), slide={"type": "section", "title": "Budget"}
        )
        assert added["slides"] == 3
        replaced = fns["update_slide"](
            deck_path=str(target),
            index=3,
            slide={"type": "section", "title": "Budget and staffing"},
        )
        assert replaced["slides"] == 3
        saved = fns["update_deck_prefs"](feedback="use our green #2E7D32")
        assert saved["saved"] and saved["preferences"]["colors"]["primary"] == "2E7D32"
        again = fns["build_deck"](
            outline=outline, output_path=str(tmp_path / "branded.pptx")
        )
        assert again["status"] == "ok"
    assert (tmp_path / "home" / "skill-preferences" / "slide-decks.json").is_file()


def test_tool_errors_are_returned_not_raised(tools, tmp_path):
    skill, fns, _r = tools
    with bind(_Host(), skill):
        assert (
            fns["build_deck"](
                outline={"slides": []}, output_path=str(tmp_path / "x.pptx")
            )["status"]
            == "error"
        )
        assert (
            fns["build_deck"](
                outline={"slides": [{"type": "title", "title": "x"}]},
                output_path=str(tmp_path / "x.docx"),
            )["status"]
            == "error"
        )
        assert (
            fns["build_deck"](
                outline={"slides": [{"type": "title", "title": "x"}]},
                output_path=str(tmp_path / "no-such-dir" / "x.pptx"),
            )["status"]
            == "error"
        )
        assert (
            fns["build_deck"](
                outline={"slides": [{"type": "title", "title": "x"}]},
                output_path=str(SKILL_DIR / "x.pptx"),
            )["status"]
            == "error"
        )
        assert (
            fns["verify_deck"](path=str(tmp_path / "missing.pptx"))["status"] == "error"
        )
        assert (
            fns["outline_from_blueprint"](name="no_such_blueprint")["status"] == "error"
        )
        assert fns["deck_outline_from_markdown"](markdown="")["status"] == "error"
        bad = fns["update_deck_prefs"](changes={"density": "huge"})
        assert bad["status"] == "error"
    assert not (SKILL_DIR / "x.pptx").exists()


def test_the_procedure_asks_at_most_once_and_never_blocks():
    """The rule for asking: one question, with a default and a short timeout."""
    import re

    body = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
    assert "request_user_input" in body
    assert re.search(r"Ask \*\*one\*\* short question \*\*only\*\*", body)
    assert "Never ask about anything a default covers" in body
    assert "references/brief-and-sources.md" in body
    ref = (SKILL_DIR / "references" / "brief-and-sources.md").read_text(
        encoding="utf-8"
    )
    timeouts = [int(t) for t in re.findall(r"timeout_seconds=(\d+)", ref)]
    assert timeouts and all(t <= 60 for t in timeouts)
    assert "default_if_no_response=" in ref
    # The example of a request that must NOT trigger a question is there too.
    assert "no question" in ref


def test_markdown_tool_and_required_items_through_the_tool_layer(tools, tmp_path):
    skill, fns, _r = tools
    md = (
        "# Harbour ferry update\n\n"
        "## Crossings rose to 1,180 a week after the timetable change\n"
        "- 1,180 weekly crossings, up from 960\n"
        "Notes: the pier repair finishes in May.\n"
    )
    with bind(_Host(), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md,
            output_path=str(tmp_path / "ferry.pptx"),
            required_items=["1,180 weekly crossings", "pier repair finishes in May"],
        )
        assert r["status"] == "built_with_issues"
        assert r["verify"]["summary"]["required_in_notes_only"] == 1
        report = fns["verify_deck"](
            path=str(tmp_path / "ferry.pptx"),
            render=False,
            required_items=["1,180 weekly crossings"],
        )
        assert report["status"] == "pass", report["issues"]
        bad = fns["build_deck_from_markdown"](
            markdown="", output_path=str(tmp_path / "x.pptx")
        )
        assert bad["status"] == "error"


def test_a_one_slide_markdown_build_is_one_slide(tools, tmp_path):
    skill, fns, _r = tools
    md = "# Ferry plan\n\n## Ferry service grows to a network\n- Four piers\n"
    with bind(_Host(), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "one.pptx"), expected_slides=1
        )
    assert r["status"] == "ok", r
    assert r["slides"] == 1


def test_the_procedure_checks_a_fresh_build_once(tools, tmp_path):
    """The build already runs the check; a second verify_deck is a wasted model turn."""
    skill, fns, _r = tools
    body = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
    assert "The build result's `verify` is the check" in body
    assert "only after\n   `add_slide` or `update_slide`" in body
    with bind(_Host(), skill):
        r = fns["build_deck_from_markdown"](
            markdown="## Ferry service grows to a network\n- Four piers\n",
            output_path=str(tmp_path / "f.pptx"),
            expected_slides=1,
        )
    assert r["verify"]["status"] == "pass"


def test_a_deck_with_no_folder_saves_in_gaia_documents(tools, tmp_path):
    skill, fns, _r = tools
    md = (
        "# Ferry plan\n\n## Crossings rose after the timetable change\n- 1,180 a week\n"
    )
    documents = tmp_path / "home" / "documents"
    with bind(_Host(), skill):
        unnamed = fns["build_deck_from_markdown"](markdown=md)
        named = fns["build_deck_from_markdown"](markdown=md, output_path="ferry.pptx")
    assert unnamed["path"] == str(documents / "Ferry plan.pptx")
    assert named["path"] == str(documents / "ferry.pptx")
    assert (documents / "Ferry plan.pptx").is_file()
    assert (documents / "ferry.pptx").is_file()


def test_an_unchanged_rebuild_of_a_deck_that_failed_its_check_is_refused(
    tools, tmp_path
):
    skill, fns, _r = tools
    md = "## Ferry crossings rose\n- 1,180 a week\n"
    out = str(tmp_path / "ferry.pptx")
    with bind(_Host(), skill):
        first = fns["build_deck_from_markdown"](
            markdown=md, output_path=out, required_items=["Pier repair finishes in May"]
        )
        again = fns["build_deck_from_markdown"](
            markdown=md, output_path=out, required_items=["Pier repair finishes in May"]
        )
        fixed = fns["build_deck_from_markdown"](
            markdown=md + "- Pier repair finishes in May\n",
            output_path=out,
            required_items=["Pier repair finishes in May"],
        )
    assert (
        first["status"] == "built_with_issues"
        and "Never call it verified" in first["next"]
    )
    assert again["status"] == "error"
    assert any("Pier repair" in p for p in again["problems"])
    assert fixed["status"] == "ok"


def test_a_build_says_when_to_stop(tools, tmp_path):
    skill, fns, _r = tools
    md = "# Ferry plan\n\n## Ferry crossings rose\n- 1,180 a week\n"
    must = ["Pier repair finishes in May"]
    with bind(_Host(), skill):
        passed = fns["build_deck_from_markdown"](markdown=md)
        first = fns["build_deck_from_markdown"](
            markdown=md + "- a\n", required_items=must
        )
        second = fns["build_deck_from_markdown"](
            markdown=md + "- b\n", required_items=must
        )
    assert "Report it to the user now" in passed["next"]
    assert first["next"].startswith("The deck failed its check")
    assert second["next"].startswith("Stop building")
    assert passed["path"] == first["path"] == second["path"]


def test_a_null_output_path_saves_in_gaia_documents(tools, tmp_path):
    skill, fns, _r = tools
    with bind(_Host(), skill):
        r = fns["build_deck_from_markdown"](
            markdown="# Plan\n\n## Crossings rose\n- 1,180 a week\n", output_path=None
        )
    assert r["status"] == "ok", r
    assert r["path"].startswith(str(tmp_path / "home" / "documents"))


_STAGES = (
    "## Ferry service grows from one pier to a network\n"
    "Diagram: process\n"
    "### One pier\n- Hourly crossing\n"
    "### Two piers\n- Half-hourly crossings\n"
    "### Network\n- Four piers, every ten minutes\n"
)
_SPLIT = (
    "# Ferry plan\n\n" + _STAGES + "\n## What grows at each stage\n"
    "| Dimension | One pier | Two piers | Network |\n"
    "|---|---|---|---|\n"
    "| Boats | 1 | 2 | 6 |\n"
    "Source: ferry_fleet.xlsx, 2026\n"
)


class _Asking(_Host):
    """A stand-in agent answering *request*, as Agent keeps it in _current_query."""

    def __init__(self, request: str = "") -> None:
        self._current_query = request


_ONE_SLIDE = "Make a crisp one-slide framework of how the ferry service grows."


def test_a_count_error_is_not_cleared_by_changing_the_count(tools, tmp_path):
    """A one-slide request built as three slides must be merged, not re-counted."""
    skill, fns, _r = tools
    out = str(tmp_path / "ferry.pptx")
    with bind(_Asking(_ONE_SLIDE), skill):
        first = fns["build_deck_from_markdown"](
            markdown=_SPLIT, output_path=out, expected_slides=1
        )
        raised = fns["build_deck_from_markdown"](
            markdown=_SPLIT, output_path=out, expected_slides=2
        )
        dropped = fns["build_deck_from_markdown"](
            markdown=_SPLIT, output_path=out, expected_slides=0
        )
        merged = fns["build_deck_from_markdown"](
            markdown=_STAGES + "Axis: More piers, boats and crossings\n",
            output_path=out,
            expected_slides=1,
        )
    assert first["status"] == "error"
    assert "'Axis:' line" in first["problems"][0]
    assert raised["status"] == "error"
    assert '("one-slide")' in raised["problems"][0]
    assert dropped["status"] == "error", "0 must apply the user's count, not skip it"
    assert merged["status"] == "ok", merged
    assert merged["slides"] == 1
    assert merged["slides_built"] == [
        "1. Ferry service grows from one pier to a network [diagram: process]"
    ]


def test_the_models_first_guess_never_overrides_the_users_count(tools, tmp_path):
    """The model misread a one-slide request as five; the user's words win."""
    skill, fns, _r = tools
    with bind(_Asking(_ONE_SLIDE), skill):
        guess = fns["build_deck_from_markdown"](
            markdown=_STAGES, output_path=str(tmp_path / "g.pptx"), expected_slides=5
        )
        right = fns["build_deck_from_markdown"](
            markdown=_STAGES, output_path=str(tmp_path / "r.pptx"), expected_slides=1
        )
    assert guess["status"] == "error"
    assert "not 5" in guess["problems"][0]
    assert right["status"] == "ok", right


def test_without_a_count_in_the_request_the_model_chooses(tools, tmp_path):
    skill, fns, _r = tools
    with bind(_Asking("Make a deck about the ferry service: 1, 2 and 6 boats."), skill):
        two = fns["build_deck_from_markdown"](
            markdown=_SPLIT, output_path=str(tmp_path / "b.pptx"), expected_slides=2
        )
    assert two["status"] == "ok", two
    assert [line.split(" [")[0] for line in two["slides_built"]] == [
        "1. Ferry service grows from one pier to a network",
        "2. What grows at each stage",
    ]


def test_verify_applies_the_users_count(tools, tmp_path):
    skill, fns, _r = tools
    out = str(tmp_path / "ferry.pptx")
    with bind(_Asking(_ONE_SLIDE), skill):
        built = fns["build_deck_from_markdown"](markdown=_STAGES, output_path=out)
        report = fns["verify_deck"](path=out, expected_slides=3, render=False)
    assert built["status"] == "ok", built
    assert report["status"] == "error"
    assert "not 3" in report["problems"][0]


@pytest.mark.parametrize(
    "request_text, count",
    [
        ("a crisp one-slide framework ... a single executive slide", 1),
        ("Make a 5-slide deck on Q3", 5),
        ("Build a deck with three PowerPoint slides", 3),
        ("Summarize one of the slides", 0),
        ("make a slide deck about our roadmap", 0),
        ("add two more slides", 0),
        ("turn slide 3 into a chart", 0),
        ("a 2 slide summary and 4 slides of appendix", 0),
    ],
)
def test_the_count_comes_from_the_users_words(tools, request_text, count):
    import tools as skill_tools  # the skill's tools.py, on sys.path via conftest

    assert skill_tools._asked_count(request_text)[0] == count


def test_skill_call_request_is_the_users_message(tools):
    skill, _fns, _r = tools
    with bind(_Asking("one slide please"), skill) as call:
        assert call.request == "one slide please"
    with bind(_Host(), skill) as call:
        assert call.request == ""


def test_an_axis_line_draws_the_boxes_as_stages(tools, tmp_path):
    """Axis: names what grows along stages, so boxes with one become a process."""
    skill, fns, _r = tools
    boxes = _STAGES.replace("Diagram: process\n", "") + "Axis: More piers and boats\n"
    with bind(_Asking(), skill):
        r = fns["build_deck_from_markdown"](
            markdown=boxes, output_path=str(tmp_path / "f.pptx"), expected_slides=1
        )
    assert r["status"] == "ok", r
    assert r["slides_built"][0].endswith("[diagram: process]")


def test_an_axis_on_a_diagram_that_cannot_draw_it_is_an_error(tools, tmp_path):
    """A cycle ignores Axis:, so the dimensions it lists would vanish from the slide."""
    skill, fns, _r = tools
    cycle = _STAGES.replace("process", "cycle") + "Axis: More piers and boats\n"
    with bind(_Asking(), skill):
        r = fns["build_deck_from_markdown"](
            markdown=cycle, output_path=str(tmp_path / "c.pptx"), expected_slides=1
        )
    assert r["status"] == "error"
    assert "Diagram: process" in r["problems"][0]


def test_an_axis_under_more_than_six_steps_is_an_error(tools, tmp_path):
    skill, fns, _r = tools
    steps = "".join(f"### Step {i}\n- Detail {i}\n" for i in range(1, 8))
    md = f"## Seven steps to launch\nDiagram: process\n{steps}Axis: More effort\n"
    with bind(_Asking(), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "s.pptx"), expected_slides=1
        )
    assert r["status"] == "error"
    assert "6 steps at most" in r["problems"][0]


_GROWS = (
    "One slide of the ferry stages. From left to right, increase: boats, crossings."
)


def test_the_users_growth_list_becomes_the_axis_when_the_model_leaves_it_out(
    tools, tmp_path
):
    """The user wrote what grows; the skill draws it instead of sending a retry."""
    from pptx import Presentation

    skill, fns, _r = tools
    out = str(tmp_path / "ferry.pptx")
    boxes = _STAGES.replace("Diagram: process\n", "")
    with bind(_Asking(_GROWS), skill):
        r = fns["build_deck_from_markdown"](
            markdown=boxes, output_path=out, expected_slides=1
        )
    assert r["status"] == "ok", r
    assert r["slides_built"][0].endswith("[diagram: process]")
    assert any("axis" in n for n in r["notes"])
    arrow = next(
        s for s in Presentation(out).slides[0].shapes if s.name == "Axis arrow"
    )
    assert arrow.text_frame.text == "More boats and crossings"


def test_a_model_axis_is_kept_when_it_names_the_list(tools, tmp_path):
    skill, fns, _r = tools
    with bind(_Asking(_GROWS), skill):
        r = fns["build_deck_from_markdown"](
            markdown=_STAGES + "Axis: More boats and crossings\n",
            output_path=str(tmp_path / "f.pptx"),
            expected_slides=1,
        )
    assert r["status"] == "ok", r
    assert not any("axis" in n for n in r.get("notes") or [])


def test_no_axis_is_needed_without_a_left_to_right_request(tools, tmp_path):
    skill, fns, _r = tools
    with bind(_Asking(_ONE_SLIDE), skill):
        r = fns["build_deck_from_markdown"](
            markdown=_STAGES, output_path=str(tmp_path / "f.pptx"), expected_slides=1
        )
    assert r["status"] == "ok", r


def test_an_axis_must_name_everything_the_user_listed(tools, tmp_path):
    """A copied example axis that drops the user's items is refused, naming them."""
    skill, fns, _r = tools
    request = (
        "One slide of the ferry stages. As we move from left to right, increase:\n"
        "\n* Boats in service\n* Crossings per hour\n* Crew on shift\n\nKeep it simple."
    )
    out = str(tmp_path / "ferry.pptx")
    with bind(_Asking(request), skill):
        partial = fns["build_deck_from_markdown"](
            markdown=_STAGES + "Axis: More boats\n", output_path=out
        )
        full = fns["build_deck_from_markdown"](
            markdown=_STAGES + "Axis: More boats, crossings and crew\n", output_path=out
        )
    assert partial["status"] == "error"
    assert "Crossings per hour; Crew on shift" in partial["problems"][0]
    assert full["status"] == "ok", full


def test_boxes_on_a_chart_slide_are_an_error_not_dropped(tools, tmp_path):
    """A chart slide has no room for '### ' boxes; they must not vanish silently."""
    skill, fns, _r = tools
    md = (
        "## Ferry demand rose at every pier\n"
        "```chart bar\nsource: ferry_counts.xlsx, 2026\nPier,Crossings\nNorth,120\nSouth,90\n```\n"
        "### North pier\n- Hourly crossing\n### South pier\n- Half-hourly\n"
    )
    with bind(_Asking(), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "c.pptx")
        )
    assert r["status"] == "error"
    assert "(North pier, South pier) would not appear" in r["problems"][0]


def test_a_chart_alone_cannot_drop_what_the_user_listed(tools, tmp_path):
    skill, fns, _r = tools
    request = (
        "One slide of the ferry stages. As we move from left to right, increase:\n"
        "\n* Boats in service\n* Crew on shift\n"
    )
    md = (
        "## Ferry crossings grow at every stage\n"
        "```chart bar\nsource: ferry_counts.xlsx, 2026\nStage,Crossings\n"
        "One pier,10\nNetwork,60\n```\n"
    )
    with bind(_Asking(request), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "c.pptx")
        )
    assert r["status"] == "error"
    assert "Boats in service; Crew on shift" in r["problems"][0]


@pytest.mark.parametrize(
    "request_text, listed",
    [
        (
            "As we move from left to right, increase:\n\n* Task complexity\n* Compute\n",
            ["Task complexity", "Compute"],
        ),
        (
            "Across the tiers these grow: cost, staff and hours.",
            ["cost", "staff", "hours"],
        ),
        ("From step to step, more of: risk, review effort", ["risk", "review effort"]),
        ("Lay the three options out left to right.", []),
        ("Revenue should increase: we need a deck on it.", []),
        ("Show what grows across the stages: cost", []),
        ("One slide: the stages; from left to right, increase boats and crew.", []),
    ],
)
def test_what_grows_comes_only_from_an_explicit_list(tools, request_text, listed):
    import tools as skill_tools

    assert skill_tools._grows_list(request_text) == listed


def test_content_under_the_only_title_is_the_slide_not_a_cover(tools, tmp_path):
    """'# Title' over KPIs with no '## ' slide used to become a cover with no KPIs."""
    skill, fns, _r = tools
    md = "# October KPIs\n\n- Revenue: $12.4M\n- Churn: 2.1%\n- NPS: 48\n"
    with bind(
        _Asking("One slide with October KPIs: $12.4M revenue, 2.1% churn, NPS 48."),
        skill,
    ):
        r = fns["build_deck_from_markdown"](
            markdown=md,
            output_path=str(tmp_path / "k.pptx"),
            required_items=["$12.4M", "2.1%", "48"],
        )
    assert r["status"] == "ok", r
    assert r["slides"] == 1


def test_content_under_a_real_cover_is_an_error_not_dropped(tools, tmp_path):
    skill, fns, _r = tools
    md = "# Ferry plan\n- Crossings up 20%\n\n## Next steps\n- Order a boat\n"
    with bind(_Asking(), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "c.pptx")
        )
    assert r["status"] == "error"
    assert "Crossings up 20%" in r["problems"][0]


_CHURN = (
    "## Churn fell every quarter\n"
    "```chart line\nsource: {src}\nQuarter,Churn %\nQ1,3.4\nQ2,2.9\nQ3,2.1\n```\n"
)


def test_made_up_numbers_must_be_labelled_on_the_slide(tools, tmp_path):
    """Invented data the user invited is fine; presenting it as real is not."""
    skill, fns, _r = tools
    out = str(tmp_path / "churn.pptx")
    with bind(_Asking("One slide on churn going down. Make up some numbers."), skill):
        bare = fns["build_deck_from_markdown"](
            markdown=_CHURN.format(src="CRM export, 2026"), output_path=out
        )
        labelled = fns["build_deck_from_markdown"](
            markdown=_CHURN.format(src="Illustrative data, not actual results"),
            output_path=out,
        )
    assert bare["status"] == "error"
    assert "Illustrative data" in bare["problems"][0]
    assert labelled["status"] == "ok", labelled


@pytest.mark.parametrize(
    "request_text, invites",
    [
        ("Make up some realistic numbers", True),
        ("use dummy data for now", True),
        ("a hypothetical scenario with sample numbers", True),
        ("Make a slide about our Q3 revenue of $5.6M", False),
        ("make sure it looks clean", False),
    ],
)
def test_only_an_invitation_to_invent_numbers_triggers_the_label(
    tools, request_text, invites
):
    import tools as skill_tools

    assert bool(skill_tools._INVENTED.search(request_text)) is invites


def test_a_chart_type_on_its_own_line_is_read(tools, tmp_path):
    """'```chart' then 'bar' on the next line used to fail as a chart with no data."""
    skill, fns, _r = tools
    md = (
        "## West led every quarter\n"
        "```chart\nbar\nsource: finance_q3.xlsx, 2026\nRegion,Revenue\nWest,5.6\nEast,4.1\n```\n"
    )
    with bind(_Asking(), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "c.pptx")
        )
    assert r["status"] == "ok", r


class _Seeing(_Asking):
    """A stand-in agent that has also seen tool results, as Agent.seen_sources keeps."""

    def __init__(self, request: str, seen=()) -> None:
        super().__init__(request)
        self._seen = [("user", {}, request), *seen]

    def seen_sources(self):
        return list(self._seen)


_REGIONS = "Q3 revenue in $M: West 5.6, East 4.1, North 3.2. Two slides."


def _share_slide(claim: str) -> str:
    return (
        f"## West drove {claim} of Q3 revenue\n"
        "```chart bar\nsource: finance_q3.xlsx\nRegion,Revenue ($M)\n"
        "West,5.6\nEast,4.1\nNorth,3.2\n```\n"
        "## Next steps\n- Hire sales reps in North\n"
    )


def test_a_number_nobody_gave_is_refused(tools, tmp_path):
    """45% is mental arithmetic (the data gives 43%): nothing grounds it."""
    skill, fns, _r = tools
    with bind(_Seeing(_REGIONS), skill):
        r = fns["build_deck_from_markdown"](
            markdown=_share_slide("45%"), output_path=str(tmp_path / "s.pptx")
        )
    assert r["status"] == "error"
    assert "45" in r["problems"][0] and "run_python" in r["problems"][0]


def test_a_number_a_tool_computed_is_grounded_even_rounded(tools, tmp_path):
    skill, fns, _r = tools
    computed = ("run_python", {"code": "5.6/12.9"}, '{"stdout": "0.4341085"}')
    with bind(_Seeing(_REGIONS, [computed]), skill):
        r = fns["build_deck_from_markdown"](
            markdown=_share_slide("43%"), output_path=str(tmp_path / "s.pptx")
        )
    assert r["status"] == "ok", r


def test_a_number_from_a_file_read_is_grounded(tools, tmp_path):
    skill, fns, _r = tools
    read = ("read_file", {"path": "kpis.txt"}, "NPS this quarter: 48")
    md = "## NPS reached 48 this quarter\n- Up from last quarter\n"
    with bind(_Seeing("One slide on our NPS from kpis.txt.", [read]), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "n.pptx")
        )
    assert r["status"] == "ok", r


def test_the_skills_own_outputs_and_instructions_do_not_ground_numbers(tools, tmp_path):
    """The SKILL.md example says 'Revenue $12.4M'; that must not excuse inventing it."""
    skill, fns, _r = tools
    own = [
        ("load_skill", {"name": "slide-decks"}, "Revenue $12.4M; West +37%"),
        ("slide-decks/build_deck_from_markdown", {}, '{"slides": 12.4}'),
    ]
    md = "## Revenue hit $12.4M\n- A strong quarter\n"
    with bind(_Seeing("One slide on how the quarter went.", own), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "r.pptx")
        )
    assert r["status"] == "error"
    assert "12.4" in r["problems"][0]


def test_a_count_of_the_slides_own_parts_in_its_title_is_fine(tools, tmp_path):
    skill, fns, _r = tools
    md = _STAGES.replace(
        "## Ferry service grows from one pier to a network",
        "## 3 stages from one pier to a network",
    )
    with bind(_Seeing("One slide: the ferry stages."), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "c.pptx")
        )
    assert r["status"] == "ok", r


def test_invented_dates_are_refused_like_any_number(tools, tmp_path):
    skill, fns, _r = tools
    md = "## The move wraps up in March\n- Jan 15: lease signed\n- Mar 10: go live\n"
    with bind(_Seeing("One slide on our office move to Austin in March."), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "d.pptx")
        )
    assert r["status"] == "error"
    assert "15" in r["problems"][0] and "10" in r["problems"][0]


def test_emoji_the_user_did_not_use_are_refused(tools, tmp_path):
    """They render as empty boxes in LibreOffice and in the exported image."""
    skill, fns, _r = tools
    md = "## Three priorities for 2027\n- 🇲🇽 Grow in Mexico\n- 📱 Launch the app\n"
    with bind(_Seeing("One slide: grow in Mexico, launch the app."), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "e.pptx")
        )
    assert r["status"] == "error"
    assert "📱" in r["problems"][0]


def test_bold_in_a_table_cell_is_bold_not_asterisks(out):
    from pptx import Presentation
    from slidedeck.build import build_deck
    from slidedeck.markdown import outline_from_markdown

    md = "## Model B is the pick\n| Model | RAM |\n|---|---|\n| A | 16 GB |\n| B | **32 GB** |\n"
    path = out("t.pptx")
    assert build_deck(outline_from_markdown(md), path)["status"] in (
        "ok",
        "built_with_issues",
    )
    table = next(s for s in Presentation(path).slides[0].shapes if s.has_table).table
    cell = table.cell(2, 1).text_frame.paragraphs[0]
    assert cell.text == "32 GB" and all(r.font.bold for r in cell.runs)


def test_export_slide_image_saves_a_png_by_default_and_jpeg_on_request(
    tools, tmp_path, soffice, monkeypatch
):
    import tempfile

    # LibreOffice's profile path must stay short; pytest's tmp_path is not.
    short = tempfile.mkdtemp(prefix="g")
    monkeypatch.setenv("GAIA_HOME", short)
    skill, fns, _r = tools
    deck = str(tmp_path / "s.pptx")
    with bind(_Asking(), skill):
        built = fns["build_deck_from_markdown"](markdown=_STAGES, output_path=deck)
        png = fns["export_slide_image"](deck_path=deck)
        jpg = fns["export_slide_image"](
            deck_path=deck, image_format="jpeg", output_path=str(tmp_path / "s.jpg")
        )
        bad = fns["export_slide_image"](deck_path=deck, image_format="gif")
        missing = fns["export_slide_image"](deck_path=deck, slide=4)
    assert built["status"] == "ok", built
    assert png["status"] == "ok", png
    p = png["images"][0]["path"]
    assert p.endswith(".png") and open(p, "rb").read(8) == b"\x89PNG\r\n\x1a\n"
    assert png["images"][0]["width"] == 1920
    assert (
        jpg["status"] == "ok"
        and open(jpg["images"][0]["path"], "rb").read(3) == b"\xff\xd8\xff"
    )
    assert bad["status"] == "error" and "png" in bad["problems"][0]
    assert missing["status"] == "error" and "has 1" in missing["problems"][0]


_FUNNEL = (
    "One slide on our October funnel: 20,000 visited signup, 6,400 created an "
    "account, 1,600 were active after 7 days."
)


def test_arithmetic_on_the_users_own_numbers_needs_no_tool(tools, tmp_path):
    """13,600 lost (68%) is the user's numbers subtracted and divided, not invented."""
    skill, fns, _r = tools
    md = (
        "## Signup loses 13,600 visitors (68%)\n"
        "- 20,000 visited, 6,400 created an account\n- 8% were active after 7 days\n"
    )
    with bind(_Seeing(_FUNNEL), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "f.pptx")
        )
    assert r["status"] == "ok", r


def test_a_share_the_users_numbers_do_not_give_is_still_refused(tools, tmp_path):
    skill, fns, _r = tools
    md = "## Signup loses 45% of visitors\n- 20,000 visited, 6,400 signed up\n"
    with bind(_Seeing(_FUNNEL), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "f.pptx")
        )
    assert r["status"] == "error" and "45" in r["problems"][0]


def test_requested_numbers_hidden_in_the_notes_are_refused(tools, tmp_path):
    """Two of three themes moved to the notes used to pass the check."""
    skill, fns, _r = tools
    request = (
        "Two slides on our interviews: setup is slow (9 of 14), calendar wanted "
        "(7 of 14), privacy worries (6 of 14)."
    )
    md = (
        "## Setup is the top pain\n- 9 of 14 said setup is slow\n"
        "Notes: also calendar (7 of 14) and privacy (6 of 14)\n"
        "## Next\n- Cut setup time\n"
    )
    with bind(_Seeing(request), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "n.pptx")
        )
    assert r["status"] == "error"
    assert "only in the speaker notes" in r["problems"][0]


def test_model_shapes_that_used_to_cost_rebuilds_now_build_first_time(tools, tmp_path):
    """Real first attempts from stress runs that the skill used to refuse or garble."""
    skill, fns, _r = tools
    outline = {
        "slides": [
            {
                "type": "diagram",
                "diagram": "boxes",
                "title": "Ferry service grows in three stages",
                "axis": "More boats",
                "items": [
                    "One pier\n\nHourly crossing",
                    "Two piers\n\nHalf-hourly",
                    "Network\n\nEvery ten minutes",
                ],
            }
        ]
    }
    md = (
        "# Ferry plan\n\n## Three stages\n### One pier\nHourly crossing\nOne boat\n"
        "### Two piers\nHalf-hourly\nTwo boats\n### Network\nEvery ten minutes\nSix boats\n"
        "\n---\n\n# Recommendation\n- Order two boats\n"
    )
    with bind(_Asking(), skill):
        j = fns["build_deck"](outline=outline, output_path=str(tmp_path / "j.pptx"))
        m = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "m.pptx")
        )
    assert j["status"] == "ok", j
    assert j["slides_built"] == [
        "1. Ferry service grows in three stages [diagram: process]"
    ]
    assert m["status"] == "ok", m
    assert [x.split(" [")[0] for x in m["slides_built"]] == [
        "1. Ferry plan",
        "2. Three stages",
        "3. Recommendation",
    ]


def test_a_chart_of_the_users_own_figures_needs_no_made_up_source(tools, tmp_path):
    """The user typed the numbers, so there is no file to cite and nothing to invent."""
    skill, fns, _r = tools
    md = "## Churn fell every quarter\n```chart line\nQuarter,Churn %\nQ1,3.4\nQ2,2.9\nQ3,2.1\n```\n"
    with bind(
        _Asking("One slide: churn was 3.4% in Q1, 2.9% in Q2, 2.1% in Q3."), skill
    ):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "u.pptx")
        )
    assert r["status"] == "ok", r
    assert not [i for i in r["verify"]["issues"] if i["check"] == "missing_source"]


def test_a_chart_of_other_figures_still_needs_a_source(tools, tmp_path):
    skill, fns, _r = tools
    md = "## Churn fell every quarter\n```chart line\nQuarter,Churn %\nQ1,3.4\nQ2,2.9\nQ3,2.1\n```\n"
    with bind(_Host(), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "f.pptx")
        )
    assert [i for i in r["verify"]["issues"] if i["check"] == "missing_source"]


def test_phase_and_step_labels_are_not_figures(tools, tmp_path):
    skill, fns, _r = tools
    md = (
        "## Rollout plan\n\n### Phase 1\n- Pilot one depot\n\n"
        "### Phase 2\n- Two regions\n\n### Phase 3\n- Every depot\n"
    )
    with bind(_Asking("One slide on the depot rollout plan in phases."), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "p.pptx")
        )
    assert r["status"] == "ok", r


def test_one_box_beside_a_table_becomes_its_takeaway(tools, tmp_path):
    skill, fns, _r = tools
    md = (
        "## West leads on revenue\n\n| Region | Revenue |\n|---|---|\n"
        "| West | 4.2 |\n| East | 3.1 |\n\n### Key insight\nKeep funding West\n"
    )
    with bind(_Asking("One slide: West 4.2, East 3.1 revenue."), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "t.pptx")
        )
    assert r["status"] == "ok", r
    assert "Keep funding West" in json.dumps(r["slides_built"])


def test_a_new_chat_never_overwrites_another_chats_deck(tools, tmp_path):
    """Same title in a new chat: a fresh file, not an overwrite prompt."""
    skill, fns, _r = tools
    md = "## Ferry plan\n- Order a boat\n- Hire crew\n"
    with bind(_Asking("one slide on the ferry plan"), skill):
        first = fns["build_deck_from_markdown"](markdown=md)
        again = fns["build_deck_from_markdown"](markdown=md)
    with bind(_Asking("one slide on the ferry plan"), skill):
        other = fns["build_deck_from_markdown"](markdown=md)
    assert first["status"] == again["status"] == other["status"] == "ok", other
    assert first["path"] == again["path"]
    assert other["path"] != first["path"]
    assert other["path"].endswith("(2).pptx")
    assert any("already exists" in n for n in other.get("notes") or [])


def test_status_words_colour_boxes_and_cards_fit_their_text(tools, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    skill, fns, _r = tools
    md = (
        "## Weekly status\n\n### Backend — Green\n- Done\n\n"
        "### Connector — Amber\n- Fix in review\n\n### Mobile UI — Red\n- Slipping\n"
    )
    out = tmp_path / "s.pptx"
    with bind(_Asking("one slide weekly status"), skill):
        r = fns["build_deck_from_markdown"](markdown=md, output_path=str(out))
    assert r["status"] == "ok", r
    boxes = [
        s for s in Presentation(str(out)).slides[0].shapes if s.name.startswith("Box ")
    ]
    fills = {str(b.fill.fore_color.rgb) for b in boxes}
    assert len(fills) == 3
    assert all(b.height < Inches(3) for b in boxes)


def test_several_points_beside_a_table_become_a_side_panel(tools, tmp_path):
    skill, fns, _r = tools
    md = (
        "## Build vs license\n\n| | Build | License |\n|---|---|---|\n"
        "| Time | 9 months | 6 weeks |\n| Cost | 4 engineers | $0.004/min |\n\n"
        "### Recommendation: license now\n- Faster to market\n- Cost is manageable\n"
        "- Revisit in 12 months\n"
    )
    with bind(
        _Asking(
            "One slide: build takes 9 months with 4 engineers; license takes 6 weeks at $0.004/min; revisit in 12 months."
        ),
        skill,
    ):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "d.pptx")
        )
    assert r["status"] == "ok", r


@pytest.mark.parametrize(
    "request_text, count",
    [
        ("Generate an image of a slide that shows our values", 1),
        ("I need a single slide on churn", 1),
        ("Make a slide deck on Q3", 0),
        ("Add a slide about risks", 0),
        ("Make a 5-slide deck, end with a slide on the ask", 5),
    ],
)
def test_a_slide_means_one_slide_but_a_slide_deck_does_not(tools, request_text, count):
    import tools as skill_tools

    assert skill_tools._asked_count(request_text)[0] == count


def test_italic_markers_do_not_show_on_the_slide(tools, tmp_path):
    skill, fns, _r = tools
    md = '## What we heard\n\n- *"I gave up on day two"*\n- **Setup** is _slow_\n'
    with bind(_Asking("one slide on what we heard"), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "i.pptx")
        )
    assert r["status"] == "ok", r
    from pptx import Presentation

    text = " ".join(
        sh.text_frame.text
        for sh in Presentation(r["path"]).slides[0].shapes
        if sh.has_text_frame
    )
    assert '"I gave up on day two"' in text and "*" not in text and "_slow_" not in text


def test_the_option_a_decision_picks_stands_out_in_its_column(tools, tmp_path):
    from pptx import Presentation

    skill, fns, _r = tools
    md = (
        "## License the speech model now\n\n| | Build in-house | License vendor |\n"
        "|---|---|---|\n| Time | 9 months | 6 weeks |\n| Control | Full | Vendor lock-in |\n"
        "\nTakeaway: License now, revisit building in 12 months\n"
    )
    out = tmp_path / "d.pptx"
    with bind(
        _Asking(
            "One slide: build takes 9 months, license takes 6 weeks; revisit in 12 months."
        ),
        skill,
    ):
        r = fns["build_deck_from_markdown"](markdown=md, output_path=str(out))
    assert r["status"] == "ok", r
    table = next(
        sh.table for sh in Presentation(str(out)).slides[0].shapes if sh.has_table
    )
    picked = {str(table.cell(i, 2).fill.fore_color.rgb) for i in (1, 2)}
    other = {str(table.cell(i, 1).fill.fore_color.rgb) for i in (1, 2)}
    assert len(picked) == 1 and not picked & other
    assert all(r.font.bold for r in table.cell(1, 2).text_frame.paragraphs[0].runs)


def test_a_count_the_user_gave_out_of_a_base_keeps_its_base(tools, tmp_path):
    """'7 of 14' shown as a bare 7 loses how many were asked."""
    skill, fns, _r = tools
    ask = "One slide: 9 of 14 said setup is slow and 7 of 14 want calendar sync."
    bare = "## Setup is the top pain\n- 9 of 14 say setup is slow\n- 7 want calendar sync\n"
    kept = bare.replace("- 7 want", "- 7 of 14 want")
    with bind(_Asking(ask), skill):
        r1 = fns["build_deck_from_markdown"](
            markdown=bare, output_path=str(tmp_path / "a.pptx")
        )
        r2 = fns["build_deck_from_markdown"](
            markdown=kept, output_path=str(tmp_path / "b.pptx")
        )
    assert r1["status"] == "error" and "7 of 14" in r1["problems"][0]
    assert r2["status"] == "ok", r2


def test_table_numbers_keep_the_separators_the_user_typed(tools, tmp_path):
    from pptx import Presentation

    skill, fns, _r = tools
    md = (
        "## Signup is the biggest leak\n\n| Stage | Users | Year |\n|---|---|---|\n"
        "| Visited | 20,000 | 2026 |\n| Signed up | 6,400 | 2026 |\n\n"
        "### Biggest drop: signup\nTwo in three visitors never sign up\n"
    )
    with bind(_Asking("One slide: 20,000 visited, 6,400 signed up, in 2026."), skill):
        r = fns["build_deck_from_markdown"](
            markdown=md, output_path=str(tmp_path / "f.pptx")
        )
    assert r["status"] == "ok", r
    shapes = Presentation(r["path"]).slides[0].shapes
    cells = [
        c.text
        for sh in shapes
        if sh.has_table
        for row in sh.table.rows
        for c in row.cells
    ]
    assert "6,400" in cells and "20,000" in cells and "2026" in cells
    text = json.dumps(r["slides_built"])
    assert "Biggest drop: signup. Two in three" in text


def test_a_rebuild_replaces_its_own_deck_even_after_the_skill_reloads(tools, tmp_path):
    """Ownership is stamped in the file, so reloading the skill loses nothing."""
    from gaia.skills.loader import register_skill_tools

    skill, fns, _r = tools
    md = "## Ferry plan\n- Order a boat\n- Hire crew\n"
    with bind(_Asking("one slide on the ferry plan"), skill):
        theirs = fns["build_deck_from_markdown"](markdown=md)
    me = _Asking("one slide on the ferry plan")
    with bind(me, skill):
        first = fns["build_deck_from_markdown"](markdown=md)
    reloaded = {
        k.split("/", 1)[1]: v["function"]
        for k, v in register_skill_tools(skill).items()
    }
    with bind(me, skill):
        again = reloaded["build_deck_from_markdown"](markdown=md)
    assert first["path"] != theirs["path"]
    assert again["path"] == first["path"], (first["path"], again["path"])
