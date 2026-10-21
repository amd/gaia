# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""verify_deck catches what is wrong with a deck — including decks this skill did not build.

Each negative test hand-builds a flawed deck with python-pptx, as a person or
another tool would, and asserts the specific check fires; the good decks from
build_deck must stay clean.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
from slidedeck.build import build_deck
from slidedeck.outline import normalize_outline
from slidedeck.verify import verify_deck

GOOD = [
    {
        "type": "title",
        "title": "Maple Street clinic network review",
        "subtitle": "Quarterly update",
    },
    {
        "type": "bullets",
        "title": "Visits grew 9% while wait times fell",
        "bullets": ["Visits up 9% on the quarter", "Median wait down to 18 minutes"],
        "notes": "Talk track.",
    },
    {
        "type": "chart",
        "title": "Visits rose every month this quarter",
        "categories": ["Jul", "Aug", "Sep"],
        "series": {"Visits": [4100, 4300, 4470]},
        "y_title": "Visits",
        "source": "visits.csv, Jul-Sep 2026",
    },
]


def checks(report):
    return {i["check"] for i in report["issues"]}


def deck(slide_fn, n=1, path=None, layout=6):
    """A bare python-pptx deck: n slides built by slide_fn(slide, prs, i)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i in range(n):
        s = prs.slides.add_slide(prs.slide_layouts[layout])
        slide_fn(s, prs, i)
    prs.save(path)
    return path


def titled(slide, text="The clinic network grew nine percent this quarter"):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = (
        Inches(0.6),
        Inches(0.45),
        Inches(12.1),
        Inches(0.9),
    )
    t.text = text
    for r in t.text_frame.paragraphs[0].runs:
        r.font.size = Pt(32)
    return t


def box(slide, x, y, w, h, text, size=18, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = text
    for p in tb.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
    if name:
        tb.name = name
    return tb


def test_good_deck_passes_with_a_high_score(out):
    r = build_deck({"slides": GOOD}, out())
    rep = verify_deck(
        r["path"],
        expected_slides=3,
        render=False,
        outline=normalize_outline({"slides": GOOD}),
    )
    assert rep["status"] == "pass", rep["issues"]
    assert rep["style_score"] >= 95
    assert rep["titles"] == [s["title"] for s in GOOD]
    assert (
        rep["slides_overview"][2].startswith("3. Visits rose")
        and "source: visits.csv" in rep["slides_overview"][2]
    )


def test_wrong_slide_count(out):
    r = build_deck({"slides": GOOD}, out())
    rep = verify_deck(r["path"], expected_slides=5, render=False)
    assert "slide_count" in checks(rep) and rep["status"] == "fail"


def test_overflowing_text(out):
    def make(s, prs, i):
        titled(s)
        box(
            s,
            0.6,
            1.6,
            6,
            1.0,
            "A paragraph that is far too long for its little box. " * 12,
        )

    rep = verify_deck(deck(make, path=out(), layout=5), render=False)
    assert "text_overflow" in checks(rep)


def test_picture_instead_of_chart(out, tmp_path):
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    png = tmp_path / "plot.png"
    fig, ax = plt.subplots()
    ax.bar(["a", "b"], [1, 2])
    fig.savefig(png)
    plt.close(fig)

    def make(s, prs, i):
        titled(s, "Visits rose every month this quarter")
        s.shapes.add_picture(str(png), Inches(0.6), Inches(1.6), Inches(8), Inches(4.5))
        box(s, 0.6, 6.8, 6, 0.4, "Source: visits.csv", size=12, name="Source")

    rep = verify_deck(deck(make, path=out(), layout=5), render=False)
    assert "picture_not_chart" in checks(rep)
    rep = verify_deck(
        deck(make, path=out("b.pptx"), layout=5),
        render=False,
        require_native_charts=False,
    )
    assert "picture_not_chart" not in checks(rep)


def test_overlapping_shapes(out):
    def make(s, prs, i):
        titled(s)
        box(s, 0.6, 1.6, 6, 2, "First block of text")
        box(s, 3.0, 2.2, 6, 2, "Second block of text")

    rep = verify_deck(deck(make, path=out(), layout=5), render=False)
    assert "overlap" in checks(rep)


def test_empty_placeholder(out):
    def make(s, prs, i):
        s.shapes.title.text = "Wait times fell to 18 minutes on average"

    rep = verify_deck(deck(make, path=out(), layout=1), render=False)
    assert "empty_placeholder" in checks(rep)


def test_missing_title_tiny_font_and_margins(out):
    def make(s, prs, i):
        box(s, 0.05, 1.0, 5, 1, "tiny footnote-sized text in the body", size=8)

    rep = verify_deck(deck(make, path=out()), render=False)
    c = checks(rep)
    assert {"no_title", "tiny_font", "margin"} <= c


def test_low_contrast_text(out):
    def make(s, prs, i):
        titled(s)
        tb = box(s, 0.6, 1.6, 8, 1, "Pale grey words on a white slide")
        tb.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

    rep = verify_deck(deck(make, path=out(), layout=5), render=False)
    assert "contrast" in checks(rep)


def test_chart_without_source_unit_and_with_single_series_legend(out):
    def make(s, prs, i):
        titled(s)
        cd = CategoryChartData()
        cd.categories = ["a", "b", "c"]
        cd.add_series("Visits", (1, 2, 3))
        gf = s.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(0.6),
            Inches(1.6),
            Inches(10),
            Inches(4.5),
            cd,
        )
        gf.chart.has_legend = True
        gf.chart.value_axis.minimum_scale = 0.5

    rep = verify_deck(deck(make, path=out(), layout=5), render=False)
    assert {
        "missing_source",
        "chart_legend",
        "bar_baseline",
        "axis_unit",
        "alt_text",
    } <= checks(rep)


def test_chart_cache_that_disagrees_with_its_data_sheet(out):
    good = build_deck({"slides": [GOOD[2]]}, out())["path"]
    # Edit the drawn values in the chart XML but not the embedded workbook.
    src = zipfile.ZipFile(good)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as dst:
        for item in src.infolist():
            blob = src.read(item.filename)
            if item.filename.startswith("ppt/charts/chart") and item.filename.endswith(
                ".xml"
            ):
                blob = blob.replace(b"<c:v>4470.0</c:v>", b"<c:v>9999</c:v>")
            dst.writestr(item, blob)
    bad = Path(out("bad.pptx"))
    bad.write_bytes(buf.getvalue())
    rep = verify_deck(
        str(bad), render=False, outline=normalize_outline({"slides": [GOOD[2]]})
    )
    assert {"chart_data_mismatch", "chart_values"} <= checks(rep)


def test_missing_notes_when_required(out):
    r = build_deck({"slides": GOOD}, out())
    rep = verify_deck(r["path"], render=False, require_notes=True)
    missing = [i for i in rep["issues"] if i["check"] == "missing_notes"]
    assert [i["slide"] for i in missing] == [
        3
    ], "the cover is exempt; slide 2 has notes"


def test_too_many_words_bullets_long_titles_and_topic_titles(out):
    wordy = {
        "type": "bullets",
        "title": "This title keeps going well past the point where a reader can take it in at a glance today",
        "bullets": [
            f"Bullet {i} carries far more words than any single bullet on a slide ever should"
            for i in range(8)
        ],
    }
    topic = {"type": "bullets", "title": "Revenue", "bullets": ["Up 4%"]}
    r = build_deck({"slides": [wordy, topic]}, out())
    c = set(r["verify"]["summary"])
    assert {
        "too_many_words",
        "too_many_bullets",
        "bullet_length",
        "title_length",
        "topic_title",
    } <= c
    # The same deck under a user's 'detailed' density and topic-title style.
    rep = verify_deck(
        r["path"], render=False, prefs={"density": "detailed", "title_style": "topic"}
    )
    assert "topic_title" not in checks(rep)


def test_unfilled_todo_slot(out):
    r = build_deck(
        {
            "slides": [
                {
                    "type": "bullets",
                    "title": "TODO: the headline result",
                    "bullets": ["TODO: three takeaways"],
                }
            ]
        },
        out(),
    )
    assert (
        "unfilled_slot" in r["verify"]["summary"] and r["status"] == "built_with_issues"
    )


def test_off_theme_colours_and_font_families(out):
    def make(s, prs, i):
        titled(s)
        for k, (font, rgb) in enumerate(
            [
                ("Comic Sans MS", "FF00FF"),
                ("Courier New", "00FF00"),
                ("Impact", "FF8800"),
            ]
        ):
            tb = box(s, 0.6 + 4 * k, 2, 3.5, 1, f"Text in {font}")
            run = tb.text_frame.paragraphs[0].runs[0]
            run.font.name = font
            run.font.color.rgb = RGBColor.from_string(rgb)

    rep = verify_deck(deck(make, path=out(), layout=5), render=False)
    assert {"font_families", "off_theme_color"} <= checks(rep)


def test_duplicate_titles_and_position_jumps(out):
    def make(s, prs, i):
        t = titled(s, "Every clinic now opens on Saturday mornings")
        if i == 3:
            t.left = Inches(2.5)
        box(s, 0.6, 1.6, 8, 1, "Short point")

    rep = verify_deck(deck(make, n=4, path=out(), layout=5), render=False)
    assert {"duplicate_title", "title_position"} <= checks(rep)


def test_style_score_orders_good_above_bad(out):
    good = build_deck({"slides": GOOD}, out())["verify"]["style_score"]

    def make(s, prs, i):
        box(s, 0.05, 0.2, 13.2, 7, "everything everywhere " * 80, size=9)

    bad = verify_deck(deck(make, n=3, path=out("bad.pptx")), render=False)[
        "style_score"
    ]
    assert good > bad + 30


def test_a_foreign_deck_with_inherited_sizes_is_measured(out):
    """Text with no explicit size inherits from the layout/master; no false 'tiny font'."""

    def make(s, prs, i):
        s.shapes.title.text = "Clinic hours expanded at four sites"
        s.placeholders[1].text_frame.text = "Saturday hours added"

    rep = verify_deck(deck(make, path=out(), layout=1), render=False)
    assert not {"tiny_font", "small_font", "text_overflow"} & checks(rep)


# ---------------------------------------------------------------- render


def test_render_confirms_a_good_deck_and_writes_pngs(out, renderer, tmp_path):
    r = build_deck({"slides": GOOD}, out())
    png_dir = tmp_path / "png"
    png_dir.mkdir()
    rep = verify_deck(
        r["path"], expected_slides=3, renderer=renderer, render_dir=str(png_dir)
    )
    assert rep["render"]["rendered"] and len(rep["render"]["pngs"]) == 3
    assert rep["status"] == "pass", rep["issues"]


def test_render_confirms_overflow_on_the_page(out, renderer):
    def make(s, prs, i):
        from pptx.enum.text import MSO_AUTO_SIZE

        titled(s)
        tb = box(
            s,
            0.6,
            1.6,
            5,
            0.8,
            "A paragraph that is far too long for its little box. " * 10,
        )
        tb.text_frame.auto_size = MSO_AUTO_SIZE.NONE

    rep = verify_deck(deck(make, path=out(), layout=5), renderer=renderer)
    assert {"text_overflow", "render_overflow"} <= checks(rep), rep["issues"]


def test_without_a_renderer_the_report_says_so(out):
    r = build_deck({"slides": GOOD}, out())
    rep = verify_deck(r["path"], render=True, renderer=None)
    assert rep["render"]["rendered"] is False and "estimates" in rep["render"]["reason"]
