# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""What makes the skill usable by a small model: Markdown in, required content shown, fixable errors.

Three failure shapes seen in practice drive these tests:

- a model that writes Markdown well but malformed JSON should still get a deck
  (the Markdown path, with chart/table/kpi blocks);
- content the request asked for must be on the slides, not only in the notes
  (``required_items``);
- an outline written the way small models write them (synonym fields, typos,
  wrappers) is either accepted or rejected with the exact field and a
  corrected example.
"""

from __future__ import annotations

import re

import pytest
from pptx import Presentation
from slidedeck.build import build_deck
from slidedeck.markdown import outline_from_markdown
from slidedeck.outline import OutlineError, normalize_outline
from slidedeck.verify import chart_series, required_item_found, verify_deck

from .conftest import SKILL_DIR

# ---------------------------------------------------------------- markdown path

MD = """# Oakfield Library Services review
Board meeting | Spring 2026
Must show: 41,200 visits; 2 branch closures; e-loans up 18%

## Visits held at 41,200 despite 2 branch closures
- 41,200 visits in the quarter, flat on last year
- 2 branch closures, both temporary
  - Roof repairs at Elm Road and Mill Lane
Source: footfall.csv, Jan-Mar 2026

## E-loans grew 18% and now outnumber print loans
- e-loans up 18% year on year
```chart line
y_title: Loans (thousands)
source: loans.xlsx, 2024-2026
Quarter,E-loans,Print
Q1 24,30.1,34.0
Q1 25,33.9,32.2
Q1 26,40.0,31.5
```

## Two branches carry most of the demand
```table
source: branches.xlsx, Q1 2026
Branch,Visits,Loans
Central,18200,21000
Riverside,9100,10400
Elm Road,2100,2300
```
- Central and Riverside take two thirds of visits

## The headline numbers for the quarter
```kpi
source: footfall.csv; loans.xlsx, Q1 2026
label,value,delta
Visits,41200,0%
E-loans (k),40.0,+18%
```
"""


def test_markdown_with_chart_table_and_kpi_blocks_builds_clean(out):
    outline = outline_from_markdown(MD)
    assert outline["required_items"] == [
        "41,200 visits",
        "2 branch closures",
        "e-loans up 18%",
    ]
    types = [s.get("type") for s in outline["slides"]]
    assert types == ["title", "bullets", "chart", "table", "kpi"]
    r = build_deck(outline, out())
    assert r["verify"]["status"] == "pass", r["verify"]["issues"]
    assert [x["slide"] for x in r["verify"]["required_items"]] == [2, 2, 3]
    chart = next(
        s for s in Presentation(r["path"]).slides[2].shapes if s.has_chart
    ).chart
    got = {s["name"]: s["values"] for s in chart_series(chart)}
    assert got == {"E-loans": [30.1, 33.9, 40.0], "Print": [34.0, 32.2, 31.5]}
    slide4 = Presentation(r["path"]).slides[3]
    assert any(s.has_table for s in slide4.shapes)
    assert "two thirds" in " ".join(
        s.text_frame.text for s in slide4.shapes if s.has_text_frame
    ), "commentary under a table stays visible"


def test_the_skill_md_quick_example_builds_clean(out):
    body = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
    md = re.search(r"````markdown\n(.*?)\n````", body, re.S).group(1)
    r = build_deck(outline_from_markdown(md), out())
    assert r["verify"]["status"] == "pass", r["verify"]["issues"]
    assert all(x["slide"] for x in r["verify"]["required_items"])


def test_skill_md_leads_with_the_no_code_rule():
    body = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8").split("\n---\n", 2)[-1]
    head = body.split("## 1.")[0]
    assert "Do not write python-pptx code" in head
    assert "build_deck_from_markdown" in head and "````markdown" in head
    assert "Never move requested content into the notes" in body
    assert "required_items" in body and "ledger" in body


def test_bad_csv_block_and_bad_json_block_say_how_to_fix():
    with pytest.raises(OutlineError, match="no rows"):
        outline_from_markdown("## A chart\n```chart\ny_title: x\n```\n")
    with pytest.raises(OutlineError, match="CSV block is easier"):
        outline_from_markdown("## A chart\n```slide\n{not json}\n```\n")


# ---------------------------------------------------------------- required items


@pytest.mark.parametrize(
    "item, text, found",
    [
        ("Revenue $12.4M", "Revenue rose to $12.4M in Q3", True),
        ("revenue $12.4m", "REVENUE $12.4M", True),
        ("3 site closures", "Three sites closed: 3 closures in total", True),
        ("41,200 visits", "Visits: 41200 this quarter", True),
        ("e-loans up 18%", "E-loans grew 18% year on year", True),
        ("Revenue $12.4M", "Revenue rose to $12.5M", False),
        ("3 site closures", "Two new sites opened", False),
        ("Supplier risk on part B", "A single supplier for part B is a risk", True),
        ("Supplier risk on part B", "Part B shipped on time", False),
    ],
)
def test_required_item_matching_is_fuzzy_but_strict_on_numbers(item, text, found):
    assert required_item_found(item, text) is found


FRAMEWORK = {
    "type": "diagram",
    "diagram": "pyramid",
    "title": "Three tiers organise our accounts",
    "items": [{"title": "Strategic"}, {"title": "Core"}, {"title": "Long tail"}],
    "notes": "The central message: retention beats acquisition for the top tier.",
}


def test_required_content_only_in_notes_fails(out):
    r = build_deck(
        {"required_items": ["retention beats acquisition"], "slides": [FRAMEWORK]},
        out(),
    )
    assert r["status"] == "built_with_issues"
    assert r["verify"]["summary"].get("required_in_notes_only") == 1
    issue = next(
        i for i in r["verify"]["issues"] if i["check"] == "required_in_notes_only"
    )
    assert "never" in issue["fix"].lower() and "notes" in issue["fix"]


def test_required_content_on_the_slide_passes(out):
    fixed = dict(FRAMEWORK, title="Retention beats acquisition for the top tier")
    r = build_deck(
        {"required_items": ["retention beats acquisition"], "slides": [fixed]}, out()
    )
    assert r["verify"]["status"] == "pass", r["verify"]["issues"]
    assert r["verify"]["required_items"] == [
        {"item": "retention beats acquisition", "slide": 1, "only_in_notes": False}
    ]


def test_missing_required_item_fails_and_verify_takes_items_directly(out):
    r = build_deck({"slides": [FRAMEWORK]}, out())
    rep = verify_deck(r["path"], render=False, required_items=["14 new customers"])
    assert rep["status"] == "fail" and rep["summary"].get("required_missing") == 1


def test_required_items_found_in_chart_labels_and_table_cells(out):
    deck = {
        "required_items": ["Riverside", "18200"],
        "slides": [
            {
                "type": "chart",
                "title": "Riverside is the second busiest branch",
                "categories": ["Central", "Riverside"],
                "series": {"Visits": [18200, 9100]},
                "y_title": "Visits",
                "source": "f.csv",
            },
        ],
    }
    r = build_deck(deck, out())
    assert r["verify"]["status"] == "pass", r["verify"]["issues"]


# ---------------------------------------------------------------- forgiving, corrective outlines


def test_small_model_field_names_are_accepted():
    n = normalize_outline(
        {
            "deck": {
                "pages": [
                    {"Title": "Cover slide"},
                    {
                        "title": "Points",
                        "bullet_points": ["one", "two"],
                        "speaker_note": "say it",
                    },
                    {
                        "title": "Prose",
                        "Content": "A paragraph of plain text for the slide.",
                    },
                    {
                        "title": "Data",
                        "chart_type": "bar",
                        "x_labels": ["a", "b"],
                        "data_series": {"S": [1, 2]},
                    },
                ]
            }
        }
    )
    types = [s["type"] for s in n["slides"]]
    assert types == ["title", "bullets", "bullets", "chart"]
    assert n["slides"][1]["notes"] == "say it"
    assert n["slides"][2]["paras"][0][0].startswith("A paragraph")


def test_typo_fields_name_the_field_and_show_a_corrected_slide():
    with pytest.raises(OutlineError) as e:
        normalize_outline(
            {
                "slides": [
                    {
                        "type": "bullets",
                        "title": "x",
                        "bullett": ["lost point", "another"],
                    }
                ]
            }
        )
    text = "\n".join(e.value.problems)
    assert "unknown field 'bullett'" in text and "did you mean 'bullets'" in text
    assert '"type": "bullets"' in text, "a corrected example is included"


def test_chart_mistakes_get_a_chart_example():
    with pytest.raises(OutlineError) as e:
        normalize_outline(
            {
                "slides": [
                    {
                        "type": "chart",
                        "title": "x",
                        "categories": ["a", "b"],
                        "serie": {"S": [1, 2]},
                    }
                ]
            }
        )
    text = "\n".join(e.value.problems)
    assert "did you mean 'series'" in text
    assert '"chart_type": "column"' in text and '"series": {' in text


def test_harmless_settings_are_ignored_not_errors():
    n = normalize_outline(
        {
            "slides": [
                {
                    "type": "bullets",
                    "title": "x",
                    "bullets": ["a"],
                    "layout": "wide",
                    "id": 3,
                }
            ]
        }
    )
    assert n["slides"][0]["type"] == "bullets"


def test_markdown_passed_as_an_outline_points_at_the_markdown_tool():
    with pytest.raises(OutlineError, match="build_deck_from_markdown"):
        normalize_outline("# My deck\n## Slide\n- point")


def test_an_empty_outline_shows_a_whole_valid_one():
    with pytest.raises(OutlineError) as e:
        normalize_outline({"slides": []})
    assert (
        '"type": "title"' in e.value.problems[0]
        and '"type": "bullets"' in e.value.problems[0]
    )
