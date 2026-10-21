# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""Native chart writers: .xlsx (the chart_into_workbook tool) and .pptx (library).

Every chart is reopened from the saved file and its values compared with the
input; these tests also prove the comparison catches a chart that drifted.
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("openpyxl")

_SKILL = Path(__file__).resolve().parents[2] / "hub" / "skills" / "report-from-sources"
if str(_SKILL) not in sys.path:
    sys.path.insert(0, str(_SKILL))

from report_from_sources_lib import xlsx_chart as X  # noqa: E402

CATS = ["2025Q3", "2025Q4", "2026Q1", "2026Q2"]
SERIES = {"Net revenue": [10.5, 11.25, None, 12.0], "Costs": [7, 8, 9, 10]}


@pytest.mark.parametrize(
    "chart_type", ["line", "column", "bar", "stacked_column", "stacked_bar", "area"]
)
def test_each_chart_type_round_trips(tmp_path, chart_type):
    out = tmp_path / f"{chart_type}.xlsx"
    result = X.write_workbook_chart(
        out,
        CATS,
        SERIES,
        chart_type=chart_type,
        title="Revenue and costs",
        y_axis_title="$ millions",
        source_note="Source: book.xlsx, 2025Q3-2026Q2",
    )
    assert result["verified"] and result["read_back"] == {
        k: list(v) for k, v in SERIES.items()
    }
    from openpyxl import load_workbook

    ws = load_workbook(out)[result["sheet"]]
    assert ws["A1"].value == "Revenue and costs"
    assert any(
        c.value == "Source: book.xlsx, 2025Q3-2026Q2"
        for row in ws.iter_rows()
        for c in row
    )
    assert ws.cell(row=4, column=2).number_format == "#,##0.00"


def test_combo_puts_secondary_series_on_a_line(tmp_path):
    out = tmp_path / "combo.xlsx"
    result = X.write_workbook_chart(
        out, CATS, SERIES, chart_type="combo", secondary=["Costs"]
    )
    assert result["verified"]
    with zipfile.ZipFile(out) as zf:
        xml = zf.read("xl/charts/chart1.xml").decode()
    assert "<c:barChart>" in xml.replace(" ", "") or "barChart" in xml
    assert "lineChart" in xml


@pytest.mark.parametrize(
    "kwargs,message",
    [
        ({"chart_type": "pie3d"}, "not one of"),
        ({"chart_type": "combo"}, "secondary_labels"),
        (
            {"chart_type": "combo", "secondary": ["Net revenue", "Costs"]},
            "some \\(not all\\)",
        ),
    ],
)
def test_bad_requests(tmp_path, kwargs, message):
    with pytest.raises(X.ChartWriteError, match=message):
        X.write_workbook_chart(tmp_path / "x.xlsx", CATS, SERIES, **kwargs)


def test_misaligned_or_empty_series(tmp_path):
    with pytest.raises(X.ChartWriteError, match="4 periods"):
        X.write_workbook_chart(tmp_path / "x.xlsx", CATS, {"A": [1, 2]})
    with pytest.raises(X.ChartWriteError, match="no values"):
        X.write_workbook_chart(tmp_path / "x.xlsx", CATS, {"A": [None] * 4})


def test_adds_a_sheet_to_an_existing_workbook_and_keeps_its_chart(tmp_path):
    out = tmp_path / "book.xlsx"
    first = X.write_workbook_chart(out, CATS, SERIES, title="First")
    second = X.write_workbook_chart(
        out, CATS, {"Costs": SERIES["Costs"]}, title="First"
    )
    assert second["sheet"] != first["sheet"] and second["verified"]
    with zipfile.ZipFile(out) as zf:
        assert sum(n.startswith("xl/charts/chart") for n in zf.namelist()) == 2


def test_refuses_workbooks_openpyxl_would_damage(tmp_path):
    out = tmp_path / "pivot.xlsx"
    X.write_workbook_chart(out, CATS, SERIES)
    with zipfile.ZipFile(out, "a") as zf:
        zf.writestr("xl/pivotTables/pivotTable1.xml", "<pivotTableDefinition/>")
    before = out.read_bytes()
    with pytest.raises(X.ChartWriteError, match="openpyxl\\s+would drop"):
        X.write_workbook_chart(out, CATS, SERIES)
    assert out.read_bytes() == before


def test_verification_catches_a_changed_value(tmp_path):
    out = tmp_path / "v.xlsx"
    result = X.write_workbook_chart(out, CATS, SERIES)
    from openpyxl import load_workbook

    wb = load_workbook(out)
    wb[result["sheet"]].cell(row=2, column=2, value=99)
    wb.save(out)
    check = X.verify_workbook_chart(out, result["sheet"], CATS, SERIES)
    assert not check["verified"] and "Net revenue" in check["mismatches"][0]


# --------------------------------------------------------------------------- #
# PowerPoint library function (consumed by the deck skill)
# --------------------------------------------------------------------------- #


def test_pptx_native_chart_round_trips(tmp_path):
    pytest.importorskip("pptx")
    from report_from_sources_lib.pptx_chart import add_native_chart

    deck = tmp_path / "deck.pptx"
    out = add_native_chart(
        deck,
        "new",
        "line",
        CATS,
        SERIES,
        "Revenue rose",
        "#,##0.0",
        "Source: book.xlsx, 2025Q3-2026Q2",
        y_axis_title="$ millions",
    )
    assert out["verified"] and out["slide"] == 1
    assert out["read_back"] == {k: list(v) for k, v in SERIES.items()}
    from pptx import Presentation

    prs = Presentation(str(deck))
    shapes = prs.slides[0].shapes
    assert sum(1 for s in shapes if s.has_chart) == 1
    assert any(
        s.has_text_frame and s.text_frame.text.startswith("Source: book.xlsx")
        for s in shapes
    )
    again = add_native_chart(
        deck, 1, "column", CATS, {"Costs": SERIES["Costs"]}, "Costs", "", ""
    )
    assert again["verified"] and len(Presentation(str(deck)).slides) == 1


def test_pptx_in_memory_presentation_is_not_saved(tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from report_from_sources_lib.pptx_chart import PptxChartError, add_native_chart

    prs = Presentation()
    out = add_native_chart(
        prs, None, "pie", CATS, {"Costs": SERIES["Costs"]}, "Mix", "0", ""
    )
    assert out["verified"] and out["file"] is None
    with pytest.raises(PptxChartError, match="slide 9"):
        add_native_chart(prs, 9, "line", CATS, SERIES, "", "", "")
    with pytest.raises(PptxChartError, match="one series"):
        add_native_chart(prs, None, "pie", CATS, SERIES, "", "", "")
