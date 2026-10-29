# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""Generated data folders with planted ground truth for the report-from-sources tools.

Every dataset here is invented. Each ``make_*`` function writes a folder and
returns a :class:`Scenario` saying what a correct tool must find: which file is
newest per dataset, the exact values of the last N periods of named series,
which lookups are ambiguous and which series do not exist.
"""

from __future__ import annotations

import csv
import json
import random
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Dict, Iterator, List, Optional, Tuple


@contextmanager
def library_call() -> Iterator[None]:
    """Run the skill's library code as a tool call would, for an agent with
    no path validator (library use: nothing is scope-checked).

    The readers open files through ``gaia.skills.context``, which works only
    inside a skill tool call.
    """
    from gaia.skills.context import bind

    skill = SimpleNamespace(name="report-from-sources", parsed_permissions=tuple)
    with bind(SimpleNamespace(), skill):
        yield


QUARTERS = [(y, q) for y in range(2021, 2027) for q in range(1, 5)]


@dataclass
class Lookup:
    """One series_from_workbook call and what it must return."""

    file: str
    labels: List[str]
    periods: str
    sheet: str = ""
    filters: str = ""
    # Exact expected output, or None when the call must fail.
    expect_periods: Optional[List[str]] = None
    expect_values: Optional[Dict[str, List[Any]]] = None  # matched label -> values
    expect_status: str = "success"  # success | ambiguous | not_found
    expect_candidates: List[str] = field(default_factory=list)


@dataclass
class Scenario:
    name: str
    folder: Path
    newest: Dict[str, str] = field(
        default_factory=dict
    )  # any file of dataset -> newest
    older: List[str] = field(default_factory=list)  # files that must be flagged older
    distinct_datasets: List[Tuple[str, str]] = field(
        default_factory=list
    )  # must NOT merge
    lookups: List[Lookup] = field(default_factory=list)
    narrative: List[str] = field(default_factory=list)


def _q(y: int, q: int) -> str:
    return f"{y}Q{q}"


def _round(rng: random.Random, lo: float, hi: float, nd: int = 1) -> float:
    return round(rng.uniform(lo, hi), nd)


def _series(
    rng: random.Random, n: int, start: float, drift: float, nd: int = 1
) -> List[float]:
    out, v = [], start
    for _ in range(n):
        v += rng.uniform(-drift, drift * 1.5)
        out.append(round(v, nd))
    return out


def _wide_workbook(
    path: Path,
    rng: random.Random,
    quarters: List[Tuple[int, int]],
    data_through: int,
    rows: List[Tuple[Optional[str], str, List[Optional[float]]]],
    *,
    title: str,
    units: str,
    spacer_between_years: bool = False,
    banner: bool = False,
    footnotes: Optional[Dict[Tuple[str, int], str]] = None,
    notes: Optional[List[str]] = None,
) -> Dict[str, Dict[str, Any]]:
    """Write a quarterly wide sheet: title, year row (merged), quarter row, units row.

    ``rows`` are (section, label, values) with values for every quarter;
    values past ``data_through`` are left blank. Returns {label: {quarter: cell value}}.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Table 1"
    ws.append([title])
    ws.append([f"Units: {units}; seasonally adjusted"])
    ws.append([])
    col_of: Dict[Tuple[int, int], int] = {}
    col = 3  # A: label, B: code/units, C.. periods
    year_cells: Dict[int, List[int]] = {}
    prev_year = None
    for y, q in quarters:
        if spacer_between_years and prev_year is not None and y != prev_year:
            col += 1
        col_of[(y, q)] = col
        year_cells.setdefault(y, []).append(col)
        prev_year = y
        col += 1
    header_row = 5 if banner else 4
    if banner:
        split = data_through
        actual = [
            c
            for (yq, c) in zip(quarters, [col_of[x] for x in quarters])
            if quarters.index(yq) < split
        ]
        ws.cell(row=4, column=actual[0], value="Actual")
        ws.merge_cells(
            start_row=4, start_column=actual[0], end_row=4, end_column=actual[-1]
        )
        rest = [col_of[x] for x in quarters[split:]]
        if rest:
            ws.cell(row=4, column=rest[0], value="Forecast")
            if len(rest) > 1:
                ws.merge_cells(
                    start_row=4, start_column=rest[0], end_row=4, end_column=rest[-1]
                )
    for y, cols in year_cells.items():
        ws.cell(row=header_row, column=cols[0], value=y)
        if len(cols) > 1:
            ws.merge_cells(
                start_row=header_row,
                start_column=cols[0],
                end_row=header_row,
                end_column=cols[-1],
            )
    ws.cell(row=header_row + 1, column=1, value="Series")
    for (y, q), c in col_of.items():
        ws.cell(row=header_row + 1, column=c, value=f"Q{q}")
    ws.cell(row=header_row + 2, column=1, value="Units")
    for c in col_of.values():
        ws.cell(row=header_row + 2, column=c, value=units)
    r = header_row + 3
    truth: Dict[str, Dict[str, Any]] = {}
    current_section = None
    for section, label, values in rows:
        if section and section != current_section:
            ws.cell(row=r, column=1, value=section)
            current_section = section
            r += 1
        ws.cell(row=r, column=1, value=("  " if section else "") + label)
        t: Dict[str, Any] = {}
        for i, ((y, q), v) in enumerate(zip(quarters, values)):
            if i >= data_through or v is None:
                continue
            cell_value: Any = v
            marker = (footnotes or {}).get((label, i))
            if marker:
                cell_value = f"{v}{marker}"
            ws.cell(row=r, column=col_of[(y, q)], value=cell_value)
            t[_q(y, q)] = v
        key = f"{section} > {label}" if section else label
        truth[key] = t
        r += 1
    r += 1
    for line in notes or []:
        ws.cell(row=r, column=1, value=line)
        r += 1
    wb.save(path)
    return truth


def _last(
    truth: Dict[str, Any], periods: List[str], n: int
) -> Tuple[List[str], List[Any]]:
    have = [p for p in periods if p in truth]
    picked = have[-n:]
    return picked, [truth[p] for p in picked]


def make_quarterly_vintage(root: Path, seed: int) -> Scenario:
    """Scenario 1: a wide quarterly workbook next to an older vintage of itself."""
    rng = random.Random(seed)
    folder = root / f"s1_vintage_{seed}"
    folder.mkdir(parents=True)
    names = [
        "Manufacturing output",
        "Construction output",
        "Services output",
        "Retail trade",
        "Transport and storage",
    ]
    rng.shuffle(names)
    names = names[:4]
    quarters = QUARTERS[: 4 * 6 - 2]  # 2021Q1 .. 2026Q2 (22 quarters)
    labels_q = [_q(*x) for x in quarters]
    base = {n: _series(rng, len(quarters), _round(rng, 50, 900), 15) for n in names}
    old_rows = [
        (None, n, [round(v - rng.uniform(0.1, 2), 1) for v in base[n]]) for n in names
    ]
    new_rows = [(None, n, base[n]) for n in names]
    stem = rng.choice(["regional-output", "regional_output", "Regional Output"])
    old_name = f"{stem}-second-quarter-2025.xlsx"
    new_name = f"{stem}_2026Q2.xlsx"
    # The old file is written second so its modified time is newer: mtime must lose.
    new_truth = _wide_workbook(
        folder / new_name,
        rng,
        quarters + [(2026, 3), (2026, 4)],
        len(quarters),
        [(s, l, v + [None, None]) for s, l, v in new_rows],
        title="Regional output by industry",
        units="$ millions",
    )
    _wide_workbook(
        folder / old_name,
        rng,
        quarters[:18],
        18,
        [(s, l, v[:18]) for s, l, v in old_rows],
        title="Regional output by industry",
        units="$ millions",
    )
    other = "household_survey_2026.csv"
    with open(folder / other, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["year", "respondents", "response_rate"])
        for y in range(2018, 2027):
            w.writerow([y, rng.randint(900, 1500), _round(rng, 40, 70)])
    from docx import Document

    doc = Document()
    doc.add_paragraph("Quarterly commentary as of 30 June 2026.")
    doc.save(folder / "commentary.docx")
    n = rng.choice([4, 5, 6])
    pick = names[:2]
    exp_vals = {}
    exp_periods = None
    for name in pick:
        exp_periods, vals = _last(new_truth[name], labels_q, n)
        exp_vals[name] = vals
    return Scenario(
        "quarterly_vintage",
        folder,
        newest={old_name: new_name, new_name: new_name},
        older=[old_name],
        distinct_datasets=[(new_name, other)],
        lookups=[
            Lookup(
                new_name,
                [p.lower() for p in pick],
                f"last {n}",
                expect_periods=exp_periods,
                expect_values=exp_vals,
            )
        ],
        narrative=["commentary.docx"],
    )


MONTHS = [(y, m) for y in range(2022, 2027) for m in range(1, 13)]


def make_long_monthly_csv(root: Path, seed: int) -> Scenario:
    """Scenario 2: long-format monthly CSV (date, series, value), CRLF + BOM, Unicode names."""
    rng = random.Random(seed)
    folder = root / f"s2_long_{seed}"
    folder.mkdir(parents=True)
    months = MONTHS[: 12 * 4 + rng.randint(1, 6)]
    series_names = [
        "Exports – goods",
        "Imports – goods",
        "Café & restaurant sales",
        "Energy prices index",
    ]
    data = {
        s: _series(rng, len(months), _round(rng, 80, 400), 6, 2) for s in series_names
    }
    gap_series = series_names[1]
    gap_at = len(months) - 3
    name = f"trade_monthly_{months[-1][0]}{months[-1][1]:02d}.csv"
    rows = []
    for i, (y, m) in enumerate(months):
        for s in series_names:
            v: Any = data[s][i]
            if s == gap_series and i == gap_at:
                v = ""
            rows.append([f"{y}-{m:02d}-01", s, v])
    rng.shuffle(rows)  # long files are not always sorted
    with open(folder / name, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(["date", "series", "value"])
        w.writerows(rows)
    n = rng.choice([6, 12])
    labels = [f"{y}-{m:02d}" for y, m in months]
    want = [series_names[0], series_names[1]]
    exp = {}
    for s in want:
        vals = [
            None if (s == gap_series and i == gap_at) else data[s][i]
            for i in range(len(months))
        ]
        exp[s] = vals[-n:]
    return Scenario(
        "long_monthly_csv",
        folder,
        newest={name: name},
        lookups=[
            Lookup(
                name,
                ["exports - goods", "Imports – goods"],
                f"last {n}",
                expect_periods=labels[-n:],
                expect_values=exp,
            )
        ],
    )


def make_messy_sheet(root: Path, seed: int) -> Scenario:
    """Scenario 3: banner + merged year headers, spacer columns, notes, footnote markers."""
    rng = random.Random(seed)
    folder = root / f"s3_messy_{seed}"
    folder.mkdir(parents=True)
    quarters = QUARTERS[8:24]  # 2023Q1..2026Q4
    through = 14  # data through 2026Q2, forecast blank afterwards
    labels_q = [_q(*x) for x in quarters]
    rows = []
    for section, items in (
        ("Revenue", ["Product sales", "Licensing", "Total"]),
        ("Costs", ["Staff", "Materials", "Total"]),
    ):
        for item in items:
            rows.append(
                (section, item, _series(rng, len(quarters), _round(rng, 20, 300), 8))
            )
    foot = {
        ("Licensing", through - 2): "(a)",
        ("Staff", through - 1): "(b)",
        ("Product sales", through - 1): "*",
    }
    name = "finance_pack_FY26_v3.xlsx"
    truth = _wide_workbook(
        folder / name,
        rng,
        quarters,
        through,
        rows,
        title="Group finance pack (USD thousands)",
        units="USD '000",
        spacer_between_years=True,
        banner=True,
        footnotes=foot,
        notes=[
            "Notes:",
            "(a) Includes a one-off settlement.",
            "(b) Restated for the new cost allocation.",
            "* Preliminary.",
            "Source: internal ledger",
        ],
    )
    n = 4
    exp = {}
    per = None
    for key in ("Revenue > Licensing", "Costs > Staff"):
        per, vals = _last(truth[key], labels_q, n)
        exp[key.split(" > ")[1]] = vals
    total_per, total_vals = _last(truth["Costs > Total"], labels_q, n)
    return Scenario(
        "messy_sheet",
        folder,
        newest={name: name},
        lookups=[
            Lookup(
                name,
                ["Licensing", "Staff"],
                f"last {n}",
                expect_periods=per,
                expect_values=exp,
            ),
            Lookup(
                name,
                ["Costs > Total"],
                "last 4",
                expect_periods=total_per,
                expect_values={"Total": total_vals},
            ),
        ],
    )


def make_fiscal_mixed(root: Path, seed: int) -> Scenario:
    """Scenario 4: fiscal-year columns, and a month column written in five date formats."""
    from openpyxl import Workbook

    rng = random.Random(seed)
    folder = root / f"s4_fiscal_{seed}"
    folder.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Annual"
    years = list(range(2019, 2027))
    style = rng.choice(["FY{y}", "FY {p}/{s}", "{p}-{s}"])
    heads = [style.format(y=y, p=y - 1, s=str(y)[2:]) for y in years]
    ws.append(["Budget outturn by fiscal year"])
    ws.append(["Line item"] + heads)
    items = {
        "Tax receipts": _series(rng, len(years), 500, 30),
        "Grants paid": _series(rng, len(years), 120, 10),
    }
    for k, v in items.items():
        ws.append([k] + v)
    ws2 = wb.create_sheet("Monthly")
    ws2.append(["Month", "Admissions", "Discharges"])
    months = MONTHS[24 : 24 + 18]  # 2024-01 .. 2025-06
    adm = [rng.randint(800, 1500) for _ in months]
    dis = [rng.randint(700, 1400) for _ in months]
    fmts = [
        lambda y, m: datetime(y, m, 1),
        lambda y, m: f"{y}-{m:02d}-28",
        lambda y, m: date(y, m, 1).strftime("%b %Y"),
        lambda y, m: f"{m:02d}/15/{y}",
        lambda y, m: date(y, m, 1).strftime("%B %Y"),
    ]
    for i, (y, m) in enumerate(months):
        ws2.append([fmts[(i + seed) % len(fmts)](y, m), adm[i], dis[i]])
    name = "public_services_2025.xlsx"
    wb.save(folder / name)
    fy_labels = [f"FY{y}" for y in years]
    m_labels = [f"{y}-{m:02d}" for y, m in months]
    return Scenario(
        "fiscal_mixed",
        folder,
        newest={name: name},
        lookups=[
            Lookup(
                name,
                ["Tax receipts"],
                "last 3",
                sheet="Annual",
                expect_periods=fy_labels[-3:],
                expect_values={"Tax receipts": items["Tax receipts"][-3:]},
            ),
            Lookup(
                name,
                ["Admissions", "Discharges"],
                "2024-10:2025-03",
                sheet="Monthly",
                expect_periods=m_labels[9:15],
                expect_values={"Admissions": adm[9:15], "Discharges": dis[9:15]},
            ),
        ],
    )


def make_ambiguous(root: Path, seed: int) -> Scenario:
    """Scenario 5: similar labels; the tool must report ambiguity, not guess."""
    rng = random.Random(seed)
    folder = root / f"s5_ambiguous_{seed}"
    folder.mkdir(parents=True)
    quarters = QUARTERS[12:22]
    pairs = rng.choice(
        [
            ("Net sales", "Net sales, adjusted", "sales"),
            ("Core inflation", "Core inflation (ex food)", "core"),
            ("Operating margin", "Operating margin excl. one-offs", "margin"),
        ]
    )
    rows = [
        (None, pairs[0], _series(rng, len(quarters), 100, 5)),
        (None, pairs[1], _series(rng, len(quarters), 100, 5)),
        ("North", "Total", _series(rng, len(quarters), 50, 5)),
        ("South", "Total", _series(rng, len(quarters), 50, 5)),
    ]
    name = "kpi_book.xlsx"
    truth = _wide_workbook(
        folder / name, rng, quarters, len(quarters), rows, title="KPI book", units="%"
    )
    labels_q = [_q(*x) for x in quarters]
    per, vals = _last(truth[pairs[0]], labels_q, 3)
    return Scenario(
        "ambiguous",
        folder,
        newest={name: name},
        lookups=[
            Lookup(
                name,
                [pairs[2]],
                "last 3",
                expect_status="ambiguous",
                expect_candidates=[pairs[0], pairs[1]],
            ),
            Lookup(
                name,
                ["Total"],
                "last 3",
                expect_status="ambiguous",
                expect_candidates=["Total"],
            ),
            # The exact name is not ambiguous even though a longer one contains it.
            Lookup(
                name,
                [pairs[0]],
                "last 3",
                expect_periods=per,
                expect_values={pairs[0]: vals},
            ),
        ],
    )


def make_negative(root: Path, seed: int) -> Scenario:
    """Scenario 6: series that are not there; a folder with nothing to merge."""
    rng = random.Random(seed)
    folder = root / f"s6_negative_{seed}"
    folder.mkdir(parents=True)
    quarters = QUARTERS[16:24]
    rows = [
        (None, "Visitor numbers", _series(rng, len(quarters), 300, 20)),
        (None, "Average stay (nights)", _series(rng, len(quarters), 3, 0.2, 2)),
    ]
    name = "tourism_2026.xlsx"
    _wide_workbook(
        folder / name, rng, quarters, 6, rows, title="Tourism", units="thousands"
    )
    unrelated = "rainfall_2026.json"
    (folder / unrelated).write_text(
        json.dumps(
            {
                "data": [
                    {"month": f"2026-{m:02d}", "rainfall_mm": rng.randint(10, 120)}
                    for m in range(1, 7)
                ]
            }
        ),
        encoding="utf-8",
    )
    return Scenario(
        "negative",
        folder,
        newest={name: name, unrelated: unrelated},
        distinct_datasets=[(name, unrelated)],
        lookups=[
            Lookup(
                name,
                [rng.choice(["Hotel occupancy rate", "Headcount", "Airport arrivals"])],
                "last 4",
                expect_status="not_found",
            ),
            Lookup(unrelated, ["snowfall_mm"], "last 3", expect_status="not_found"),
        ],
    )


SCENARIOS = [
    make_quarterly_vintage,
    make_long_monthly_csv,
    make_messy_sheet,
    make_fiscal_mixed,
    make_ambiguous,
    make_negative,
]
