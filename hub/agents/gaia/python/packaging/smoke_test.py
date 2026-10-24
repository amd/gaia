# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""
No-Python smoke test for the frozen GAIA flagship agent.

One binary serves two transports, and BOTH are spawned in production, so both
are checked here against the FROZEN BINARY (not ``python -m ...``).

HTTP -- what the daemon and the Agent UI use:

  1. Launch the binary as a subprocess -- binary only, no interpreter available
     to it. The argv is the daemon's own (``--host``/``--port``, no ``--serve``),
     because that is the invocation that ships.
  2. Poll ``GET /health`` until ready (dependency-free readiness probe).
  3. ``GET /health``       -> 200 ``{"status": "ok", ...}``.
  4. ``GET /version``      -> 200 with BOTH ``apiVersion`` and ``agentVersion``
     non-empty. The daemon reads exactly these two keys to decide whether to
     attach, so an empty one is a contract break, not cosmetic.
  5. ``GET /openapi.json`` -> contains ``/v1/gaia/query`` and ``/v1/gaia/init``.

stdio JSONL -- what the TUI spawns as a child process:

  6. ``--json-events --help`` reaches the STDIO parser. The uvicorn-only binary
     answered that flag with ``unrecognized arguments: --json-events``, which is
     the whole reason no shipped build could drive the TUI.
  7. Spawned the way the TUI spawns it -- bare, no arguments -- it writes its
     opening ``status`` event as the FIRST line of stdout, then exits when stdin
     closes. This is the handshake the TUI reads, and it proves three things at
     once: the agent constructed inside the frozen binary (a hidden import
     PyInstaller missed shows up as an ``error`` event naming it), stdout
     carries JSON and nothing else, and the process tears down when its parent
     leaves.
  8. The same handshake with ``--use-claude`` names the ``claude`` backend.
     The TUI forwards that flag, and the Anthropic client's import is guarded,
     so a binary built without it boots fine and fails only here.

run_python -- the binary is the only Python a release user has:

  9. ``--gaia-run-python-file <script>`` runs a script that imports the
     bundled office/data libraries, writes an ``.xlsx``, starts a process
     pool, and starts ``sys.executable -c`` / ``-m pip`` itself; a failing
     script's exit code comes back unchanged. Before the fix the script path
     reached the stdio parser and exited 2 with ``unrecognized arguments``.

Tokenizer -- token counts drive history trimming and context eviction:

 10. With tiktoken's download host unreachable and no cache folder set, the
     agent's tokenizer is still ``tiktoken:cl100k_base``: the encoding plugin
     and its data are in the binary, not fetched at runtime (#5066).

Linux only -- the binary carries its build distro's OpenSSL:

 11. A ``run_python`` script's default SSL context holds CA certificates, so
     HTTPS from a script verifies on a distro other than the one it was
     built on (``rthook_linux_ca.py``).

The handshake deliberately stops short of a real turn: an answer needs a running
Lemonade Server and a downloaded model, which the release runners do not have.
It still fails loudly on a broken binary -- the opening event is emitted with
Lemonade absent, and reports it as unreachable rather than erroring.

This harness runs under Python (it is a test driver), but the process under test
is the frozen binary. Uses only the stdlib so it has no install requirements.

Exit code 0 = PASS, non-zero = FAIL. Verbose ``[smoke]`` logging throughout.
"""

from __future__ import annotations

import argparse
import json
import os
import queue
import socket
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
from pathlib import Path

HOST = "127.0.0.1"
PORT = 8142  # 8141 is the gaia sidecar's runtime port; 8131 is email. NEVER 4001.
BASE = f"http://{HOST}:{PORT}"

REQUIRED_PATHS = {"/v1/gaia/query", "/v1/gaia/init"}

#: The flagship's catalog entry declares no ``BinaryArgs``, so the TUI spawns
#: the binary bare and only appends flags for dev mode / Claude / full access
#: (``tui/internal/client/factory.go``). Bare argv IS the shipping invocation,
#: which is why "no arguments" has to mean the stdio wire.
STDIO_ARGV: list[str] = []

#: Agent construction (embedding validation, two FAISS index rebuilds, the
#: scratchpad DB, the filesystem index) runs before the opening event, and a
#: cold one-file binary unpacks itself first. Generous on purpose: a slow
#: runner must not read as a broken binary.
STDIO_HANDSHAKE_DEADLINE_S = 300.0

#: ``gaia_agent_chat.python_runner.RUN_PYTHON_FILE_FLAG``, copied because this
#: harness imports nothing outside the stdlib (a unit test pins the two).
RUN_PYTHON_FILE_FLAG = "--gaia-run-python-file"

#: What a typical ``run_python`` call does: office files, a table, a
#: chart, and stdlib modules nothing in GAIA itself imports.
RUN_PYTHON_SCRIPT = """\
import sys, statistics, sqlite3, zipfile, xml.etree.ElementTree, csv, decimal
import subprocess
from concurrent.futures import ProcessPoolExecutor

import openpyxl, docx, pptx, pandas, numpy, reportlab, pypdf, pymupdf, xlsxwriter
import matplotlib.pyplot as plt


def py(*argv):
    # A script starting sys.executable -- the binary -- as an interpreter.
    return subprocess.run([sys.executable, *argv], capture_output=True, text=True)


def main():
    wb = openpyxl.Workbook()
    wb.active.append(["part", "qty"])
    wb.active.append(["bolt", 4])
    wb.save("bom.xlsx")
    assert openpyxl.load_workbook("bom.xlsx").active["B2"].value == 4
    docx.Document().save("notes.docx")
    pptx.Presentation().save("deck.pptx")
    plt.plot([1, 2, 3])
    plt.savefig("chart.png")
    frame = pandas.DataFrame({"qty": [4, 6]})
    # Workers re-launch the binary. Under spawn (macOS) they re-import this
    # file as __mp_main__, as CPython does, so the work sits under the guard.
    with ProcessPoolExecutor(2) as pool:
        assert list(pool.map(abs, [-1, -2])) == [1, 2]
    nested = py("-c", "import pypdf, pymupdf, xlsxwriter; print('nested ok')")
    assert nested.stdout.strip().endswith("nested ok"), (nested.stdout, nested.stderr)
    pip = py("-m", "pip", "install", "pypdf")
    assert pip.returncode == 1 and "can't be installed" in pip.stderr, pip.stderr
    total = int(frame["qty"].sum())
    print("RUN_PYTHON_OK", sys.argv[1:], total, statistics.mean([1, 3]))


if __name__ == "__main__":
    main()
"""


def log(msg: str) -> None:
    print(f"[smoke] {msg}", flush=True)


def _get(path: str, timeout: float = 5.0):
    req = urllib.request.Request(BASE + path, method="GET")
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.status, json.loads(resp.read().decode("utf-8"))


def _kill_tree(proc: subprocess.Popen) -> None:
    """Kill the server AND its children.

    PyInstaller's one-file bootloader spawns a child process; terminating the
    parent orphans the child and leaves the socket bound. A host app must do the
    same on shutdown.
    """
    if proc.poll() is not None:
        return
    if sys.platform == "win32":
        subprocess.run(
            ["taskkill", "/F", "/T", "/PID", str(proc.pid)],
            capture_output=True,
            check=False,
        )
    else:
        proc.terminate()
    try:
        proc.wait(timeout=10)
    except subprocess.TimeoutExpired:
        proc.kill()


def _drain(proc: subprocess.Popen, label: str, tail_chars: int = 4000) -> None:
    try:
        if proc.stdout:
            out = proc.stdout.read()
            if out:
                log(f"{label}:\n{out[-tail_chars:]}")
    except OSError as exc:
        log(f"could not drain server output ({label}): {exc}")


def wait_for_health(proc: subprocess.Popen, deadline_s: float = 90.0) -> bool:
    start = time.time()
    while time.time() - start < deadline_s:
        if proc.poll() is not None:
            log(f"server process exited early with code {proc.returncode}")
            return False
        try:
            status, body = _get("/health", timeout=2.0)
            if status == 200 and body.get("status") == "ok":
                log(f"/health ready after {time.time() - start:.1f}s -> {body}")
                return True
            log(f"/health answered but not ready: HTTP {status} {body}")
        except (urllib.error.URLError, ConnectionError, OSError):
            pass
        time.sleep(0.5)
    log(f"/health not ready within {deadline_s}s")
    return False


def check_health() -> bool:
    status, body = _get("/health", timeout=5.0)
    log(f"/health -> HTTP {status} {body}")
    if status != 200:
        log(f"FAIL: /health returned HTTP {status}")
        return False
    if body.get("status") != "ok":
        log(f"FAIL: /health status is {body.get('status')!r}, expected 'ok'")
        return False
    log("health check PASS")
    return True


def check_version() -> bool:
    status, body = _get("/version", timeout=5.0)
    log(f"/version -> HTTP {status} {body}")
    if status != 200:
        log(f"FAIL: /version returned HTTP {status}")
        return False
    missing = [k for k in ("apiVersion", "agentVersion") if not body.get(k)]
    if missing:
        log(f"FAIL: /version has empty/absent field(s): {missing}")
        return False
    log(
        f"version check PASS (apiVersion={body['apiVersion']}, "
        f"agentVersion={body['agentVersion']})"
    )
    return True


def check_openapi() -> bool:
    status, spec = _get("/openapi.json", timeout=15.0)
    if status != 200:
        log(f"FAIL: /openapi.json returned HTTP {status}")
        return False
    paths = set(spec.get("paths", {}).keys())
    log(f"openapi paths: {sorted(paths)}")
    missing = REQUIRED_PATHS - paths
    if missing:
        log(f"FAIL: missing gaia paths in openapi: {sorted(missing)}")
        return False
    log("openapi check PASS -- all required gaia paths present")
    return True


def _log_stderr_tail(proc: subprocess.Popen, tail_chars: int = 4000) -> None:
    """Report the child's stderr. Only safe once the child is dead."""
    try:
        if proc.stderr:
            err = proc.stderr.read()
            if err:
                log(f"stdio stderr tail:\n{err[-tail_chars:]}")
    except (OSError, ValueError) as exc:
        log(f"could not read stderr: {exc}")


def check_stdio_argv(binary: Path) -> bool:
    """The TUI's flags must reach the stdio parser, not the HTTP one.

    ``--help`` is answered by whichever parser the dispatch picked, so the help
    TEXT is the evidence: the HTTP parser cannot print flags it does not own.
    """
    result = subprocess.run(
        [str(binary), "--json-events", "--help"],
        capture_output=True,
        text=True,
        timeout=120,
    )
    if result.returncode != 0:
        log(
            f"FAIL: `--json-events --help` exited {result.returncode}. "
            f"stdout: {result.stdout[-800:]!r} stderr: {result.stderr[-800:]!r}"
        )
        return False
    missing = [flag for flag in ("--json-events", "--dev") if flag not in result.stdout]
    if missing:
        log(
            f"FAIL: the help text is missing {missing} -- `--json-events` was "
            "parsed by the HTTP transport instead of the stdio one."
        )
        log(f"help text was:\n{result.stdout[-2000:]}")
        return False
    log("stdio argv check PASS -- the TUI's flags reach the stdio parser")
    return True


def _first_stdout_line(proc: subprocess.Popen, deadline_s: float):
    """The first line the child writes, or ``None`` if it never writes one.

    Read on a thread because the child may exit, block, or print nothing at all,
    and a bare ``readline()`` on any of those hangs the harness forever.
    """
    lines: "queue.Queue[str]" = queue.Queue()

    def _pump() -> None:
        line = proc.stdout.readline()
        if line:
            lines.put(line)

    threading.Thread(target=_pump, daemon=True).start()
    try:
        return lines.get(timeout=deadline_s)
    except queue.Empty:
        return None


def check_stdio_handshake(
    binary: Path,
    argv: list[str] = STDIO_ARGV,
    env: dict | None = None,
    backend: str | None = None,
) -> bool:
    """The opening event of a real stdio session, read off the real pipe.

    *backend*, when given, is the ``model_backend`` the opening event must name.
    """
    cmd = [str(binary), *argv]
    log(f"command: {' '.join(cmd) if argv else str(binary) + ' (no arguments)'}")
    t0 = time.time()
    proc = subprocess.Popen(
        cmd,
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        bufsize=1,
        env=env,
    )
    try:
        line = _first_stdout_line(proc, STDIO_HANDSHAKE_DEADLINE_S)
        if line is None:
            state = (
                f"exited {proc.returncode}"
                if proc.poll() is not None
                else "still running"
            )
            log(
                f"FAIL: no event on stdout within "
                f"{STDIO_HANDSHAKE_DEADLINE_S}s (process {state})"
            )
            # Killed first: the child may still be alive, and reading a live
            # pipe to EOF would hang the harness instead of reporting.
            _kill_tree(proc)
            _log_stderr_tail(proc)
            return False
        log(f"first stdout line after {time.time() - t0:.1f}s: {line.strip()[:400]}")

        try:
            event = json.loads(line)
        except ValueError as exc:
            log(
                f"FAIL: the first line of stdout is not JSON ({exc}). stdout is "
                "the event wire -- the TUI reads a stray line as a malformed "
                "event, so nothing may print to it before the handshake."
            )
            return False
        if not isinstance(event, dict) or not event.get("type"):
            log(f"FAIL: the opening event carries no 'type': {event!r}")
            return False
        if event["type"] != "status":
            # An 'error' here is the frozen binary failing to build its agent --
            # a missing hidden import or uncollected data file. Its detail names
            # the cause, so print it rather than a generic verdict. The 'detail'
            # is only the exception's str() (see _terminal_error/stdio.py's
            # construction-failure handler) -- the full traceback goes to
            # stderr via traceback.format_exc(), and lives only there. Dumping
            # it here is what turned a bare "IndexError: 4" (v0.2.0's release,
            # useless on its own) into an actionable finding.
            log(
                f"FAIL: expected the opening 'status' event, got "
                f"{event['type']!r}: {event.get('detail') or event!r}"
            )
            _kill_tree(proc)
            _log_stderr_tail(proc)
            return False

        # Reported, not asserted: the release runners have no model server, and
        # the handshake is emitted either way.
        log(
            f"opening status: model={event.get('model_id')!r} "
            f"backend={event.get('model_backend')!r} "
            f"lemonade_reachable={event.get('lemonade_reachable')}"
        )
        if backend is not None and event.get("model_backend") != backend:
            log(
                f"FAIL: expected model_backend={backend!r}, got "
                f"{event.get('model_backend')!r}"
            )
            return False

        # stdin closing is how the TUI says goodbye; a child that ignores it
        # holds the model slot for the life of the machine.
        proc.stdin.close()
        try:
            proc.wait(timeout=60)
        except subprocess.TimeoutExpired:
            log("FAIL: the agent did not exit within 60s of stdin closing")
            return False
        log(f"stdio handshake check PASS -- exited {proc.returncode} on stdin close")
        return True
    finally:
        _kill_tree(proc)


def _run_script(binary: Path, workdir: Path, body: str, *args: str, env=None):
    script = workdir / "snippet.py"
    script.write_text(body, encoding="utf-8")
    cmd = [str(binary), RUN_PYTHON_FILE_FLAG, str(script), *args]
    log(f"command: {' '.join(cmd)}")
    t0 = time.time()
    # The tool's own call shape: piped output, closed stdin, script's cwd.
    result = subprocess.run(
        cmd,
        cwd=str(workdir),
        capture_output=True,
        stdin=subprocess.DEVNULL,
        encoding="utf-8",
        errors="replace",
        timeout=300,
        check=False,
        env=env,
    )
    log(f"exited {result.returncode} after {time.time() - t0:.1f}s")
    return result


def check_run_python(binary: Path) -> bool:
    """A script runs on the bundled libraries, and its exit code comes back."""
    with tempfile.TemporaryDirectory(prefix="gaia-smoke-run-python-") as tmp:
        workdir = Path(tmp)
        ok = _run_script(binary, workdir, RUN_PYTHON_SCRIPT, "--rows", "2")
        if ok.returncode != 0 or "RUN_PYTHON_OK ['--rows', '2'] 10 2" not in ok.stdout:
            log(
                "FAIL: the run_python script did not run on the frozen binary. "
                f"stdout: {ok.stdout[-1500:]!r} stderr: {ok.stderr[-3000:]!r}"
            )
            return False
        missing = [
            name
            for name in ("bom.xlsx", "notes.docx", "deck.pptx", "chart.png")
            if not (workdir / name).is_file()
        ]
        if missing:
            log(f"FAIL: the script ran but did not write {missing}")
            return False

        failed = _run_script(binary, workdir, "import sys\nsys.exit(3)\n")
        if failed.returncode != 3:
            log(
                f"FAIL: a script exiting 3 came back as {failed.returncode}. "
                f"stderr: {failed.stderr[-1500:]!r}"
            )
            return False
    log("run_python check PASS -- scripts run on the bundled libraries")
    return True


#: Prints the estimator the agent's token counts use, from inside the binary.
TOKENIZER_SCRIPT = """import sys
from gaia.agents.base.turn_metrics import _tokenizer
name = _tokenizer().name
print("TOKENIZER", name)
sys.exit(0 if name == "tiktoken:cl100k_base" else 5)
"""

#: Every request through a closed port, so a download attempt fails fast.
_UNREACHABLE_PROXY = "http://127.0.0.1:9"


def check_tokenizer(binary: Path) -> bool:
    """cl100k_base loads from the binary alone, with its download host unreachable."""
    dropped = {"TIKTOKEN_CACHE_DIR", "DATA_GYM_CACHE_DIR", "NO_PROXY"}
    env = {k: v for k, v in os.environ.items() if k.upper() not in dropped}
    for name in ("HTTPS_PROXY", "HTTP_PROXY", "ALL_PROXY"):
        env[name] = env[name.lower()] = _UNREACHABLE_PROXY
    with tempfile.TemporaryDirectory(prefix="gaia-smoke-tokenizer-") as tmp:
        # A fresh temp dir, so tiktoken's default cache cannot hold the file.
        env["TMPDIR"] = env["TEMP"] = env["TMP"] = tmp
        result = _run_script(binary, Path(tmp), TOKENIZER_SCRIPT, env=env)
    log(f"tokenizer: {result.stdout.strip()[-300:]}")
    if result.returncode != 0:
        log(
            "FAIL: the agent's token counts are an estimate, not cl100k_base: "
            "the binary is missing tiktoken's encoding plugin or its data. "
            f"stderr: {result.stderr[-3000:]!r}"
        )
        return False
    log("tokenizer check PASS -- cl100k_base loads offline from the binary")
    return True


#: Prints how many CA certificates a script's default SSL context loaded.
TLS_TRUST_SCRIPT = """import ssl, sys
count = ssl.create_default_context().cert_store_stats()["x509_ca"]
print("CA_CERTS", count, ssl.get_default_verify_paths())
sys.exit(0 if count else 4)
"""


def check_tls_trust(binary: Path) -> bool:
    """A script can verify HTTPS: its default context found the host's CAs."""
    with tempfile.TemporaryDirectory(prefix="gaia-smoke-tls-") as tmp:
        result = _run_script(binary, Path(tmp), TLS_TRUST_SCRIPT)
    log(f"tls: {result.stdout.strip()[-600:]}")
    if result.returncode != 0:
        log(
            "FAIL: a run_python script's default SSL context holds no CA "
            "certificates, so its HTTPS calls fail verification on this distro. "
            f"stderr: {result.stderr[-1500:]!r}"
        )
        return False
    log("tls trust check PASS -- scripts find the host's CA certificates")
    return True


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="Smoke-test the frozen GAIA agent binary (both transports)."
    )
    parser.add_argument(
        "binary", help="Path to the frozen gaia-agent executable to test."
    )
    args = parser.parse_args(argv)

    binary = Path(args.binary).resolve()
    if not binary.exists():
        log(f"FAIL: binary not found: {binary}")
        return 2

    # Preflight: refuse if the port is already bound -- otherwise we would
    # health-check a stale server and report a false PASS.
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.settimeout(2.0)
        if s.connect_ex((HOST, PORT)) == 0:
            log(f"FAIL: port {PORT} already in use -- kill the stale server first.")
            return 2

    # The daemon's own argv: bind flags, no --serve. That is what ships, so that
    # is what is tested; the --serve spelling is covered by the unit tests.
    cmd = [str(binary), "--host", HOST, "--port", str(PORT)]
    log(f"launching frozen binary: {binary}")
    log(f"command: {' '.join(cmd)}")
    t0 = time.time()
    proc = subprocess.Popen(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
    )

    results: dict[str, bool] = {}
    try:
        ready = wait_for_health(proc)
        log(f"startup time to /health: {time.time() - t0:.1f}s")
        if not ready:
            return 3  # the finally block kills the tree and drains the log
        results["health"] = check_health()
        results["version"] = check_version()
        results["openapi"] = check_openapi()
    finally:
        log("shutting down server")
        _kill_tree(proc)
        _drain(proc, "server output tail")

    # Same binary, other transport. Run second so a broken HTTP surface is
    # reported before the slower agent construction.
    log("checking the stdio transport")
    results["stdio_argv"] = check_stdio_argv(binary)
    results["stdio_handshake"] = check_stdio_handshake(binary)

    # Building the Claude client makes no network call, so a placeholder key is
    # enough to prove the client is bundled.
    log("checking the --use-claude backend")
    claude_env = dict(os.environ)
    claude_env.setdefault("ANTHROPIC_API_KEY", "smoke-test-placeholder")
    results["claude_handshake"] = check_stdio_handshake(
        binary, ["--use-claude"], env=claude_env, backend="claude"
    )

    log("checking run_python on the frozen binary")
    results["run_python"] = check_run_python(binary)
    log("checking the agent's tokenizer offline")
    results["tokenizer"] = check_tokenizer(binary)
    expected = 8
    if sys.platform.startswith("linux"):
        log("checking a script's TLS trust store")
        results["tls_trust"] = check_tls_trust(binary)
        expected += 1

    log(f"results: {results}")
    ok = len(results) == expected and all(results.values())
    log("VERDICT: PASS" if ok else "VERDICT: FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
