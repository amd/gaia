# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""Put a native, editable chart on a PowerPoint slide and prove its numbers.

A library function, not a tool: the deck-building skill owns .pptx output and
imports :func:`add_native_chart` from here (or moves it). Its input is exactly
the ``chart_input`` block ``series_from_workbook`` returns — ``categories``,
``series`` ({name: values}) and ``source_note`` — so no number is retyped
between reading the source and drawing the chart.

Stable signature::

    add_native_chart(prs_or_path, slide, chart_type, categories, series,
                     title, number_format, source_note, *,
                     y_axis_title="", x_axis_title="") -> dict
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Union

from gaia.logger import get_logger

log = get_logger(__name__)

PPTX_CHART_TYPES = (
    "line",
    "column",
    "bar",
    "stacked_column",
    "stacked_bar",
    "area",
    "pie",
)


class PptxChartError(ValueError):
    """A slide chart that cannot be written as asked."""


def _xl_type(chart_type: str):
    from pptx.enum.chart import XL_CHART_TYPE

    return {
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
        "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
        "area": XL_CHART_TYPE.AREA,
        "pie": XL_CHART_TYPE.PIE,
    }[chart_type]


def add_native_chart(
    prs_or_path: Union[str, Path, Any],
    slide: Union[int, str, None],
    chart_type: str,
    categories: Sequence[Any],
    series: Dict[str, Sequence[Any]],
    title: str,
    number_format: str,
    source_note: str,
    *,
    y_axis_title: str = "",
    x_axis_title: str = "",
) -> Dict[str, Any]:
    """Add a native chart to a slide, then reread it and compare every value.

    Args:
        prs_or_path: A ``pptx.Presentation`` (the caller saves it) or a .pptx
            path (opened, or created at 16:9 when missing, and saved here).
        slide: 1-based index of an existing slide, or ``"new"``/None for a new
            title-only slide.
        chart_type: One of :data:`PPTX_CHART_TYPES`.
        categories: Period labels, oldest first.
        series: ``{series name: values}`` aligned to *categories*.
        title: Slide title (set when the slide's title is empty) and chart title.
        number_format: Value-axis and data-label format, e.g. ``'#,##0.0'``;
            empty leaves PowerPoint's default.
        source_note: Footnote placed under the chart, e.g. the ``provenance``
            string from ``series_from_workbook``.

    Returns:
        ``{"status": "success", "slide": n, "chart_name", "verified": True,
        "read_back": {name: values}, "categories": [...]}``.

    Raises:
        PptxChartError: bad input, or a saved chart whose values differ.
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    chart_type = (chart_type or "line").lower()
    if chart_type not in PPTX_CHART_TYPES:
        raise PptxChartError(
            f"chart_type {chart_type!r} is not one of {', '.join(PPTX_CHART_TYPES)}."
        )
    if not categories or not series:
        raise PptxChartError(
            "Nothing to chart: categories and series must be non-empty."
        )
    for name, values in series.items():
        if len(values) != len(categories):
            raise PptxChartError(
                f"Series {name!r} has {len(values)} values for {len(categories)} categories."
            )
    if chart_type == "pie" and len(series) != 1:
        raise PptxChartError("A pie chart shows one series; pass exactly one.")

    path: Optional[Path] = None
    if isinstance(prs_or_path, (str, Path)):
        path = Path(prs_or_path)
        if path.suffix.lower() != ".pptx":
            raise PptxChartError(f"{path.name}: charts go into a .pptx file.")
        if path.exists():
            prs = Presentation(str(path))
        else:
            prs = Presentation()
            prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    else:
        prs = prs_or_path

    if slide in (None, "new", ""):
        layout = (
            prs.slide_layouts[5]
            if len(prs.slide_layouts) > 5
            else prs.slide_layouts[-1]
        )
        target = prs.slides.add_slide(layout)
    else:
        index = int(slide)
        if not 1 <= index <= len(prs.slides):
            raise PptxChartError(
                f"slide {index} does not exist; the deck has {len(prs.slides)} slide(s)."
            )
        target = prs.slides[index - 1]
    slide_number = list(prs.slides).index(target) + 1
    if (
        target.shapes.title is not None
        and not target.shapes.title.text.strip()
        and title
    ):
        target.shapes.title.text = title

    data = CategoryChartData()
    data.categories = [str(c) for c in categories]
    for name, values in series.items():
        data.add_series(str(name), list(values), number_format=number_format or None)
    width, height = prs.slide_width, prs.slide_height
    left, top = int(width * 0.06), int(height * 0.2)
    cw, ch = int(width * 0.88), int(height * 0.66)
    frame = target.shapes.add_chart(_xl_type(chart_type), left, top, cw, ch, data)
    chart_name = f"Chart: {title}"[:120] if title else f"Chart {len(target.shapes)}"
    frame.name = chart_name
    chart = frame.chart
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
    chart.has_legend = len(series) > 1 or chart_type == "pie"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if chart_type != "pie":
        value_axis = chart.value_axis
        if number_format:
            value_axis.tick_labels.number_format = number_format
            value_axis.tick_labels.number_format_is_linked = False
        if y_axis_title:
            value_axis.has_title = True
            value_axis.axis_title.text_frame.text = y_axis_title
        if x_axis_title:
            chart.category_axis.has_title = True
            chart.category_axis.axis_title.text_frame.text = x_axis_title
    if source_note:
        box = target.shapes.add_textbox(
            left, top + ch + int(height * 0.02), cw, int(height * 0.06)
        )
        box.name = "Source note"
        box.text_frame.text = source_note
        box.text_frame.word_wrap = True
        for run in box.text_frame.paragraphs[0].runs:
            run.font.size = Pt(12)

    buffer = io.BytesIO()
    prs.save(buffer)
    if path is not None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(buffer.getvalue())
    buffer.seek(0)
    check = read_back_chart(Presentation(buffer), slide_number, chart_name)
    mismatches = _compare(check, categories, series)
    if mismatches:
        raise PptxChartError(
            "The chart does not read back as written: " + "; ".join(mismatches[:5])
        )
    return {
        "status": "success",
        "file": str(path) if path else None,
        "slide": slide_number,
        "chart_name": chart_name,
        "chart_type": chart_type,
        "verified": True,
        "categories": check["categories"],
        "read_back": check["series"],
        "source_note": source_note,
    }


def read_back_chart(prs: Any, slide_number: int, chart_name: str) -> Dict[str, Any]:
    """Categories and series values of the named chart on a slide."""
    slide = prs.slides[slide_number - 1]
    for shape in slide.shapes:
        if shape.has_chart and shape.name == chart_name:
            plot = shape.chart.plots[0]
            return {
                "categories": [str(c) for c in plot.categories],
                "series": {
                    s.name: list(s.values) for p in shape.chart.plots for s in p.series
                },
            }
    return {"categories": [], "series": {}, "missing": True}


def _compare(
    check: Dict[str, Any], categories: Sequence[Any], series: Dict[str, Sequence[Any]]
) -> List[str]:
    if check.get("missing"):
        return ["no chart with that name on the slide after saving"]
    out = []
    if check["categories"] != [str(c) for c in categories]:
        out.append(f"categories {check['categories']} != {list(categories)}")
    for name, values in series.items():
        got = check["series"].get(str(name))
        if got is None:
            out.append(f"series {name!r} missing")
        elif [None if v is None else float(v) for v in got] != [
            None if v is None else float(v) for v in values
        ]:
            out.append(f"{name}: {got} != {list(values)}")
    return out
