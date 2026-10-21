# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""Build a .pptx from a declarative outline with python-pptx.

Every deck is built on a real slide master — the skill's theme generated from
the design tokens, a user template, or (for incremental edits) the deck being
edited. Every slide gets a title placeholder, every chart is native and
editable, every run gets an explicit font and size, and every block of text is
measured to fit its box before it is written. Text that cannot fit at the
minimum size is written anyway and reported (or split onto continuation
slides when the outline allows) — never cut.

Without the GAIA tool wrapper (e.g. from ``run_python``), put the skill folder
on ``sys.path`` and call ``slidedeck.build.build_deck(outline, out_path)``.
"""

from __future__ import annotations

import io
import json
import math
import re
import time
import uuid
import zipfile
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from lxml import etree
from pptx import Presentation
from pptx.chart.axis import ValueAxis
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .outline import OutlineError, normalize_outline, normalize_slide
from .style import (
    DENSITY,
    LAYOUT,
    SIZES,
    PrefsError,
    best_text_on,
    merge_prefs,
    mix,
    resolve_tokens,
    tokens_from_theme_xml,
    validate_prefs,
)
from .text import (
    EMU_PER_PT,
    LINE_SPACING,
    MEASURER,
    FontSpec,
    Para,
    block_height_pt,
    fit_size,
    format_number,
    is_number,
    parse_number,
)
from .theme import make_theme_presentation

# ----------------------------------------------------------------------
# Colours
# ----------------------------------------------------------------------


@dataclass(frozen=True)
class Paint:
    """An RGB colour with an optional tint toward white (or shade, if negative)."""

    rgb: str
    tint: float = 0.0

    def tinted(self, tint: float) -> "Paint":
        return replace(self, tint=tint)

    @property
    def hex(self) -> str:
        return mix(self.rgb, self.tint) if self.tint else self.rgb


def _apply_color(color_format, paint: Paint) -> None:
    color_format.rgb = RGBColor.from_string(paint.hex)


def fill_shape(shape, paint: Optional[Paint]) -> None:
    if paint is None:
        shape.fill.background()
        return
    shape.fill.solid()
    _apply_color(shape.fill.fore_color, paint)


@dataclass
class Theme:
    """Resolved colours and fonts the builder paints with."""

    name: str
    font: str
    heading_font: str
    background: Paint
    title: Paint
    text: Paint
    muted: Paint
    panel: Paint
    line: Paint
    primary: Paint
    highlight: Paint
    context: Paint
    good: Paint
    bad: Paint
    palette: List[Paint]

    def accent(self, i: int) -> Paint:
        return self.palette[i % len(self.palette)]

    def text_on(self, fill: Paint) -> Paint:
        """The readable text colour on *fill* (highest WCAG contrast)."""
        return Paint(
            best_text_on(
                fill.hex, [self.text.hex, self.background.hex, "FFFFFF", "111827"]
            )
        )


def theme_from_tokens(tokens: Dict[str, Any]) -> Theme:
    c = tokens["colors"]

    def P(key: str, default: str) -> Paint:
        return Paint(c.get(key) or default)

    return Theme(
        name=tokens["theme"],
        font=tokens["fonts"]["body"],
        heading_font=tokens["fonts"]["title"],
        background=P("background", "FFFFFF"),
        title=P("title", "1F2937"),
        text=P("text", "1F2937"),
        muted=P("muted", "4B5563"),
        panel=P("panel", "F3F4F6"),
        line=P("line", "D1D5DB"),
        primary=P("primary", "0072B2"),
        highlight=P("highlight", "D55E00"),
        context=P("context", "9CA3AF"),
        good=P("good", "047857"),
        bad=P("bad", "B91C1C"),
        palette=[Paint(h) for h in tokens["palette"]] or [P("primary", "0072B2")],
    )


# ----------------------------------------------------------------------
# Geometry and small helpers
# ----------------------------------------------------------------------


@dataclass
class Box:
    x: int
    y: int
    w: int
    h: int

    @property
    def right(self) -> int:
        return self.x + self.w

    @property
    def bottom(self) -> int:
        return self.y + self.h

    def inset(self, dx: int, dy: Optional[int] = None) -> "Box":
        dy = dx if dy is None else dy
        return Box(
            self.x + dx, self.y + dy, max(1, self.w - 2 * dx), max(1, self.h - 2 * dy)
        )


def pt(emu: float) -> float:
    return emu / EMU_PER_PT


#: Text-frame insets every builder text box uses, so measuring is exact.
INSET_X = Inches(0.08)
INSET_Y = Inches(0.04)
#: Fraction of the measured box text may fill: covers renderer metric drift.
FIT_SLACK = 0.92

CHART_XL = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "column_stacked_100": XL_CHART_TYPE.COLUMN_STACKED_100,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "bar_stacked": XL_CHART_TYPE.BAR_STACKED,
    "bar_stacked_100": XL_CHART_TYPE.BAR_STACKED_100,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "line_plain": XL_CHART_TYPE.LINE,
    "area": XL_CHART_TYPE.AREA,
    "area_stacked": XL_CHART_TYPE.AREA_STACKED,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
    "scatter_lines": XL_CHART_TYPE.XY_SCATTER_LINES,
}

_FAMILY = {
    "column": "column",
    "column_stacked": "column",
    "column_stacked_100": "column",
    "bar": "bar",
    "bar_stacked": "bar",
    "bar_stacked_100": "bar",
    "line": "line",
    "line_plain": "line",
    "area": "area",
    "area_stacked": "area",
}

CHART_LABELS = {
    "column": "Column chart",
    "column_stacked": "Stacked column chart",
    "column_stacked_100": "100% stacked column chart",
    "bar": "Bar chart",
    "bar_stacked": "Stacked bar chart",
    "bar_stacked_100": "100% stacked bar chart",
    "line": "Line chart",
    "line_plain": "Line chart",
    "area": "Area chart",
    "area_stacked": "Stacked area chart",
    "pie": "Pie chart",
    "doughnut": "Doughnut chart",
    "scatter": "Scatter chart",
    "scatter_lines": "Scatter chart",
    "combo": "Combination chart",
}


def _split_bold(text: str) -> List[Tuple[str, bool]]:
    """``"Revenue **up 12%** in Q3"`` -> runs with the marked part bold."""
    parts = text.split("**")
    if len(parts) % 2 == 0:  # unbalanced markers: keep the text literal
        return [(text, False)]
    return [(p, i % 2 == 1) for i, p in enumerate(parts) if p]


def _plain(text: str) -> str:
    return "".join(t for t, _b in _split_bold(text))


def _open_template(path: Path):
    """Open a .pptx/.potx (a template's content type is rewritten in memory)."""
    if path.suffix.lower() not in (".potx", ".potm"):
        return Presentation(str(path))
    out = io.BytesIO()
    with (
        zipfile.ZipFile(str(path)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for item in src.infolist():
            blob = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                blob = re.sub(
                    rb"application/vnd\.(?:openxmlformats-officedocument\.presentationml\.template|ms-powerpoint\.template\.macroEnabled)\.main\+xml",
                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                    blob,
                )
            dst.writestr(item, blob)
    return Presentation(io.BytesIO(out.getvalue()))


def _chart_as_table(chart: Dict[str, Any]) -> Dict[str, Any]:
    """The same numbers as a table, for users who prefer tables."""
    cols = [chart.get("x_title") or ""] + [s["name"] for s in chart["series"]]
    rows = [
        [c] + [s["values"][i] for s in chart["series"]]
        for i, c in enumerate(chart["categories"])
    ]
    fmt = chart.get("number_format")
    return {
        "columns": cols,
        "rows": rows,
        "formats": {
            j: (s.get("number_format") or fmt)
            for j, s in enumerate(chart["series"], 1)
            if (s.get("number_format") or fmt)
        },
        "align": {},
        "total_row": False,
        "highlight_rows": [],
        "col_widths": None,
        "source": chart.get("source", ""),
    }


def _default_number_code(tokens: Dict[str, Any]) -> Optional[str]:
    nf = tokens.get("number_format") or {}
    if not nf:
        return None
    decimals = int(nf.get("decimals", 0) or 0)
    return ("#,##0" if nf.get("thousands", True) else "0") + (
        "." + "0" * decimals if decimals else ""
    )


# ----------------------------------------------------------------------
# The builder
# ----------------------------------------------------------------------


class DeckBuilder:
    """Turns normalized slides into slides on a master."""

    def __init__(
        self,
        outline: Dict[str, Any],
        theme: Optional[str] = None,
        template_path: Optional[str] = None,
        prefs: Optional[Dict[str, Any]] = None,
        *,
        existing_deck: Optional[str] = None,
    ) -> None:
        self.outline = outline
        self.warnings: List[str] = []
        self.notes: List[str] = []
        prefs = dict(prefs or {})
        style = outline.get("style") if isinstance(outline.get("style"), dict) else {}
        problems = validate_prefs(prefs) + validate_prefs(style)
        if problems:
            raise OutlineError(problems)
        request = dict(style)
        if isinstance(outline.get("theme"), str):
            request.setdefault("theme", outline["theme"])
        elif isinstance(outline.get("theme"), dict):
            request = merge_prefs(request, outline["theme"])
        if theme:
            request["theme"] = theme
        if outline.get("footer"):
            request["footer"] = merge_prefs(
                request.get("footer") or {}, {"text": outline["footer"]}
            )
        if outline.get("slide_numbers") is False:
            request["footer"] = merge_prefs(
                request.get("footer") or {}, {"slide_numbers": False}
            )
        saved_template = prefs.pop("template_path", None)
        template = existing_deck or template_path or saved_template
        if existing_deck:
            self.template_source = "the deck being edited"
        elif template_path:
            self.template_source = "template (this request)"
        elif saved_template:
            self.template_source = "template (saved preferences)"
        else:
            self.template_source = "skill theme"
        try:
            if template:
                if not Path(template).is_file():
                    raise OutlineError([f"template {template} does not exist."])
                self.prs = _open_template(Path(template))
                theme_blob = self.prs.slide_master.part.part_related_by(RT.THEME).blob
                template_tokens = tokens_from_theme_xml(theme_blob)
                own = template_tokens.get("theme_name", "").startswith("slide-decks:")
                if existing_deck and own and not request.get("theme"):
                    # Editing one of our decks: keep its theme and its colours.
                    request["theme"] = template_tokens["theme_name"].split(":", 1)[1]
                    request = merge_prefs(
                        request,
                        {
                            "colors": template_tokens["colors"],
                            "fonts": template_tokens["fonts"],
                        },
                    )
                self.tokens = resolve_tokens(
                    request,
                    prefs,
                    template_tokens,
                    template_explicit=bool(template_path or existing_deck),
                )
                if not existing_deck:
                    removed = len(self.prs.slides)
                    sld_ids = (
                        self.prs.slides._sldIdLst
                    )  # noqa: SLF001 - python-pptx has no public delete
                    for sld in list(sld_ids):
                        self.prs.part.drop_rel(sld.rId)
                        sld_ids.remove(sld)
                    if removed:
                        self.notes.append(
                            f"removed {removed} sample slide(s) that came with the template"
                        )
                if request.get("aspect") or prefs.get("aspect"):
                    self.notes.append(
                        "aspect preference not applied: a template keeps its own slide size"
                    )
            else:
                self.tokens = resolve_tokens(request, prefs)
                self.prs = make_theme_presentation(self.tokens)
        except PrefsError as exc:
            raise OutlineError([str(exc)]) from exc
        self.theme = theme_from_tokens(self.tokens)
        self._read_template_styles()
        self.W, self.H = int(self.prs.slide_width), int(self.prs.slide_height)
        density = self.tokens.get("density", "balanced")
        self.body_min = float(DENSITY.get(density, DENSITY["balanced"])["min_body_pt"])
        self.dense_min = float(min(SIZES["dense_min"], self.body_min))
        self.min_pt = float(SIZES["footnote_min"])
        self.font = FontSpec(self.theme.font)
        self.hfont = FontSpec(self.theme.heading_font)
        self.number_code = _default_number_code(self.tokens)
        self.logo = self.tokens.get("logo_path")
        if self.logo and not Path(str(self.logo)).is_file():
            self.warnings.append(
                f"logo {self.logo} does not exist; no logo was placed."
            )
            self.logo = None
        self._layouts = self._pick_layouts()
        self._geometry()
        self.slide_map: List[Tuple[int, List[int]]] = []
        self._current = len(self.prs.slides)
        self._cover = False

    # ------------------------------------------------------------------ setup

    def _read_template_styles(self) -> None:
        self.body_marl: Dict[int, float] = {}
        styles = self.prs.slide_master._element.find(qn("p:txStyles"))  # noqa: SLF001
        if styles is None:
            return
        bs = styles.find(qn("p:bodyStyle"))
        if bs is not None:
            for lvl in range(1, 6):
                ppr = bs.find(qn(f"a:lvl{lvl}pPr"))
                if ppr is not None and ppr.get("marL"):
                    self.body_marl[lvl - 1] = pt(int(ppr.get("marL")))

    def _pick_layouts(self) -> Dict[str, Any]:
        """Map slide kinds onto the master's layouts by placeholder types, then name."""
        layouts = list(self.prs.slide_layouts)
        TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
        BODY_TYPES = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
        OTHER = (
            PP_PLACEHOLDER.PICTURE,
            PP_PLACEHOLDER.CHART,
            PP_PLACEHOLDER.TABLE,
            PP_PLACEHOLDER.MEDIA_CLIP,
            PP_PLACEHOLDER.ORG_CHART,
            PP_PLACEHOLDER.BITMAP,
            PP_PLACEHOLDER.SUBTITLE,
            PP_PLACEHOLDER.VERTICAL_BODY,
            PP_PLACEHOLDER.VERTICAL_OBJECT,
            PP_PLACEHOLDER.VERTICAL_TITLE,
        )

        def kinds(layout):
            return [ph.placeholder_format.type for ph in layout.placeholders]

        def find(pred, prefer=()):
            cands = [lay for lay in layouts if pred(kinds(lay))]
            for word in prefer:
                for lay in cands:
                    if word in lay.name.lower():
                        return lay
            return cands[0] if cands else None

        def has_title(k):
            return any(t in TITLE_TYPES for t in k)

        def n_body(k):
            return sum(1 for t in k if t in BODY_TYPES)

        def clean(k):
            return not any(t in OTHER for t in k)

        title_only = find(
            lambda k: has_title(k)
            and n_body(k) == 0
            and clean(k)
            and PP_PLACEHOLDER.CENTER_TITLE not in k,
            ("title only",),
        )
        content = find(
            lambda k: has_title(k)
            and n_body(k) == 1
            and clean(k)
            and PP_PLACEHOLDER.CENTER_TITLE not in k,
            ("title and content", "content", "text"),
        )
        two = find(
            lambda k: has_title(k) and n_body(k) == 2 and clean(k),
            ("two content", "two"),
        )
        comparison = find(
            lambda k: has_title(k) and n_body(k) == 4 and clean(k), ("comparison",)
        )
        title = find(
            lambda k: PP_PLACEHOLDER.CENTER_TITLE in k, ("title slide", "title")
        ) or find(has_title, ("title slide", "cover"))
        section = find(has_title, ("section",))
        if section is not None and "section" not in section.name.lower():
            section = None
        fallback = title_only or content or find(has_title) or layouts[0]
        if title_only is None:
            self.notes.append(
                f"the template has no 'Title Only' layout; charts and diagrams use '{fallback.name}' with its body placeholder removed"
            )
        if title is None:
            self.notes.append(
                f"the template has no title-slide layout; the cover uses '{fallback.name}'"
            )
        return {
            "title": title or fallback,
            "section": section or title or fallback,
            "content": content or fallback,
            "title_only": title_only or fallback,
            "two": two,
            "comparison": comparison,
        }

    def _geometry(self) -> None:
        W, H = self.W, self.H
        ref = next(
            (
                ph
                for ph in self._layouts["title_only"].placeholders
                if ph.placeholder_format.type
                in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
            ),
            None,
        )
        if ref is not None and ref.width:
            self.title_box = Box(
                int(ref.left), int(ref.top), int(ref.width), int(ref.height)
            )
        else:
            mx = Inches(LAYOUT["margin_x"])
            self.title_box = Box(
                mx,
                Inches(LAYOUT["title_top"]),
                W - 2 * mx,
                Inches(LAYOUT["title_bottom"] - LAYOUT["title_top"]),
            )
        band_top = H - Inches(0.75)
        for ph in self.prs.slide_master.placeholders:
            if (
                ph.placeholder_format.type
                in (
                    PP_PLACEHOLDER.FOOTER,
                    PP_PLACEHOLDER.DATE,
                    PP_PLACEHOLDER.SLIDE_NUMBER,
                )
                and ph.top
            ):
                band_top = min(band_top, int(ph.top))
        self.footer_top = max(band_top, int(H * 0.82))
        self.footer_bottom = min(H - Inches(0.25), self.footer_top + Inches(0.5))
        self.content_x, self.content_w = self.title_box.x, self.title_box.w
        self.unit = Inches(LAYOUT["unit"])
        self.gap = Inches(LAYOUT["gutter"])
        self.content_top = self.title_box.bottom + Inches(0.2)
        self.logo_box = None
        if self.logo:
            from PIL import Image

            with Image.open(self.logo) as im:
                iw, ih = im.size
            lh = min(int(self.title_box.h * 0.6), Inches(0.6))
            lw = min(int(lh * iw / max(1, ih)), Inches(2.0))
            self.logo_box = Box(
                self.title_box.right - lw,
                self.title_box.y + (self.title_box.h - lh) // 2,
                lw,
                lh,
            )

    def grid(self, start: int, span: int, box: Box) -> Tuple[int, int]:
        """x and width of grid columns ``start .. start+span-1`` (12 columns) inside *box*."""
        n, g = LAYOUT["columns"], self.gap
        col = (box.w - g * (n - 1)) / n
        return int(box.x + start * (col + g)), int(col * span + g * (span - 1))

    # ------------------------------------------------------------------ helpers

    def warn(self, msg: str) -> None:
        self.warnings.append(f"slide {self._current}: {msg}")

    def _new_slide(self, kind: str):
        slide = self.prs.slides.add_slide(self._layouts[kind])
        self._current = len(self.prs.slides)
        return slide

    def _title_ph(self, slide):
        for ph in slide.placeholders:
            if ph.placeholder_format.type in (
                PP_PLACEHOLDER.TITLE,
                PP_PLACEHOLDER.CENTER_TITLE,
                PP_PLACEHOLDER.VERTICAL_TITLE,
            ):
                return ph
        return None

    def _ph(self, slide, types) -> List[Any]:
        return [ph for ph in slide.placeholders if ph.placeholder_format.type in types]

    def _delete_last_slide(self) -> None:
        """Remove the slide just added (and its relationship, so no orphan part is saved)."""
        sld_ids = self.prs.slides._sldIdLst  # noqa: SLF001
        last = sld_ids[-1]
        self.prs.part.drop_rel(last.rId)
        sld_ids.remove(last)
        self._current = len(self.prs.slides)

    def _finish(self, slide) -> None:
        """Drop empty placeholders, place the logo, and fix the reading order."""
        for ph in list(slide.placeholders):
            if ph.has_text_frame and not ph.text_frame.text.strip():
                ph._element.getparent().remove(ph._element)  # noqa: SLF001
        if self.logo_box is not None and not self._cover:
            b = self.logo_box
            pic = slide.shapes.add_picture(str(self.logo), b.x, b.y, b.w, b.h)
            pic.name = "Logo"
            pic._element.nvPicPr.cNvPr.set("descr", "Logo")  # noqa: SLF001
        # Screen readers follow z-order: title first, content next, footer last.
        tree = slide.shapes._spTree  # noqa: SLF001
        title = self._title_ph(slide)
        anchor = tree.find(qn("p:grpSpPr"))
        if title is not None and anchor is not None:
            anchor.addnext(title._element)  # noqa: SLF001
        for shp in list(slide.shapes):
            if shp.name in ("Source", "Slide number", "Logo"):
                tree.append(shp._element)  # noqa: SLF001

    @staticmethod
    def _frame(
        tf, *, anchor=MSO_ANCHOR.TOP, wrap=True, inset_x=INSET_X, inset_y=INSET_Y
    ) -> None:
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = inset_x
        tf.margin_top = tf.margin_bottom = inset_y
        tf.vertical_anchor = anchor

    def _indent_pt(self, level: int, placeholder: bool) -> float:
        if level < 0:
            return 0.0
        if placeholder and level in self.body_marl:
            return self.body_marl[level]
        return 20.0 + 23.0 * level

    def paras(
        self,
        items: Sequence[Tuple[str, int, bool]],
        *,
        placeholder: bool = False,
        space: float = 0.4,
    ) -> List[Para]:
        out = []
        for text, level, bold in items:
            scale = 1.0 if level <= 0 else 0.85
            out.append(
                Para(
                    _plain(text),
                    level=level,
                    bold=bold,
                    scale=scale,
                    indent_pt=self._indent_pt(level, placeholder),
                    space_before=space,
                )
            )
        return out

    def fit(
        self,
        paras: Sequence[Para],
        box_w: int,
        box_h: int,
        max_pt: float,
        min_pt: Optional[float] = None,
        *,
        font: Optional[FontSpec] = None,
        inset_x=INSET_X,
        inset_y=INSET_Y,
    ) -> Tuple[float, bool, float, float]:
        """``(size, fits, needed_pt, available_pt)`` for *paras* in a box."""
        min_pt = self.dense_min if min_pt is None else max(min_pt, self.min_pt)
        max_pt = max(max_pt, min_pt)
        w = pt(box_w - 2 * inset_x) * 0.98
        h = pt(box_h - 2 * inset_y) * FIT_SLACK
        size, ok, need = fit_size(
            paras,
            w,
            h,
            font or self.font,
            max_pt,
            min_pt,
            step=1.0,
            floor=min(min_pt, self.dense_min),
        )
        return size, ok, need, h

    def write(
        self,
        shape,
        paras: Sequence[Para],
        size: float,
        *,
        color: Optional[Paint] = None,
        align=None,
        anchor=MSO_ANCHOR.TOP,
        font: Optional[FontSpec] = None,
        placeholder: bool = False,
        colors: Optional[List[Optional[Paint]]] = None,
        inset_x=INSET_X,
        inset_y=INSET_Y,
        raw_texts: Optional[List[str]] = None,
        floor: Optional[float] = None,
    ) -> None:
        """Write *paras* into *shape* at base *size*; every run gets font, size and colour.

        No run goes below *floor* (the dense-text minimum unless a footnote
        passes the footnote minimum) — the same clamp ``fit`` measured with.
        """
        floor = self.dense_min if floor is None else floor
        tf = shape.text_frame
        tf.clear()
        self._frame(tf, anchor=anchor, inset_x=inset_x, inset_y=inset_y)
        f = font or self.font
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if align is not None:
                p.alignment = align
            run_size = max(size * para.scale, floor)
            if i and para.space_before:
                p.space_before = Pt(round(para.space_before * run_size, 1))
            ppr = p._p.get_or_add_pPr()  # noqa: SLF001
            if para.level < 0:
                ppr.set("marL", "0")
                ppr.set("indent", "0")
                etree.SubElement(ppr, qn("a:buNone"))
            elif placeholder:
                p.level = para.level
            else:
                ppr.set("marL", str(int(para.indent_pt * EMU_PER_PT)))
                ppr.set("indent", str(int(-14 * EMU_PER_PT)))
                etree.SubElement(ppr, qn("a:buFont")).set("typeface", "Arial")
                etree.SubElement(ppr, qn("a:buChar")).set(
                    "char", "•" if para.level % 2 == 0 else "–"
                )
            text = raw_texts[i] if raw_texts is not None else para.text
            for chunk, bold in _split_bold(text):
                r = p.add_run()
                r.text = chunk
                r.font.size = Pt(run_size)
                r.font.bold = bool(bold or para.bold)
                r.font.name = (para.font or f).family
                paint = (
                    colors[i]
                    if colors and i < len(colors) and colors[i] is not None
                    else color
                )
                if paint is not None:
                    _apply_color(r.font.color, paint)

    def textbox(self, slide, box: Box, name: str):
        tb = slide.shapes.add_textbox(box.x, box.y, box.w, box.h)
        tb.name = name
        return tb

    def shape(
        self,
        slide,
        kind,
        box: Box,
        name: str,
        fill: Optional[Paint],
        line: Optional[Paint] = None,
        *,
        round_adj: Optional[float] = None,
    ):
        shp = slide.shapes.add_shape(kind, box.x, box.y, box.w, box.h)
        shp.name = name
        fill_shape(shp, fill)
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.width = Pt(1)
            _apply_color(shp.line.color, line)
        shp.shadow.inherit = False
        if round_adj is not None and len(shp.adjustments):
            shp.adjustments[0] = round_adj
        return shp

    def arrow(
        self,
        slide,
        x1,
        y1,
        x2,
        y2,
        name: str,
        paint: Optional[Paint] = None,
        *,
        begin=None,
        end=None,
        elbow: bool = False,
    ):
        c = slide.shapes.add_connector(
            MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT,
            int(x1),
            int(y1),
            int(x2),
            int(y2),
        )
        c.name = name
        if begin is not None:
            c.begin_connect(*begin)
        if end is not None:
            c.end_connect(*end)
        c.line.width = Pt(1.75)
        _apply_color(c.line.color, paint or self.theme.muted)
        ln = c.line._get_or_add_ln()  # noqa: SLF001
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
        return c

    # ------------------------------------------------------------------ chrome

    def set_title(
        self,
        slide,
        text: str,
        box: Optional[Box] = None,
        *,
        max_pt: Optional[float] = None,
        min_pt: Optional[float] = None,
        align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.BOTTOM,
    ) -> None:
        ph = self._title_ph(slide)
        if ph is None:
            ph = self.textbox(slide, box or self.title_box, "Title")
        if box is not None:
            ph.left, ph.top, ph.width, ph.height = box.x, box.y, box.w, box.h
        elif self.logo_box is not None:
            b = self.title_box
            ph.left, ph.top, ph.width, ph.height = (
                b.x,
                b.y,
                b.w - self.logo_box.w - self.gap,
                b.h,
            )
        if not text:
            return
        mx = max_pt or SIZES["title_max"]
        mn = min_pt or SIZES["title_min"]
        hf = FontSpec(self.hfont.family, True)
        para = [Para(_plain(text), level=-1, bold=True, font=hf)]
        size, ok, need, avail = self.fit(
            para, int(ph.width), int(ph.height), mx, mn, font=hf
        )
        if not ok:
            size, ok, need, avail = self.fit(
                para, int(ph.width), int(ph.height), mn, SIZES["title_floor"], font=hf
            )
            self.warn(
                f"title is too long for {mn}pt ({len(text.split())} words), set at {size:.0f}pt. "
                "An action title is one claim in at most ~15 words."
                + ("" if ok else " It still does not fit: shorten it.")
            )
        self.write(
            ph,
            para,
            size,
            color=self.theme.title,
            align=align,
            anchor=anchor,
            font=hf,
            placeholder=True,
            raw_texts=[text],
        )

    def footer(self, slide, s: Dict[str, Any], *, number: bool = True) -> int:
        """Source/footnote line and slide number. Returns the content bottom (EMU)."""
        lines = []
        if s.get("footnote"):
            lines.append(s["footnote"])
        if s.get("source"):
            src = s["source"]
            lines.append(
                src
                if re.match(r"^\s*(sources?|data)\s*:", src, re.I)
                else f"Source: {src}"
            )
        foot = self.tokens.get("footer") or {}
        right_text = str(foot.get("text") or "")
        show_number = number and foot.get("slide_numbers", True) is not False
        fs = float(SIZES["footnote"])
        right_w = 0
        if show_number or right_text:
            right_w = Inches(0.6) + (
                int(MEASURER.width_pt(right_text + "  ", self.font, fs) * EMU_PER_PT)
                if right_text
                else 0
            )
        bottom, top = self.footer_bottom, self.footer_top
        if lines:
            w = self.content_w - right_w - self.gap
            paras = [Para(t, level=-1) for t in lines]
            size = fs
            need = block_height_pt(
                paras, size, pt(w - 2 * INSET_X) * 0.98, self.font
            ) / FIT_SLACK + pt(2 * INSET_Y)
            if need * EMU_PER_PT > bottom - top:
                size = float(SIZES["footnote_min"])
                need = block_height_pt(
                    paras, size, pt(w - 2 * INSET_X) * 0.98, self.font
                ) / FIT_SLACK + pt(2 * INSET_Y)
            h = max(int(need * EMU_PER_PT), bottom - top)
            top = min(top, bottom - h)
            tb = self.textbox(
                slide, Box(self.content_x, top, w, bottom - top), "Source"
            )
            self.write(
                tb,
                paras,
                size,
                color=self.theme.muted,
                anchor=MSO_ANCHOR.BOTTOM,
                floor=self.min_pt,
            )
        if show_number or right_text:
            nb = self.textbox(
                slide,
                Box(
                    self.content_x + self.content_w - right_w,
                    self.footer_top,
                    right_w,
                    bottom - self.footer_top,
                ),
                "Slide number",
            )
            tf = nb.text_frame
            self._frame(tf, anchor=MSO_ANCHOR.BOTTOM)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT
            if right_text:
                r = p.add_run()
                r.text = right_text + "  "
                r.font.size = Pt(fs)
                r.font.name = self.font.family
                _apply_color(r.font.color, self.theme.muted)
            if show_number:
                fld = etree.SubElement(p._p, qn("a:fld"))  # noqa: SLF001
                fld.set("id", "{" + str(uuid.uuid4()).upper() + "}")
                fld.set("type", "slidenum")
                rpr = etree.SubElement(fld, qn("a:rPr"))
                rpr.set("lang", "en-US")
                rpr.set("sz", str(int(fs * 100)))
                sf = etree.SubElement(rpr, qn("a:solidFill"))
                etree.SubElement(sf, qn("a:srgbClr")).set("val", self.theme.muted.hex)
                etree.SubElement(rpr, qn("a:latin")).set("typeface", self.font.family)
                etree.SubElement(fld, qn("a:t")).text = str(self._current)
        return min(top, self.footer_top) - Inches(0.2)

    def notes_for(self, slide, s: Dict[str, Any]) -> None:
        if not s.get("notes"):
            return
        if self.tokens.get("speaker_notes") is False:
            msg = "speaker notes are off (preference): notes in the outline were not written"
            if msg not in self.notes:
                self.notes.append(msg)
            return
        slide.notes_slide.notes_text_frame.text = s["notes"]

    def content_box(self, bottom: int) -> Box:
        return Box(
            self.content_x,
            self.content_top,
            self.content_w,
            max(Inches(1), bottom - self.content_top),
        )

    def _allow_split(self, s: Dict[str, Any]) -> bool:
        if s.get("allow_split") is not None:
            return bool(s["allow_split"])
        return bool(self.outline.get("allow_split"))

    # ------------------------------------------------------------------ build

    def build(self):
        for s in self.outline["slides"]:
            self.build_one(s)
        return self.prs

    def build_one(self, s: Dict[str, Any]) -> List[int]:
        start = len(self.prs.slides) + 1
        self._current = start
        self._cover = s["type"] == "title"
        kind = s["type"]
        if (
            kind == "chart"
            and s.get("data_slide")
            and self.tokens.get("prefer") == "tables"
        ):
            s = dict(s, type="table", table=_chart_as_table(s["chart"]))
            kind = "table"
            self.notes.append(
                f"slide {start}: shown as a table (preference: tables over charts)"
            )
        if s.get("todo"):
            self.warnings.append(
                f"slide {start}: slot '{s.get('slot') or s['title']}' is still a TODO; fill it before handing the deck over."
            )
        getattr(self, f"_slide_{kind}")(s)
        nums = list(range(start, len(self.prs.slides) + 1))
        self.slide_map.append((s["index"], nums))
        if kind in ("chart", "table", "kpi") and not s.get("source"):
            self.warnings.append(
                f"slide {start}: a {kind} slide with no 'source'; say where the numbers come from (file, sheet, period)."
            )
        return nums

    def _standard(self, s, kind="title_only", *, number=True):
        slide = self._new_slide(kind)
        self.set_title(slide, s["title"])
        bottom = self.footer(slide, s, number=number)
        self.notes_for(slide, s)
        return slide, self.content_box(bottom)

    def _side_text(self, slide, items, box: Box, name: str, max_pt: float):
        tb = self.textbox(slide, box, name)
        paras = self.paras(items)
        size, ok, need, avail = self.fit(paras, box.w, box.h, max_pt, self.body_min)
        if not ok:
            size, ok, need, avail = self.fit(
                paras, box.w, box.h, max_pt, self.dense_min
            )
            if not ok:
                self.warn(
                    f"{name.lower()} do not fit beside the content; keep them to 1-3 short lines."
                )
        colors = [self.theme.text if p.level <= 0 else self.theme.muted for p in paras]
        self.write(tb, paras, size, colors=colors, raw_texts=[t for t, _l, _b in items])
        return tb

    # title ------------------------------------------------------------

    def _slide_title(self, s):
        slide = self._new_slide("title")
        H = self.H
        ph = self._title_ph(slide)
        box = (
            Box(int(ph.left), int(ph.top), int(ph.width), int(ph.height))
            if ph is not None and ph.width
            else Box(self.content_x, int(H * 0.3), self.content_w, int(H * 0.22))
        )
        self.set_title(
            slide,
            s["title"] or self.outline.get("title"),
            box,
            max_pt=SIZES["cover_title_max"],
            min_pt=SIZES["title_min"],
        )
        texts = [t for t in (s.get("subtitle"), s.get("meta")) if t]
        sub = self._ph(slide, (PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY))
        if texts:
            ph2 = (
                sub[0]
                if sub
                else self.textbox(
                    slide,
                    Box(box.x, box.bottom + self.unit // 2, box.w, int(H * 0.18)),
                    "Subtitle",
                )
            )
            paras = [
                Para(t, level=-1, scale=1.0 if i == 0 else 0.8)
                for i, t in enumerate(texts)
            ]
            size, ok, need, avail = self.fit(
                paras, int(ph2.width), int(ph2.height), SIZES["subtitle"], self.body_min
            )
            if not ok:
                self.warn("subtitle does not fit; shorten it.")
            colors = [self.theme.text] + [self.theme.muted] * (len(paras) - 1)
            self.write(
                ph2, paras, size, colors=colors, align=PP_ALIGN.LEFT, placeholder=True
            )
        if self.logo:
            from PIL import Image

            with Image.open(self.logo) as im:
                iw, ih = im.size
            lh = Inches(0.8)
            lw = min(int(lh * iw / max(1, ih)), Inches(2.5))
            pic = slide.shapes.add_picture(
                str(self.logo), self.content_x, self.footer_top - lh, lw, lh
            )
            pic.name = "Logo"
            pic._element.nvPicPr.cNvPr.set("descr", "Logo")  # noqa: SLF001
        if s.get("source") or s.get("footnote"):
            self.footer(slide, s, number=False)
        self.notes_for(slide, s)
        self._finish(slide)

    # section ------------------------------------------------------------

    def _slide_section(self, s):
        slide = self._new_slide("section")
        title = f"{s['number']}  {s['title']}" if s.get("number") else s["title"]
        ph = self._title_ph(slide)
        box = (
            Box(int(ph.left), int(ph.top), int(ph.width), int(ph.height))
            if ph is not None and ph.width
            else None
        )
        self.set_title(
            slide, title, box, max_pt=SIZES["title_max"] + 4, min_pt=SIZES["title_min"]
        )
        body = self._ph(
            slide, (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.OBJECT)
        )
        if s.get("subtitle"):
            target = (
                body[0]
                if body
                else self.textbox(
                    slide,
                    Box(
                        self.content_x,
                        (box.bottom if box else self.content_top) + self.unit,
                        self.content_w,
                        Inches(1),
                    ),
                    "Subtitle",
                )
            )
            paras = [Para(s["subtitle"], level=-1)]
            size, ok, *_ = self.fit(
                paras,
                int(target.width),
                int(target.height),
                SIZES["subtitle"],
                self.body_min,
            )
            if not ok:
                self.warn("section subtitle does not fit; shorten it.")
            self.write(target, paras, size, color=self.theme.muted, placeholder=True)
        self.footer(slide, s, number=True)
        self.notes_for(slide, s)
        self._finish(slide)

    # agenda ------------------------------------------------------------

    def _slide_agenda(self, s):
        slide, box = self._standard(s)
        items = s["items"]
        hl = s.get("highlight")
        paras = [
            Para(f"{i + 1}.  {_plain(t)}", level=-1, bold=(hl == i), space_before=0.6)
            for i, t in enumerate(items)
        ]
        tb = self.textbox(slide, box, "Agenda")
        size, ok, need, avail = self.fit(
            paras, box.w, box.h, SIZES["body_max"] + 4, self.body_min
        )
        if not ok:
            self.warn(
                f"agenda needs {need:.0f}pt of {avail:.0f}pt at {size:.0f}pt; use fewer or shorter items."
            )
        colors = [
            self.theme.primary if hl == i else self.theme.text
            for i in range(len(items))
        ]
        self.write(tb, paras, size, colors=colors)
        self._finish(slide)

    # bullets ------------------------------------------------------------

    def _slide_bullets(self, s, paras_in=None, title=None):
        items = paras_in if paras_in is not None else s["paras"]
        use_ph = self._layouts["content"] is not self._layouts["title_only"]
        slide = self._new_slide("content" if use_ph else "title_only")
        self.set_title(slide, title if title is not None else s["title"])
        bottom = self.footer(slide, s)
        self.notes_for(slide, s)
        body = (
            self._ph(slide, (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT))
            if use_ph
            else []
        )
        if body:
            shp = body[0]
            if int(shp.top) + int(shp.height) > bottom:
                shp.height = max(Inches(1), bottom - int(shp.top))
            placeholder = True
        else:
            shp = self.textbox(slide, self.content_box(bottom), "Body")
            placeholder = False
        paras = self.paras(items, placeholder=placeholder)
        size, ok, need, avail = self.fit(
            paras, int(shp.width), int(shp.height), SIZES["body_max"], self.body_min
        )
        if not ok and self._allow_split(s) and len(items) > 1 and paras_in is None:
            chunks = self._split_to_fit(
                items, int(shp.width), int(shp.height), placeholder
            )
            if len(chunks) > 1:
                self._delete_last_slide()
                for k, chunk in enumerate(chunks):
                    self._slide_bullets(
                        s, chunk, s["title"] if k == 0 else f"{s['title']} (cont.)"
                    )
                self.notes.append(
                    f"outline slide {s['index']}: split into {len(chunks)} slides to keep body text at {self.body_min:.0f}pt or larger"
                )
                return
        if not ok:
            self.warn(
                f"body text needs {need:.0f}pt of height at the minimum {size:.0f}pt but the box has {avail:.0f}pt; it will overflow. Cut words, split the slide (allow_split), or move extra detail (never requested content) to the notes."
            )
        colors = [self.theme.text if p.level <= 0 else self.theme.muted for p in paras]
        self.write(
            shp,
            paras,
            size,
            colors=colors,
            placeholder=placeholder,
            raw_texts=[t for t, _l, _b in items],
        )
        self._finish(slide)

    def _split_to_fit(
        self, items, w, h, placeholder
    ) -> List[List[Tuple[str, int, bool]]]:
        chunks: List[List[Tuple[str, int, bool]]] = []
        cur: List[Tuple[str, int, bool]] = []
        for it in items:
            trial = cur + [it]
            _size, ok, _n, _a = self.fit(
                self.paras(trial, placeholder=placeholder),
                w,
                h,
                self.body_min + 2,
                self.body_min + 2,
            )
            if ok or not cur:
                cur = trial
            else:
                chunks.append(cur)
                cur = [it]
        if cur:
            chunks.append(cur)
        return chunks

    # two column ------------------------------------------------------------

    def _slide_two_column(self, s):
        cols = s["columns"]
        headed = all(c["heading"] for c in cols)
        kind = (
            "comparison"
            if headed and self._layouts.get("comparison") is not None
            else ("two" if self._layouts.get("two") is not None else "title_only")
        )
        slide = self._new_slide(kind)
        self.set_title(slide, s["title"])
        bottom = self.footer(slide, s)
        self.notes_for(slide, s)
        heads: List[Any] = []
        bodies: List[Any] = []
        if kind == "comparison":
            phs = self._ph(slide, (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT))
            heads = sorted(
                [p for p in phs if p.placeholder_format.type == PP_PLACEHOLDER.BODY],
                key=lambda p: p.left,
            )[:2]
            bodies = sorted(
                [p for p in phs if p.placeholder_format.type == PP_PLACEHOLDER.OBJECT],
                key=lambda p: p.left,
            )[:2]
            if len(heads) < 2 or len(bodies) < 2:
                heads, bodies = [], sorted(phs, key=lambda p: (p.left, p.top))[:2]
        elif kind == "two":
            bodies = sorted(
                self._ph(slide, (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)),
                key=lambda p: p.left,
            )[:2]
        for ph in bodies:
            if int(ph.top) + int(ph.height) > bottom:
                ph.height = max(Inches(1), bottom - int(ph.top))
        placeholder = len(bodies) == 2
        if not placeholder:
            box = self.content_box(bottom)
            lx, lw = self.grid(0, 6, box)
            rx, rw = self.grid(6, 6, box)
            bodies = [
                self.textbox(slide, Box(lx, box.y, lw, box.h), "Left column"),
                self.textbox(slide, Box(rx, box.y, rw, box.h), "Right column"),
            ]
        all_paras, all_raw = [], []
        for col in cols:
            ps, raws = [], []
            if col["heading"] and not heads:
                ps.append(
                    Para(
                        _plain(col["heading"]),
                        level=-1,
                        bold=True,
                        scale=1.0,
                        font=FontSpec(self.hfont.family, True),
                    )
                )
                raws.append(col["heading"])
            ps += self.paras(col["paras"], placeholder=placeholder)
            raws += [t for t, _l, _b in col["paras"]]
            all_paras.append(ps)
            all_raw.append(raws)
        fits = [
            self.fit(
                ps,
                int(shp.width),
                int(shp.height),
                SIZES["body_max"] - 2,
                self.body_min,
            )
            for shp, ps in zip(bodies, all_paras)
        ]
        size = min(f[0] for f in fits)
        for f, side in zip(fits, ("left", "right")):
            if not f[1]:
                self.warn(
                    f"{side} column needs {f[2]:.0f}pt of {f[3]:.0f}pt at {f[0]:.0f}pt; shorten it."
                )
        for shp, ps, raws in zip(bodies, all_paras, all_raw):
            colors = [
                (
                    self.theme.primary
                    if (p.bold and p.level < 0)
                    else (self.theme.text if p.level <= 0 else self.theme.muted)
                )
                for p in ps
            ]
            self.write(
                shp, ps, size, colors=colors, placeholder=placeholder, raw_texts=raws
            )
        for ph, col in zip(heads, cols):
            hp = [
                Para(
                    _plain(col["heading"]),
                    level=-1,
                    bold=True,
                    font=FontSpec(self.hfont.family, True),
                )
            ]
            hs, ok, *_ = self.fit(
                hp, int(ph.width), int(ph.height), SIZES["body_max"], self.body_min
            )
            if not ok:
                self.warn(f"column heading {col['heading']!r} is too long; shorten it.")
            self.write(
                ph,
                hp,
                hs,
                color=self.theme.primary,
                placeholder=True,
                anchor=MSO_ANCHOR.BOTTOM,
                raw_texts=[col["heading"]],
            )
        self._finish(slide)

    # chart ------------------------------------------------------------

    def _slide_chart(self, s):
        slide, box = self._standard(s)
        spec = s["chart"]
        if s.get("paras"):
            cx, cw = self.grid(0, 8, box)
            tx, tw = self.grid(8, 4, box)
            self._side_text(
                slide,
                s["paras"],
                Box(tx, box.y, tw, box.h),
                "Takeaways",
                SIZES["body_max"] - 2,
            )
            cbox = Box(cx, box.y, cw, box.h)
        else:
            cbox = box
        self.add_chart(slide, spec, cbox, s["title"])
        self._finish(slide)

    def add_chart(self, slide, spec: Dict[str, Any], box: Box, title: str):
        ctype = spec["chart_type"]
        series = spec["series"]
        scatter = ctype in ("scatter", "scatter_lines")
        base = ctype
        if ctype == "combo":
            first = next((x["type"] for x in series if x["type"]), "column")
            base = (
                "column"
                if any((x["type"] or "column") == "column" for x in series)
                else first
            )
        fmt = spec.get("number_format") or self.number_code
        if scatter:
            cd = XyChartData()
            for ser in series:
                xs = cd.add_series(
                    ser["name"], number_format=ser.get("number_format") or fmt
                )
                for x, y in ser["points"]:
                    xs.add_data_point(x, y)
        else:
            cd = CategoryChartData(number_format=fmt or "General")
            cd.categories = spec["categories"]
            for ser in series:
                cd.add_series(
                    ser["name"],
                    ser["values"],
                    number_format=ser.get("number_format") or fmt,
                )
        gf = slide.shapes.add_chart(CHART_XL[base], box.x, box.y, box.w, box.h, cd)
        gf.name = f"Chart: {title[:40]}" if title else "Chart"
        gf._element.nvGraphicFramePr.cNvPr.set(
            "descr", self._chart_alt(spec, title)
        )  # noqa: SLF001
        chart = gf.chart
        family = _FAMILY.get(base)
        secondary_axis = None
        if family and not scatter:
            secondary_axis = self._overlay_groups(chart, base, series)
        self._declutter(chart)
        small = float(SIZES["dense_min"])
        chart.font.size = Pt(small)
        chart.font.name = self.font.family
        chart.font.color.rgb = RGBColor.from_string(self.theme.text.hex)
        if spec.get("chart_title"):
            chart.has_title = True
            chart.chart_title.text_frame.text = spec["chart_title"]
            for p in chart.chart_title.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(16)
        else:
            chart.has_title = False
        tk = self.tokens.get("chart") or {}
        legend = spec.get("legend", "auto")
        if legend == "auto":
            legend = tk.get("legend", "auto")
        show_legend = legend not in ("none", False) and (
            legend != "auto" or len(series) > 1 or ctype in ("pie", "doughnut")
        )
        chart.has_legend = show_legend
        if show_legend:
            positions = {
                "bottom": XL_LEGEND_POSITION.BOTTOM,
                "right": XL_LEGEND_POSITION.RIGHT,
                "top": XL_LEGEND_POSITION.TOP,
                "left": XL_LEGEND_POSITION.LEFT,
            }
            chart.legend.position = positions.get(
                legend,
                (
                    XL_LEGEND_POSITION.RIGHT
                    if ctype in ("pie", "doughnut")
                    else XL_LEGEND_POSITION.BOTTOM
                ),
            )
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(small)
        n_points = sum(len(x["values"]) for x in series)
        dl = spec.get("data_labels", "auto")
        if dl == "auto":
            dl = tk.get("data_labels", "auto")
        if dl == "auto":
            if ctype in ("pie", "doughnut"):
                dl = True
            elif scatter:
                dl = False
            elif len(chart.plots) > 1:
                dl = False  # labels from two chart groups collide
            elif family == "line":
                dl = len(series) <= 2 and n_points <= 16
            else:
                dl = n_points <= 24
        self._color_series(chart, ctype, spec, bool(dl))
        for plot in chart.plots:
            if not dl:
                if (
                    plot.__class__.__name__ != "XyPlot"
                ):  # python-pptx cannot clear XY labels
                    plot.has_data_labels = False
                continue
            kind = plot.__class__.__name__
            names = [x.name for x in plot.series]
            specs = [x for x in series if x["name"] in names]
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.font.size = Pt(small)
            if ctype in ("pie", "doughnut") and (
                dl == "percent" or not spec.get("number_format")
            ):
                labels.show_percentage = True
                labels.show_value = False
                labels.number_format = "0%"
                labels.number_format_is_linked = False
            else:
                label_fmt = next(
                    (x.get("number_format") for x in specs if x.get("number_format")),
                    None,
                )
                if (
                    label_fmt is None
                    and specs
                    and all(x["axis"] == "secondary" for x in specs)
                ):
                    label_fmt = spec.get("y2_number_format")
                label_fmt = label_fmt or fmt
                if label_fmt:
                    labels.number_format = label_fmt
                    labels.number_format_is_linked = False
                labels.show_value = True
            stacked = False
            try:
                if kind == "BarPlot":
                    grouping = plot._element.find(qn("c:grouping"))  # noqa: SLF001
                    stacked = grouping is not None and grouping.get("val") in (
                        "stacked",
                        "percentStacked",
                    )
                    labels.position = (
                        XL_DATA_LABEL_POSITION.CENTER
                        if stacked
                        else XL_DATA_LABEL_POSITION.OUTSIDE_END
                    )
                elif kind == "LinePlot":
                    labels.position = XL_DATA_LABEL_POSITION.ABOVE
                elif kind == "PiePlot":
                    labels.position = XL_DATA_LABEL_POSITION.BEST_FIT
            except (ValueError, AttributeError):
                pass
            if stacked:
                labels.font.color.rgb = RGBColor.from_string("FFFFFF")
        if ctype in ("pie", "doughnut"):
            if len(spec["categories"]) > 5:
                self.warn(
                    f"a {ctype} with {len(spec['categories'])} slices is hard to read; use a sorted bar chart (pie only for up to 4-5 parts of a whole)."
                )
            return gf
        gridlines = (
            spec.get("gridlines", True) is not False
            and tk.get("gridlines", True) is not False
        )
        try:
            cat_ax = chart.category_axis
            cat_ax.tick_labels.font.size = Pt(small)
            cat_ax.has_major_gridlines = False
            cat_ax.format.line.color.rgb = RGBColor.from_string(self.theme.line.hex)
            if spec.get("x_title"):
                cat_ax.has_title = True
                cat_ax.axis_title.text_frame.text = spec["x_title"]
                self._axis_title_size(cat_ax)
        except (ValueError, AttributeError):
            pass
        val_ax = chart.value_axis
        val_ax.tick_labels.font.size = Pt(small)
        if fmt:
            val_ax.tick_labels.number_format = fmt
            val_ax.tick_labels.number_format_is_linked = False
        val_ax.has_major_gridlines = gridlines
        if gridlines:
            val_ax.major_gridlines.format.line.color.rgb = RGBColor.from_string(
                mix(self.theme.line.hex, 0.35)
                if self.theme.background.hex == "FFFFFF"
                else self.theme.line.hex
            )
            val_ax.major_gridlines.format.line.width = Pt(0.75)
        val_ax.format.line.fill.background()
        if family in ("column", "bar") and spec.get("y_min") not in (None, 0):
            self.warn(
                "bars must start at zero; the y_min on this bar/column chart exaggerates differences. Use a line chart to show a narrow range."
            )
        if spec.get("y_min") is not None:
            val_ax.minimum_scale = float(spec["y_min"])
        elif family in ("column", "bar", "area"):
            val_ax.minimum_scale = min(
                0.0,
                min(
                    (v for x in series for v in x["values"] if v is not None),
                    default=0.0,
                ),
            )
        if spec.get("y_max") is not None:
            val_ax.maximum_scale = float(spec["y_max"])
        if spec.get("y_title"):
            val_ax.has_title = True
            val_ax.axis_title.text_frame.text = spec["y_title"]
            self._axis_title_size(val_ax)
        if secondary_axis is not None:
            sec_fmt = spec.get("y2_number_format") or next(
                (
                    x.get("number_format")
                    for x in series
                    if x["axis"] == "secondary" and x.get("number_format")
                ),
                None,
            )
            secondary_axis.tick_labels.font.size = Pt(small)
            if sec_fmt:
                secondary_axis.tick_labels.number_format = sec_fmt
                secondary_axis.tick_labels.number_format_is_linked = False
            if spec.get("y2_title"):
                secondary_axis.has_title = True
                secondary_axis.axis_title.text_frame.text = spec["y2_title"]
                self._axis_title_size(secondary_axis)
        if len(series) > 5:
            self.warn(
                f"{len(series)} series on one chart; keep to 5 or fewer (show the rest in grey, or split the chart)."
            )
        return gf

    @staticmethod
    def _chart_alt(spec: Dict[str, Any], title: str) -> str:
        kind = CHART_LABELS.get(spec["chart_type"], "Chart")
        names = ", ".join(s["name"] for s in spec["series"][:5])
        cats = spec.get("categories") or []
        span = f" over {len(cats)} categories ({cats[0]} to {cats[-1]})" if cats else ""
        text = f"{kind}: {names}{span}."
        if title:
            text = f"{title}. {text}"
        return text[:400]

    @staticmethod
    def _declutter(chart) -> None:
        """No border on the chart area and no fill behind it."""
        cs = chart._chartSpace  # noqa: SLF001
        sppr = cs.find(qn("c:spPr"))
        if sppr is None:
            sppr = etree.Element(qn("c:spPr"))
            anchor = cs.find(qn("c:txPr"))
            if anchor is None:
                anchor = cs.find(qn("c:externalData"))
            if anchor is not None:
                anchor.addprevious(sppr)
            else:
                cs.append(sppr)
        for child in list(sppr):
            sppr.remove(child)
        etree.SubElement(sppr, qn("a:noFill"))
        ln = etree.SubElement(sppr, qn("a:ln"))
        etree.SubElement(ln, qn("a:noFill"))

    def _axis_title_size(self, axis) -> None:
        for p in axis.axis_title.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(SIZES["dense_min"])
                r.font.bold = False

    def _overlay_groups(
        self, chart, base: str, series: List[Dict[str, Any]]
    ) -> Optional[ValueAxis]:
        """Move series whose type or axis differ from the base into their own chart groups.

        python-pptx builds one chart group; a combo (bars + line) or a
        secondary value axis is a second group in the same plot area, which
        PowerPoint renders and edits natively.
        """
        family = _FAMILY[base]
        groups: Dict[Tuple[str, str], List[int]] = {}
        for i, ser in enumerate(series):
            t = ser["type"] or family
            if family == "bar" and t in ("bar", "column"):
                t = "bar"
            if t == family and ser["axis"] == "primary":
                continue
            groups.setdefault((t, ser["axis"]), []).append(i)
        if not groups:
            return None
        if family == "bar":
            raise OutlineError(
                [
                    f"slide {self._current}: per-series types and secondary axes need a column, line or area base chart, not horizontal bars."
                ]
            )
        plot_area = chart._chartSpace.chart.plotArea  # noqa: SLF001
        base_el = plot_area.xpath("c:barChart|c:lineChart|c:areaChart")[0]
        base_ax = [a.get("val") for a in base_el.findall(qn("c:axId"))]
        sers = base_el.findall(qn("c:ser"))
        C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
        secondary_ids = None
        new_axes = []
        for (t, axis), idxs in groups.items():
            if axis == "secondary":
                if secondary_ids is None:
                    secondary_ids = ("50000001", "50000101")
                    cat_id, val_id = secondary_ids
                    new_axes = [
                        parse_xml(
                            f'<c:catAx xmlns:c="{C}"><c:axId val="{cat_id}"/><c:scaling><c:orientation val="minMax"/></c:scaling>'
                            f'<c:delete val="1"/><c:axPos val="b"/><c:majorTickMark val="none"/><c:minorTickMark val="none"/>'
                            f'<c:tickLblPos val="nextTo"/><c:crossAx val="{val_id}"/><c:crosses val="autoZero"/><c:auto val="1"/>'
                            f'<c:lblAlgn val="ctr"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>'
                        ),
                        parse_xml(
                            f'<c:valAx xmlns:c="{C}"><c:axId val="{val_id}"/><c:scaling><c:orientation val="minMax"/></c:scaling>'
                            f'<c:delete val="0"/><c:axPos val="r"/><c:numFmt formatCode="General" sourceLinked="1"/>'
                            f'<c:majorTickMark val="out"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
                            f'<c:crossAx val="{cat_id}"/><c:crosses val="max"/><c:crossBetween val="between"/></c:valAx>'
                        ),
                    ]
                ax_ids = secondary_ids
            else:
                ax_ids = tuple(base_ax)
            if t == "line":
                grp = parse_xml(
                    f'<c:lineChart xmlns:c="{C}"><c:grouping val="standard"/><c:varyColors val="0"/></c:lineChart>'
                )
            elif t == "area":
                grp = parse_xml(
                    f'<c:areaChart xmlns:c="{C}"><c:grouping val="standard"/><c:varyColors val="0"/></c:areaChart>'
                )
            else:
                grp = parse_xml(
                    f'<c:barChart xmlns:c="{C}"><c:barDir val="col"/><c:grouping val="clustered"/><c:varyColors val="0"/></c:barChart>'
                )
            for i in idxs:
                ser = sers[i]
                base_el.remove(ser)
                self._convert_ser(ser, t)
                grp.append(ser)
            if t == "line":
                etree.SubElement(grp, qn("c:marker")).set("val", "1")
            elif t == "column":
                etree.SubElement(grp, qn("c:gapWidth")).set("val", "150")
            for a in ax_ids:
                etree.SubElement(grp, qn("c:axId")).set("val", a)
            if t == "line" or (t == "area" and family == "column"):
                base_el.addnext(grp)
            else:
                base_el.addprevious(grp)
        if not base_el.findall(qn("c:ser")):
            raise OutlineError(
                [
                    f"slide {self._current}: every series was moved off the base chart; give at least one series the base type on the primary axis."
                ]
            )
        if new_axes:
            last_ax = plot_area.xpath("c:catAx|c:valAx|c:dateAx|c:serAx")[-1]
            for ax in reversed(new_axes):
                last_ax.addnext(ax)
            return ValueAxis(new_axes[1])
        return None

    @staticmethod
    def _convert_ser(ser, target: str) -> None:
        for tag in (
            "c:invertIfNegative",
            "c:shape",
            "c:marker",
            "c:smooth",
            "c:explosion",
            "c:pictureOptions",
        ):
            for el in ser.findall(qn(tag)):
                ser.remove(el)
        anchor = ser.find(qn("c:spPr"))
        if anchor is None:
            anchor = ser.find(qn("c:tx"))
        if anchor is None:
            anchor = ser.find(qn("c:order"))
        if target == "line":
            marker = etree.Element(qn("c:marker"))
            etree.SubElement(marker, qn("c:symbol")).set("val", "circle")
            etree.SubElement(marker, qn("c:size")).set("val", "7")
            anchor.addnext(marker)
            etree.SubElement(ser, qn("c:smooth")).set("val", "0")
        elif target == "column":
            inv = etree.Element(qn("c:invertIfNegative"))
            inv.set("val", "0")
            anchor.addnext(inv)

    def _series_paints(self, spec: Dict[str, Any], labelled: bool) -> Dict[str, Paint]:
        """Colour per series: the message series in the accent, the rest grey.

        Without a highlight, the palette in order; the two low-contrast
        Okabe-Ito hues (orange, sky blue) only when values are labelled.
        """
        series = spec["series"]
        palette = self.theme.palette if labelled else self.theme.palette[:4]
        hl = spec.get("highlight")
        names = [s["name"] for s in series]
        out: Dict[str, Paint] = {}
        if hl is not None and str(hl) in names and len(series) > 1:
            for k, n in enumerate(names):
                out[n] = (
                    self.theme.primary
                    if n == str(hl)
                    else self.theme.context.tinted(min(0.5, 0.15 * k))
                )
            return out
        for k, s in enumerate(series):
            if s.get("color") and re.fullmatch(r"#?[0-9A-Fa-f]{6}", str(s["color"])):
                out[s["name"]] = Paint(str(s["color"]).lstrip("#").upper())
            elif k < len(palette):
                out[s["name"]] = palette[k]
            else:
                out[s["name"]] = self.theme.context.tinted(
                    min(0.6, 0.15 * (k - len(palette)))
                )
        return out

    def _color_series(
        self, chart, ctype: str, spec: Dict[str, Any], labelled: bool
    ) -> None:
        from pptx.enum.chart import XL_MARKER_STYLE

        paints = self._series_paints(spec, labelled)
        cats = spec.get("categories") or []
        hl = spec.get("highlight")
        hl_point = cats.index(str(hl)) if hl is not None and str(hl) in cats else None
        if hl is not None and hl_point is None and str(hl) not in paints:
            self.warn(
                f"chart highlight {hl!r} is neither a series name nor a category."
            )
        for plot in chart.plots:
            kind = plot.__class__.__name__
            for ser in plot.series:
                paint = paints.get(ser.name, self.theme.primary)
                if ctype in ("pie", "doughnut"):
                    pal = self.theme.palette if labelled else self.theme.palette[:4]
                    for j, point in enumerate(ser.points):
                        point.format.fill.solid()
                        if hl_point is not None:
                            p = (
                                self.theme.highlight
                                if hl_point == j
                                else self.theme.context.tinted(min(0.6, 0.12 * j))
                            )
                        else:
                            p = (
                                pal[j]
                                if j < len(pal)
                                else self.theme.context.tinted(
                                    min(0.6, 0.2 * (j - len(pal)))
                                )
                            )
                        _apply_color(point.format.fill.fore_color, p)
                        point.format.line.color.rgb = RGBColor.from_string(
                            self.theme.background.hex
                        )
                    continue
                if kind in ("LinePlot", "XyPlot"):
                    ser.format.line.width = Pt(2.5)
                    _apply_color(ser.format.line.color, paint)
                    ser.marker.style = XL_MARKER_STYLE.CIRCLE
                    ser.marker.size = 7
                    ser.marker.format.fill.solid()
                    _apply_color(ser.marker.format.fill.fore_color, paint)
                    ser.marker.format.line.fill.background()
                    if kind == "XyPlot" and ctype == "scatter":
                        ser.format.line.fill.background()
                    if hl_point is not None and kind == "LinePlot":
                        point = ser.points[hl_point]
                        point.marker.style = XL_MARKER_STYLE.CIRCLE
                        point.marker.size = 11
                        point.marker.format.fill.solid()
                        _apply_color(
                            point.marker.format.fill.fore_color, self.theme.highlight
                        )
                else:
                    ser.format.fill.solid()
                    _apply_color(ser.format.fill.fore_color, paint)
                    if hl_point is not None and len(spec["series"]) == 1:
                        for j, point in enumerate(ser.points):
                            point.format.fill.solid()
                            _apply_color(
                                point.format.fill.fore_color,
                                (
                                    self.theme.highlight
                                    if j == hl_point
                                    else self.theme.primary
                                ),
                            )

    # table ------------------------------------------------------------

    def _cell_text(self, value, fmt: Optional[str]) -> str:
        if value is None:
            return ""
        if is_number(value):
            return format_number(value, fmt or self.number_code)
        if fmt and isinstance(value, str):
            n = parse_number(value)
            if n is not None and re.fullmatch(r"\s*[-+]?[\d,]*\.?\d+\s*", value):
                return format_number(n, fmt)
        return str(value)

    def _slide_table(self, s, rows_override=None, title=None):
        t = s["table"]
        rows_all = t["rows"] if rows_override is None else rows_override
        slide = self._new_slide("title_only")
        self.set_title(slide, title if title is not None else s["title"])
        bottom = self.footer(slide, s)
        self.notes_for(slide, s)
        box = self.content_box(bottom)
        header = t["columns"]
        ncols = len(header)
        texts = [
            [self._cell_text(v, t["formats"].get(j)) for j, v in enumerate(r)]
            for r in rows_all
        ]
        numeric_cols = set()
        for j in range(ncols):
            vals = [r[j] for r in rows_all if r[j] not in (None, "")]
            nums = [
                v
                for v in vals
                if is_number(v) or (isinstance(v, str) and parse_number(v) is not None)
            ]
            if vals and len(nums) / len(vals) >= 0.6 and j > 0:
                numeric_cols.add(j)
        has_header = any(h for h in header)
        grid = ([header] if has_header else []) + texts
        bold_font = FontSpec(self.font.family, True)
        cell_mx, cell_my = Inches(0.1), Inches(0.05)

        def layout(size: float):
            natural = []
            for j in range(ncols):
                w = 0.0
                for i, row in enumerate(grid):
                    f = (
                        bold_font
                        if (has_header and i == 0)
                        or (t["total_row"] and i == len(grid) - 1)
                        else self.font
                    )
                    for line in str(row[j]).split("\n"):
                        w = max(w, MEASURER.width_pt(line, f, size))
                natural.append(w + pt(2 * cell_mx) + 4)
            total_w = pt(box.w)
            if t.get("col_widths") and len(t["col_widths"]) == ncols:
                rel = [max(0.01, float(x)) for x in t["col_widths"]]
                widths = [total_w * r / sum(rel) for r in rel]
            elif sum(natural) <= total_w:
                extra = total_w - sum(natural)
                widths = [n + extra * n / sum(natural) for n in natural]
            else:
                # Water-fill: narrow columns keep their width, the rest share and wrap.
                widths = [0.0] * ncols
                remaining = list(range(ncols))
                budget = total_w
                while remaining:
                    share = budget / len(remaining)
                    fixed = [j for j in remaining if natural[j] <= share]
                    if not fixed:
                        for j in remaining:
                            widths[j] = share
                        break
                    for j in fixed:
                        widths[j] = natural[j]
                        budget -= natural[j]
                        remaining.remove(j)
            heights = []
            for i, row in enumerate(grid):
                f = bold_font if (has_header and i == 0) else self.font
                lines = max(
                    (
                        MEASURER.line_count(
                            str(row[j]), f, size, widths[j] - pt(2 * cell_mx) - 2
                        )
                        for j in range(ncols)
                    ),
                    default=1,
                )
                heights.append(lines * size * LINE_SPACING + pt(2 * cell_my) + 2)
            return widths, heights

        size = float(SIZES["body_min"])
        while True:
            widths, heights = layout(size)
            if sum(heights) <= pt(box.h) * 0.97 or size <= self.dense_min:
                break
            size -= 1
        fits = sum(heights) <= pt(box.h) * 0.97
        if (
            not fits
            and rows_override is None
            and self._allow_split(s)
            and len(rows_all) > 2
        ):
            body_h = sum(heights) - (heights[0] if has_header else 0)
            avail = pt(box.h) * 0.97 - (heights[0] if has_header else 0)
            per = max(1, int(len(rows_all) * avail / max(1.0, body_h)))
            self._delete_last_slide()
            chunks = [rows_all[i : i + per] for i in range(0, len(rows_all), per)]
            for k, chunk in enumerate(chunks):
                self._slide_table(
                    s, chunk, s["title"] if k == 0 else f"{s['title']} (cont.)"
                )
            self.notes.append(
                f"outline slide {s['index']}: table split over {len(chunks)} slides ({per} rows each) to stay at {self.dense_min:.0f}pt or larger"
            )
            return
        if not fits:
            self.warn(
                f"table needs {sum(heights):.0f}pt of height at {size:.0f}pt but {pt(box.h):.0f}pt is available; it will run off the slide. Cut rows (keep at most ~7 rows x 5 columns; put the full table in an appendix), or set allow_split."
            )
        nrows = len(grid)
        gf = slide.shapes.add_table(
            nrows, ncols, box.x, box.y, box.w, int(sum(heights) * EMU_PER_PT)
        )
        gf.name = "Table"
        table = gf.table
        table.first_row = has_header
        table.horz_banding = True
        for j, w in enumerate(widths):
            table.columns[j].width = int(w * EMU_PER_PT)
        drift = box.w - sum(int(w * EMU_PER_PT) for w in widths)
        table.columns[ncols - 1].width = int(table.columns[ncols - 1].width + drift)
        for i, h in enumerate(heights):
            table.rows[i].height = int(h * EMU_PER_PT)
        zebra = self.theme.panel
        for i, row in enumerate(grid):
            is_header = has_header and i == 0
            data_i = i - (1 if has_header else 0)
            is_total = t["total_row"] and i == nrows - 1
            hl = data_i in t["highlight_rows"]
            for j in range(ncols):
                cell = table.cell(i, j)
                cell.margin_left = cell.margin_right = cell_mx
                cell.margin_top = cell.margin_bottom = cell_my
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = ""
                r = p.add_run()
                r.text = str(row[j])
                r.font.size = Pt(size)
                r.font.bold = bool(is_header or is_total or hl)
                r.font.name = self.font.family
                align = t["align"].get(j)
                if align == "center":
                    p.alignment = PP_ALIGN.CENTER
                elif align == "right" or (align is None and j in numeric_cols):
                    p.alignment = PP_ALIGN.RIGHT
                else:
                    p.alignment = PP_ALIGN.LEFT
                if is_header:
                    fill_shape(cell, self.theme.primary)
                    _apply_color(r.font.color, self.theme.text_on(self.theme.primary))
                else:
                    fill = (
                        zebra
                        if (is_total or hl or data_i % 2 == 1)
                        else self.theme.background
                    )
                    fill_shape(cell, fill)
                    _apply_color(r.font.color, self.theme.text_on(fill))
        gf._element.nvGraphicFramePr.cNvPr.set(
            "descr",
            f"Table: {', '.join(h for h in header if h)[:200]}; {len(rows_all)} rows.",
        )  # noqa: SLF001
        self._finish(slide)

    # kpi ------------------------------------------------------------

    def _slide_kpi(self, s):
        slide, box = self._standard(s)
        metrics = s["metrics"]
        n = len(metrics)
        if n > 4:
            self.warn(f"{n} KPI tiles; 1-4 numbers land, more read as a table.")
        if s.get("paras"):
            tiles_h = int(box.h * 0.6)
            tbox = Box(
                box.x, box.y + tiles_h + self.unit, box.w, box.h - tiles_h - self.unit
            )
        else:
            tiles_h, tbox = box.h, None
        cols = n if n <= 4 else math.ceil(n / 2)
        rows = math.ceil(n / cols)
        gap = self.gap
        tw = (box.w - gap * (cols - 1)) // cols
        th = min((tiles_h - gap * (rows - 1)) // rows, Inches(3.2))
        top = box.y + (tiles_h - (th * rows + gap * (rows - 1))) // 2
        ix, iy = Inches(0.15), Inches(0.12)
        hf = FontSpec(self.hfont.family, True)
        tiles = []
        for k, m in enumerate(metrics):
            value = m["value"]
            vtext = (
                format_number(value, m["format"] or self.number_code)
                if is_number(value)
                else str(value)
            )
            if m["unit"]:
                tight = m["unit"] in (
                    "%",
                    "x",
                    "×",
                    "K",
                    "k",
                    "M",
                    "m",
                    "B",
                    "bn",
                    "pp",
                    "pts",
                )
                vtext = f"{vtext}{'' if tight else ' '}{m['unit']}"
            lines: List[Tuple[str, float, bool, Paint]] = [
                (vtext, 0.0, True, self.theme.primary)
            ]
            lines.append((m["label"], self.body_min, False, self.theme.text))
            if m["delta"]:
                d = m["delta"].strip()
                if d.startswith(("-", "−", "▼", "↓")):
                    sign: Optional[float] = -1.0
                elif d.startswith(("+", "▲", "↑")):
                    sign = 1.0
                else:
                    first = parse_number(d.split()[0]) if d.split() else None
                    sign = (
                        None
                        if first is None
                        else (1.0 if first > 0 else (-1.0 if first < 0 else 0.0))
                    )
                if sign is None or sign == 0 or m["good"] in ("neutral", "none"):
                    paint = self.theme.muted
                else:
                    paint = (
                        self.theme.good
                        if (sign > 0) == (m["good"] != "down")
                        else self.theme.bad
                    )
                # Direction is in the text as well, never in the colour alone.
                if sign and not d.startswith(("▲", "▼", "↑", "↓")):
                    d = ("▲ " if sign > 0 else "▼ ") + d
                lines.append((d, max(self.dense_min, 16.0), True, paint))
            if m["note"]:
                lines.append((m["note"], self.dense_min, False, self.theme.muted))
            r, c = divmod(k, cols)
            row_n = cols if r < rows - 1 else n - cols * (rows - 1)
            offset = (box.w - (tw * row_n + gap * (row_n - 1))) // 2
            tile = self.shape(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Box(box.x + offset + c * (tw + gap), top + r * (th + gap), tw, th),
                f"KPI tile {k + 1}",
                self.theme.panel,
                round_adj=0.06,
            )
            tiles.append((tile, lines))

        def paras_at(v: float, lines) -> List[Para]:
            return [
                Para(
                    text,
                    level=-1,
                    bold=bold,
                    scale=1.0 if j == 0 else fixed / v,
                    space_before=0.15 if j else 0.0,
                    font=hf if j == 0 else None,
                )
                for j, (text, fixed, bold, _p) in enumerate(lines)
            ]

        w_pt = pt(tw - 2 * ix) * 0.98
        h_pt = pt(th - 2 * iy) * FIT_SLACK

        def fits_at(v: float) -> bool:
            for _tile, lines in tiles:
                if MEASURER.line_count(lines[0][0], hf, v, w_pt) > 1:
                    return False
                if (
                    block_height_pt(
                        paras_at(v, lines), v, w_pt, self.font, min_pt=self.min_pt
                    )
                    > h_pt
                ):
                    return False
            return True

        v = float(SIZES["kpi_max"])
        while v > SIZES["kpi_min"] and not fits_at(v):
            v -= 2
        if not fits_at(v):
            while v > 28 and not fits_at(v):
                v -= 2
            self.warn(
                f"KPI numbers set at {v:.0f}pt (below the {SIZES['kpi_min']}pt numeral size); shorten values or labels, or show fewer tiles."
            )
            if not fits_at(v):
                self.warn("KPI tile text does not fit; shorten the labels and notes.")
        for tile, lines in tiles:
            self.write(
                tile,
                paras_at(v, lines),
                v,
                colors=[p for *_x, p in lines],
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
                inset_x=ix,
                inset_y=iy,
            )
        if tbox is not None:
            self._side_text(slide, s["paras"], tbox, "Takeaways", SIZES["body_max"] - 2)
        self._finish(slide)

    # quote ------------------------------------------------------------

    def _slide_quote(self, s):
        slide, box = self._standard(s)
        x, w = self.grid(1, 10, box)
        inner = Box(x, box.y + self.unit, w, box.h - 2 * self.unit)
        quote = s["quote"].strip().strip('"“”')
        qf = FontSpec(self.hfont.family)
        paras = [Para(f"“{quote}”", level=-1, font=qf)]
        colors = [self.theme.text]
        if s.get("attribution"):
            paras.append(
                Para(f"— {s['attribution']}", level=-1, scale=0.6, space_before=0.8)
            )
            colors.append(self.theme.muted)
        tb = self.textbox(slide, inner, "Quote")
        size, ok, need, avail = self.fit(
            paras, inner.w, inner.h, 36, self.body_min, font=qf
        )
        if not ok:
            self.warn("quote does not fit; trim it to the line that matters.")
        self.write(tb, paras, size, colors=colors, anchor=MSO_ANCHOR.MIDDLE, font=qf)
        self.shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            Box(
                inner.x - Inches(0.3),
                inner.y + Inches(0.2),
                Inches(0.08),
                inner.h - Inches(0.4),
            ),
            "Accent bar",
            self.theme.primary,
        )
        self._finish(slide)

    # image ------------------------------------------------------------

    def _slide_image(self, s):
        from PIL import Image

        slide, box = self._standard(s)
        cap_h = 0
        if s.get("caption"):
            cap_paras = [Para(s["caption"], level=-1)]
            cap_h = int(
                (
                    block_height_pt(cap_paras, self.dense_min, pt(box.w) - 8, self.font)
                    / FIT_SLACK
                    + pt(2 * INSET_Y)
                    + 4
                )
                * EMU_PER_PT
            )
        area = Box(box.x, box.y, box.w, box.h - cap_h - (Inches(0.08) if cap_h else 0))
        if s.get("paras"):
            ix, iw_ = self.grid(0, 7, box)
            tx, tw = self.grid(7, 5, box)
            area = Box(ix, box.y, iw_, area.h)
            self._side_text(
                slide,
                s["paras"],
                Box(tx, box.y, tw, box.h),
                "Body",
                SIZES["body_max"] - 2,
            )
        with Image.open(s["path"]) as im:
            iw, ih = im.size
        scale = min(area.w / iw, area.h / ih)
        w, h = int(iw * scale), int(ih * scale)
        x = area.x + (area.w - w) // 2 if not s.get("paras") else area.x
        y = area.y + (area.h - h) // 2
        pic = slide.shapes.add_picture(s["path"], x, y, w, h)
        pic.name = f"Picture: {Path(s['path']).name}"
        pic._element.nvPicPr.cNvPr.set("descr", s.get("alt") or "")  # noqa: SLF001
        if not s.get("alt"):
            self.warn("image has no alt text; add 'alt' describing what it shows.")
        if cap_h:
            cb = self.textbox(
                slide, Box(area.x, y + h + Inches(0.08), area.w, cap_h), "Caption"
            )
            self.write(
                cb,
                [Para(s["caption"], level=-1)],
                self.dense_min,
                color=self.theme.muted,
                align=PP_ALIGN.CENTER if not s.get("paras") else PP_ALIGN.LEFT,
            )
        self._finish(slide)

    # diagrams ------------------------------------------------------------

    def _hl(self, s, i: int, title: str) -> bool:
        h = s.get("highlight")
        if h is None:
            return False
        if isinstance(h, (int, float)) and not isinstance(h, bool):
            return int(h) == i
        return str(h).strip().lower() == title.strip().lower()

    def _card_paras(self, item, title_scale=1.0, body_scale=0.85, title_font=None):
        ps, raws = [], []
        if item.get("title"):
            ps.append(
                Para(
                    _plain(item["title"]),
                    level=-1,
                    bold=True,
                    scale=title_scale,
                    font=title_font,
                )
            )
            raws.append(item["title"])
        if item.get("text"):
            ps.append(
                Para(
                    _plain(item["text"]), level=-1, scale=body_scale, space_before=0.25
                )
            )
            raws.append(item["text"])
        for text, lvl, bold in item.get("bullets") or []:
            ps.append(
                Para(
                    _plain(text),
                    level=max(0, lvl),
                    bold=bold,
                    scale=body_scale,
                    indent_pt=self._indent_pt(max(0, lvl), False),
                    space_before=0.2,
                )
            )
            raws.append(text)
        return ps, raws

    def _write_group(self, entries, max_pt: float, what: str, **kw):
        """Fit a group of (shape, paras, raws, colors) at one common size."""
        if not entries:
            return
        ix = kw.pop("inset_x", INSET_X)
        iy = kw.pop("inset_y", INSET_Y)
        # Shapes with their own inner text rectangle (chevrons, ovals,
        # trapezoids) are measured narrower but written with the normal inset.
        write_ix = kw.pop("write_inset_x", ix)
        results = [
            self.fit(
                ps,
                int(shp.width),
                int(shp.height),
                max_pt,
                None,
                inset_x=ix,
                inset_y=iy,
            )
            for shp, ps, _r, _c in entries
        ]
        size = min(r[0] for r in results)
        bad = [shp.name for (shp, *_x), r in zip(entries, results) if not r[1]]
        if bad:
            self.warn(
                f"{what}: text in {', '.join(bad[:4])}{'…' if len(bad) > 4 else ''} does not fit at {size:.0f}pt; shorten it."
            )
        for shp, ps, raws, colors in entries:
            self.write(
                shp,
                ps,
                size,
                colors=colors,
                raw_texts=raws,
                inset_x=write_ix,
                inset_y=iy,
                **kw,
            )

    def _slide_diagram(self, s):
        slide, box = self._standard(s)
        if s.get("paras"):
            tx, tw = self.grid(8, 4, box)
            self._side_text(
                slide,
                s["paras"],
                Box(tx, box.y, tw, box.h),
                "Body",
                SIZES["body_max"] - 2,
            )
            dx, dw = self.grid(0, 8, box)
            box = Box(dx, box.y, dw, box.h)
        getattr(self, f"_diagram_{s['layout']}")(slide, s, box)
        self._finish(slide)

    def _fill_for(self, s, k: int, title: str) -> Paint:
        return self.theme.highlight if self._hl(s, k, title) else self.theme.primary

    def _diagram_process(self, slide, s, box: Box):
        items = s["items"]
        n = len(items)
        if n > 6:
            self.warn(
                f"{n} process steps; 3-6 read at a glance (group the rest into phases)."
            )
        per_row = n if n <= 6 else math.ceil(n / 2)
        rows = math.ceil(n / per_row)
        style = s.get("style") or "chevrons"
        row_gap = self.unit
        row_h = (box.h - row_gap * (rows - 1)) // rows
        has_detail = any(it.get("text") or it.get("bullets") for it in items)
        head_h = min(Inches(1.0), int(row_h * (0.32 if has_detail else 0.6)))
        heads, details = [], []
        for k, it in enumerate(items):
            r, c = divmod(k, per_row)
            gap = Inches(0.08) if style == "chevrons" else Inches(0.45)
            sw = (box.w - gap * (per_row - 1)) // per_row
            x = box.x + c * (sw + gap)
            y = box.y + r * (row_h + row_gap)
            kind = (
                (MSO_SHAPE.PENTAGON if c == 0 else MSO_SHAPE.CHEVRON)
                if style == "chevrons"
                else MSO_SHAPE.ROUNDED_RECTANGLE
            )
            paint = self._fill_for(s, k, it["title"])
            head = self.shape(
                slide,
                kind,
                Box(x, y, sw, head_h),
                f"Step {k + 1}",
                paint,
                round_adj=0.1 if style != "chevrons" else None,
            )
            numbered = s.get("numbered", True) is not False and not re.match(
                r"^\s*\d+[.)]", it["title"]
            )
            label = f"{k + 1}. {it['title']}" if numbered else it["title"]
            heads.append(
                (
                    head,
                    [Para(_plain(label), level=-1, bold=True)],
                    [label],
                    [self.theme.text_on(paint)],
                )
            )
            if style != "chevrons" and c < per_row - 1 and k < n - 1:
                ay = y + head_h // 2
                self.arrow(
                    slide,
                    x + sw + Inches(0.06),
                    ay,
                    x + sw + gap - Inches(0.06),
                    ay,
                    f"Arrow {k + 1}",
                )
            if has_detail:
                d = Box(
                    x,
                    y + head_h + Inches(0.12),
                    sw - (Inches(0.12) if style == "chevrons" else 0),
                    row_h - head_h - Inches(0.12),
                )
                tb = self.textbox(slide, d, f"Step {k + 1} detail")
                ps, raws = self._card_paras(
                    {"text": it.get("text"), "bullets": it.get("bullets")},
                    body_scale=1.0,
                )
                details.append((tb, ps, raws, [self.theme.text] * len(ps)))
        point = int(head_h * 0.5)
        self._write_group(
            heads,
            20,
            "process steps",
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            inset_x=int(point * 1.05) + INSET_X if style == "chevrons" else INSET_X,
            write_inset_x=INSET_X,
        )
        self._write_group(details, 18, "step details")

    def _diagram_boxes(self, slide, s, box: Box, cols: Optional[int] = None):
        items = s["items"]
        n = len(items)
        cols = cols or s.get("columns_n") or (n if n <= 5 else math.ceil(n / 2))
        rows = math.ceil(n / cols)
        gap = self.gap
        bw = (box.w - gap * (cols - 1)) // cols
        bh = (box.h - gap * (rows - 1)) // rows
        entries = []
        hf = FontSpec(self.hfont.family, True)
        for k, it in enumerate(items):
            r, c = divmod(k, cols)
            hl = self._hl(s, k, it["title"])
            fill = self.theme.primary if hl else self.theme.panel
            shp = self.shape(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Box(box.x + c * (bw + gap), box.y + r * (bh + gap), bw, bh),
                f"Box {k + 1}",
                fill,
                round_adj=0.05,
            )
            ps, raws = self._card_paras(
                it, title_scale=1.1, body_scale=0.9, title_font=hf
            )
            on = self.theme.text_on(fill)
            colors = [
                (
                    on
                    if hl
                    else (
                        self.theme.primary
                        if j == 0 and it.get("title")
                        else self.theme.text
                    )
                )
                for j in range(len(ps))
            ]
            entries.append((shp, ps, raws, colors))
        self._write_group(
            entries,
            20,
            "boxes",
            align=PP_ALIGN.LEFT,
            inset_x=Inches(0.15),
            inset_y=Inches(0.12),
        )

    def _diagram_grid(self, slide, s, box: Box):
        items = s["items"]
        if not any("row" in it or "col" in it for it in items):
            return self._diagram_boxes(
                slide,
                s,
                box,
                cols=s.get("columns_n") or min(4, math.ceil(math.sqrt(len(items)))),
            )
        nrows = max(int(it.get("row", 0)) + int(it.get("rowspan", 1)) for it in items)
        ncols = max(int(it.get("col", 0)) + int(it.get("colspan", 1)) for it in items)
        gap = Inches(0.1)
        cw = (box.w - gap * (ncols - 1)) / ncols
        rh = (box.h - gap * (nrows - 1)) / nrows
        entries = []
        occupied: Dict[Tuple[int, int], str] = {}
        hf = FontSpec(self.hfont.family, True)
        for k, it in enumerate(items):
            r, c = int(it.get("row", 0)), int(it.get("col", 0))
            rs, cs = int(it.get("rowspan", 1)), int(it.get("colspan", 1))
            for rr in range(r, r + rs):
                for cc in range(c, c + cs):
                    if (rr, cc) in occupied:
                        self.warn(
                            f"grid cells '{occupied[(rr, cc)]}' and '{it['title']}' both cover row {rr}, column {cc}."
                        )
                    occupied[(rr, cc)] = it["title"]
            b = Box(
                int(box.x + c * (cw + gap)),
                int(box.y + r * (rh + gap)),
                int(cw * cs + gap * (cs - 1)),
                int(rh * rs + gap * (rs - 1)),
            )
            hl = self._hl(s, k, it["title"])
            fill = self.theme.primary if hl else self.theme.panel
            shp = self.shape(
                slide, MSO_SHAPE.RECTANGLE, b, f"Cell {k + 1}", fill, self.theme.line
            )
            ps, raws = self._card_paras(
                it, title_scale=1.0, body_scale=0.9, title_font=hf
            )
            on = self.theme.text_on(fill)
            colors = [
                (
                    on
                    if hl
                    else (
                        self.theme.primary
                        if j == 0 and it.get("title")
                        else self.theme.text
                    )
                )
                for j in range(len(ps))
            ]
            entries.append((shp, ps, raws, colors))
        self._write_group(
            entries,
            16,
            "grid cells",
            align=PP_ALIGN.LEFT,
            inset_x=Inches(0.1),
            inset_y=Inches(0.08),
        )

    def _diagram_stack(self, slide, s, box: Box):
        """Layered stack: full-width bands, top layer first (architecture, maturity levels)."""
        items = s["items"]
        n = len(items)
        gap = Inches(0.08)
        bh = (box.h - gap * (n - 1)) // n
        entries = []
        hf = FontSpec(self.hfont.family, True)
        for k, it in enumerate(items):
            hl = self._hl(s, k, it["title"])
            fill = (
                self.theme.highlight
                if hl
                else self.theme.primary.tinted(min(0.75, 0.18 * k))
            )
            shp = self.shape(
                slide,
                MSO_SHAPE.RECTANGLE,
                Box(box.x, box.y + k * (bh + gap), box.w, bh),
                f"Layer {k + 1}",
                fill,
            )
            ps, raws = self._card_paras(
                it, title_scale=1.0, body_scale=0.85, title_font=hf
            )
            entries.append((shp, ps, raws, [self.theme.text_on(fill)] * len(ps)))
        self._write_group(
            entries,
            20,
            "stack layers",
            align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.MIDDLE,
            inset_x=Inches(0.2),
            inset_y=Inches(0.06),
        )

    def _diagram_matrix(self, slide, s, box: Box):
        items = s["items"]
        xa, ya = s.get("x_axis") or {}, s.get("y_axis") or {}
        xa = {"label": xa} if isinstance(xa, str) else xa
        ya = {"label": ya} if isinstance(ya, str) else ya
        lab_w, lab_h = Inches(0.45), Inches(0.42)
        grid = Box(
            box.x + lab_w + Inches(0.1),
            box.y,
            box.w - lab_w - Inches(0.1),
            box.h - lab_h - Inches(0.08),
        )
        gap = Inches(0.08)
        qw, qh = (grid.w - gap) // 2, (grid.h - gap) // 2
        points = s.get("points") or []
        entries = []
        hf = FontSpec(self.hfont.family, True)
        for k in range(4):
            r, c = divmod(k, 2)
            b = Box(grid.x + c * (qw + gap), grid.y + r * (qh + gap), qw, qh)
            it = (
                items[k] if k < len(items) else {"title": "", "text": "", "bullets": []}
            )
            hl = self._hl(s, k, it.get("title", ""))
            fill = self.theme.primary if hl else self.theme.panel
            on = self.theme.text_on(fill)
            if points:
                self.shape(
                    slide, MSO_SHAPE.RECTANGLE, b, f"Background quadrant {k + 1}", fill
                )
                if it.get("title"):
                    tb = self.textbox(
                        slide,
                        Box(b.x, b.y, b.w, Inches(0.5)),
                        f"Quadrant {k + 1} label",
                    )
                    entries.append(
                        (
                            tb,
                            [Para(_plain(it["title"]), level=-1, bold=True)],
                            [it["title"]],
                            [on if hl else self.theme.primary],
                        )
                    )
            else:
                shp = self.shape(
                    slide, MSO_SHAPE.RECTANGLE, b, f"Quadrant {k + 1}", fill
                )
                ps, raws = self._card_paras(
                    it, title_scale=1.1, body_scale=0.85, title_font=hf
                )
                colors = [
                    (on if hl else (self.theme.primary if j == 0 else self.theme.text))
                    for j in range(len(ps))
                ]
                entries.append((shp, ps, raws, colors))
        self._write_group(
            entries,
            20,
            "matrix quadrants",
            align=PP_ALIGN.LEFT,
            inset_x=Inches(0.15),
            inset_y=Inches(0.12),
        )
        self.arrow(
            slide,
            grid.x - Inches(0.06),
            grid.bottom,
            grid.x - Inches(0.06),
            grid.y,
            "Y axis",
        )
        self.arrow(
            slide,
            grid.x,
            grid.bottom + Inches(0.06),
            grid.right,
            grid.bottom + Inches(0.06),
            "X axis",
        )

        def axis_text(a):
            return "   ".join(
                t
                for t in (
                    a.get("low") and f"← {a['low']}",
                    a.get("label"),
                    a.get("high") and f"{a['high']} →",
                )
                if t
            )

        x_text, y_text = axis_text(xa), axis_text(ya)
        if x_text:
            xb = self.textbox(
                slide,
                Box(grid.x, grid.bottom + Inches(0.1), grid.w, lab_h - Inches(0.02)),
                "X axis label",
            )
            self.write(
                xb,
                [Para(x_text, level=-1)],
                self.dense_min,
                color=self.theme.muted,
                align=PP_ALIGN.CENTER,
            )
        if y_text:
            # A rotated box turns about its centre: size it for the rotated span.
            span = grid.h
            yb = self.textbox(
                slide,
                Box(
                    box.x + lab_w // 2 - span // 2,
                    grid.y + grid.h // 2 - lab_w // 2,
                    span,
                    lab_w,
                ),
                "Y axis label",
            )
            yb.rotation = 270
            self.write(
                yb,
                [Para(y_text, level=-1)],
                self.dense_min,
                color=self.theme.muted,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        for k, p in enumerate(points):
            if not isinstance(p, dict):
                self.warn(
                    f"matrix point {p!r} must be {{label, x, y}} with x, y between 0 and 1."
                )
                continue
            try:
                px, py = float(p.get("x")), float(p.get("y"))
            except (TypeError, ValueError):
                self.warn(
                    f"matrix point {p.get('label')!r} needs numeric x and y between 0 and 1."
                )
                continue
            if not (0 <= px <= 1 and 0 <= py <= 1):
                self.warn(
                    f"matrix point {p.get('label')!r} is outside 0..1; it was placed at the edge."
                )
                px, py = min(1, max(0, px)), min(1, max(0, py))
            d = Inches(0.2)
            cx = min(max(grid.x + int(px * grid.w), grid.x + d), grid.right - d)
            cy = min(
                max(grid.bottom - int(py * grid.h), grid.y + Inches(0.55)),
                grid.bottom - d,
            )
            self.shape(
                slide,
                MSO_SHAPE.OVAL,
                Box(cx - d // 2, cy - d // 2, d, d),
                f"Point {k + 1}",
                self.theme.highlight,
            )
            label = str(p.get("label") or "")
            lw = int(
                (MEASURER.width_pt(label, self.font, self.dense_min) + 14) * EMU_PER_PT
            )
            lx = cx + d // 2 + Inches(0.04)
            if lx + lw > grid.right:
                lx = cx - d // 2 - Inches(0.04) - lw
            lb = self.textbox(
                slide,
                Box(lx, cy - Inches(0.17), lw, Inches(0.34)),
                f"Point {k + 1} label",
            )
            self.write(
                lb,
                [Para(label, level=-1)],
                self.dense_min,
                color=self.theme.text,
                anchor=MSO_ANCHOR.MIDDLE,
                inset_y=0,
            )

    def _diagram_pyramid(self, slide, s, box: Box):
        items = s["items"]
        n = len(items)
        has_text = any(it.get("text") or it.get("bullets") for it in items)
        pw = int(box.w * (0.5 if has_text else 0.8))
        px = box.x if has_text else box.x + (box.w - pw) // 2
        gap = Inches(0.06)
        lh = (box.h - gap * (n - 1)) // n
        entries, side = [], []
        for k, it in enumerate(items):
            y = box.y + k * (lh + gap)
            # Each level's widths follow the triangle at its top and bottom edges.
            top_w = pw * (k * (lh + gap)) / box.h
            bot_w = pw * (k * (lh + gap) + lh) / box.h
            x = px + (pw - bot_w) / 2
            paint = (
                self.theme.highlight
                if self._hl(s, k, it["title"])
                else self.theme.primary.tinted(min(0.6, 0.14 * k))
            )
            if k == 0:
                shp = self.shape(
                    slide,
                    MSO_SHAPE.ISOSCELES_TRIANGLE,
                    Box(int(x), int(y), int(bot_w), int(lh)),
                    f"Level {k + 1}",
                    paint,
                )
                inner_w = (
                    bot_w * 0.46
                )  # the triangle's own text box is its lower-middle half
            else:
                shp = self.shape(
                    slide,
                    MSO_SHAPE.TRAPEZOID,
                    Box(int(x), int(y), int(bot_w), int(lh)),
                    f"Level {k + 1}",
                    paint,
                )
                shp.adjustments[0] = max(
                    0.0, min(1.0, ((bot_w - top_w) / 2) / min(bot_w, lh))
                )
                inner_w = top_w
            ix = max(INSET_X, int((bot_w - inner_w) / 2 + inner_w * 0.04))
            entries.append(
                (
                    shp,
                    [Para(_plain(it["title"]), level=-1, bold=True)],
                    [it["title"]],
                    [self.theme.text_on(paint)],
                    ix,
                )
            )
            if has_text:
                sx = px + pw + self.unit
                tb = self.textbox(
                    slide,
                    Box(sx, int(y), box.right - sx, int(lh)),
                    f"Level {k + 1} detail",
                )
                ps, raws = self._card_paras(
                    {"text": it.get("text"), "bullets": it.get("bullets")},
                    body_scale=1.0,
                )
                side.append((tb, ps, raws, [self.theme.text] * len(ps)))
        sizes = [
            self.fit(
                ps,
                int(shp.width),
                int(shp.height),
                20,
                None,
                inset_x=ix,
                inset_y=Inches(0.04),
            )
            for shp, ps, raws, colors, ix in entries
        ]
        size = min(x[0] for x in sizes)
        for (shp, ps, raws, colors, ix), f in zip(entries, sizes):
            if not f[1]:
                self.warn(
                    f"{shp.name} label does not fit inside the pyramid; shorten it."
                )
            self.write(
                shp,
                ps,
                size,
                colors=colors,
                raw_texts=raws,
                align=PP_ALIGN.CENTER,
                anchor=(
                    MSO_ANCHOR.BOTTOM if shp.name == "Level 1" else MSO_ANCHOR.MIDDLE
                ),
                inset_y=Inches(0.04),
            )
        self._write_group(side, 18, "pyramid details", anchor=MSO_ANCHOR.MIDDLE)

    def _diagram_cycle(self, slide, s, box: Box):
        items = s["items"]
        n = len(items)
        cx, cy = box.x + box.w // 2, box.y + box.h // 2
        bw = min(Inches(2.6), int(box.w / (2.2 if n <= 4 else 3.0)))
        bh = min(Inches(1.2), int(box.h / (2.6 if n <= 4 else 3.4)))
        rx, ry = (box.w - bw) / 2, (box.h - bh) / 2
        centers = [
            (
                cx + rx * math.cos(-math.pi / 2 + 2 * math.pi * k / n),
                cy + ry * math.sin(-math.pi / 2 + 2 * math.pi * k / n),
            )
            for k in range(n)
        ]
        entries = []
        for k, (it, (x, y)) in enumerate(zip(items, centers)):
            paint = self._fill_for(s, k, it["title"])
            shp = self.shape(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Box(int(x - bw / 2), int(y - bh / 2), bw, bh),
                f"Stage {k + 1}",
                paint,
                round_adj=0.2,
            )
            ps, raws = self._card_paras(it, title_scale=1.0, body_scale=0.8)
            entries.append((shp, ps, raws, [self.theme.text_on(paint)] * len(ps)))
        self._write_group(
            entries, 18, "cycle stages", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE
        )

        def edge_point(x0, y0, x1, y1):
            dx, dy = x1 - x0, y1 - y0
            tx = (bw / 2) / abs(dx) if dx else float("inf")
            ty = (bh / 2) / abs(dy) if dy else float("inf")
            t = min(tx, ty)
            return x0 + dx * t, y0 + dy * t

        for k in range(n):
            (x0, y0), (x1, y1) = centers[k], centers[(k + 1) % n]
            ax, ay = edge_point(x0, y0, x1, y1)
            bx, by = edge_point(x1, y1, x0, y0)
            dist = math.hypot(bx - ax, by - ay)
            if dist < Inches(0.3):
                continue
            shrink = Inches(0.08) / dist
            self.arrow(
                slide,
                ax + (bx - ax) * shrink,
                ay + (by - ay) * shrink,
                bx - (bx - ax) * shrink,
                by - (by - ay) * shrink,
                f"Cycle arrow {k + 1}",
                self.theme.muted,
            )
        if s.get("center"):
            cwid = int(min(rx * 2 - bw, ry * 2 - bh) * 0.9)
            if cwid > Inches(1.0):
                cb = self.shape(
                    slide,
                    MSO_SHAPE.OVAL,
                    Box(cx - cwid // 2, cy - cwid // 2, cwid, cwid),
                    "Cycle centre",
                    self.theme.panel,
                )
                ps = [Para(_plain(s["center"]), level=-1, bold=True)]
                size, ok, *_ = self.fit(ps, int(cwid * 0.66), int(cwid * 0.66), 20)
                if not ok:
                    self.warn("the cycle's centre label does not fit; shorten it.")
                self.write(
                    cb,
                    ps,
                    size,
                    color=self.theme.primary,
                    align=PP_ALIGN.CENTER,
                    anchor=MSO_ANCHOR.MIDDLE,
                )
            else:
                self.warn(
                    "no room for the cycle's centre label; put it in the title instead."
                )

    def _diagram_flow(self, slide, s, box: Box, direction: Optional[str] = None):
        items = s["items"]
        direction = direction or s.get("direction", "right")
        ids = [str(it.get("id") or it["title"]) for it in items]
        index = {i: k for k, i in enumerate(ids)}
        edges = [
            (index[a], index[b], label)
            for a, b, label in s["edges"]
            if a in index and b in index
        ]
        if not edges:
            edges = [(k, k + 1, "") for k in range(len(items) - 1)]
        n = len(items)
        layer = [0] * n
        order = self._topo(n, edges)
        pos = {v: i for i, v in enumerate(order)}
        for v in order:
            for a, b, _l in edges:
                if a == v and pos[b] > pos[a]:
                    layer[b] = max(layer[b], layer[a] + 1)
        nl = max(layer) + 1
        by_layer: Dict[int, List[int]] = {}
        for v in order:
            by_layer.setdefault(layer[v], []).append(v)
        width_n = max(len(v) for v in by_layer.values())
        main = box.w if direction == "right" else box.h
        cross = box.h if direction == "right" else box.w
        label_need = max(
            (
                MEASURER.width_pt(lbl, self.font, self.dense_min)
                for _a, _b, lbl in edges
                if lbl
            ),
            default=0,
        )
        min_gap = Inches(0.45) + (
            int(label_need * EMU_PER_PT)
            if direction == "right"
            else (Inches(0.3) if label_need else 0)
        )
        node_main = min(
            Inches(2.8) if direction == "right" else Inches(1.1),
            int((main - min_gap * (nl - 1)) / nl),
        )
        gap_main = (main - node_main * nl) / max(1, nl - 1) if nl > 1 else 0
        node_cross = min(
            Inches(1.3) if direction == "right" else Inches(2.6),
            int(cross / width_n - Inches(0.3)),
        )
        if node_main < Inches(0.6) or node_cross < Inches(0.5):
            self.warn(
                f"{n} nodes over {nl} layers leave boxes too small to read; split the flow or use fewer nodes."
            )
        shapes: Dict[int, Any] = {}
        entries = []
        for L, members in by_layer.items():
            m = len(members)
            for j, v in enumerate(members):
                c_off = cross * (j + 0.5) / m - node_cross / 2
                m_off = L * (node_main + gap_main)
                if direction == "right":
                    b = Box(
                        int(box.x + m_off),
                        int(box.y + c_off),
                        int(node_main),
                        int(node_cross),
                    )
                else:
                    b = Box(
                        int(box.x + c_off),
                        int(box.y + m_off),
                        int(node_cross),
                        int(node_main),
                    )
                it = items[v]
                paint = self._fill_for(s, v, it["title"])
                shp = self.shape(
                    slide,
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    b,
                    f"Node {v + 1}",
                    paint,
                    round_adj=0.12,
                )
                shapes[v] = shp
                ps, raws = self._card_paras(it, title_scale=1.0, body_scale=0.8)
                entries.append((shp, ps, raws, [self.theme.text_on(paint)] * len(ps)))
        self._write_group(
            entries, 18, "flow nodes", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE
        )
        # Connection sites on a rectangle: 0 top, 1 left, 2 bottom, 3 right.
        out_site, in_site = (3, 1) if direction == "right" else (2, 0)
        for k, (a, b, label) in enumerate(edges):
            back = layer[b] <= layer[a]
            sa, sb = (
                (out_site, in_site)
                if not back
                else ((0, 0) if direction == "right" else (3, 3))
            )
            elbow = (
                back
                or layer[b] - layer[a] > 1
                or len(by_layer[layer[a]]) != len(by_layer[layer[b]])
            )
            conn = self.arrow(
                slide,
                0,
                0,
                1,
                1,
                f"Connector {k + 1}",
                begin=(shapes[a], sa),
                end=(shapes[b], sb),
                elbow=elbow,
            )
            if label:
                mx = (conn.begin_x + conn.end_x) // 2
                my = (conn.begin_y + conn.end_y) // 2
                lw = int(
                    (MEASURER.width_pt(label, self.font, self.dense_min) + 12)
                    * EMU_PER_PT
                )
                lh = Inches(0.3)
                lbox = (
                    Box(mx - lw // 2, my - lh - Inches(0.02), lw, lh)
                    if direction == "right"
                    else Box(mx + Inches(0.06), my - lh // 2, lw, lh)
                )
                lb = self.textbox(slide, lbox, f"Edge label {k + 1}")
                self.write(
                    lb,
                    [Para(label, level=-1)],
                    self.dense_min,
                    color=self.theme.muted,
                    align=PP_ALIGN.CENTER,
                    inset_y=0,
                )

    def _diagram_tree(self, slide, s, box: Box):
        return self._diagram_flow(slide, s, box, direction="down")

    @staticmethod
    def _topo(n: int, edges) -> List[int]:
        indeg = [0] * n
        adj: Dict[int, List[int]] = {}
        for a, b, _l in edges:
            adj.setdefault(a, []).append(b)
            indeg[b] += 1
        order: List[int] = []
        seen = set()
        ready = [v for v in range(n) if indeg[v] == 0]
        while len(order) < n:
            if not ready:  # a cycle: start from the lowest unvisited node
                ready = [min(v for v in range(n) if v not in seen)]
            v = ready.pop(0)
            if v in seen:
                continue
            seen.add(v)
            order.append(v)
            for b in adj.get(v, []):
                indeg[b] -= 1
                if indeg[b] <= 0 and b not in seen:
                    ready.append(b)
        return order

    # timeline ------------------------------------------------------------

    def _slide_timeline(self, s):
        slide, box = self._standard(s)
        if s.get("lanes"):
            self._roadmap(slide, s, box)
        else:
            self._milestones(slide, s, box)
        self._finish(slide)

    def _milestones(self, slide, s, box: Box):
        items = s["items"]
        n = len(items)
        alternate = n > 4
        y_line = box.y + (box.h // 2 if alternate else int(box.h * 0.22))
        pad = Inches(0.3)
        span = box.w - 2 * pad
        xs = [
            box.x + pad + (span * k / (n - 1) if n > 1 else span / 2) for k in range(n)
        ]
        self.shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            Box(box.x, y_line - Inches(0.025), box.w, Inches(0.05)),
            "Timeline axis",
            self.theme.line,
        )
        d = Inches(0.26)
        gap = Inches(0.15)
        step = 2 if alternate else 1  # labels on one side are this many markers apart

        def span_for(k: int) -> Tuple[int, int]:
            """Room for label k: halfway to its same-side neighbours, inside the box."""
            left = box.x if k - step < 0 else int((xs[k - step] + xs[k]) / 2 + gap / 2)
            right = (
                box.right
                if k + step >= n
                else int((xs[k] + xs[k + step]) / 2 - gap / 2)
            )
            half = min(xs[k] - left, right - xs[k], Inches(1.6))
            lo, hi = int(max(left, xs[k] - half)), int(min(right, xs[k] + half))
            if k == 0 or k == n - 1:  # an end label may lean inwards
                lo, hi = (
                    (int(left), int(min(right, left + 2 * half)))
                    if k == 0
                    else (int(max(left, right - 2 * half)), int(right))
                )
            return lo, hi

        entries = []
        for k, (it, x) in enumerate(zip(items, xs)):
            hl = self._hl(s, k, it["title"])
            self.shape(
                slide,
                MSO_SHAPE.OVAL,
                Box(int(x - d / 2), int(y_line - d / 2), d, d),
                f"Milestone {k + 1} marker",
                self.theme.highlight if hl else self.theme.primary,
            )
            above = alternate and k % 2 == 0
            lo, hi = span_for(k)
            if above:
                lb = Box(lo, box.y, hi - lo, int(y_line - d / 2 - Inches(0.1) - box.y))
            else:
                top = int(y_line + d / 2 + Inches(0.1))
                lb = Box(lo, top, hi - lo, box.bottom - top)
            tb = self.textbox(slide, lb, f"Milestone {k + 1}")
            ps, raws, colors = [], [], []
            if it.get("date"):
                ps.append(Para(it["date"], level=-1, bold=True, scale=0.9))
                raws.append(it["date"])
                colors.append(self.theme.highlight if hl else self.theme.primary)
            p2, r2 = self._card_paras(
                {
                    "title": it["title"],
                    "text": it.get("text"),
                    "bullets": it.get("bullets"),
                },
                title_scale=1.0,
                body_scale=0.85,
            )
            ps += p2
            raws += r2
            colors += [self.theme.text] + [self.theme.muted] * (len(p2) - 1)
            entries.append((tb, ps, raws, colors, above))
        results = [
            self.fit(ps, int(tb.width), int(tb.height), 18) for tb, ps, *_r in entries
        ]
        size = min(r[0] for r in results)
        for (tb, ps, raws, colors, above), r in zip(entries, results):
            if not r[1]:
                self.warn(f"{tb.name} text does not fit; shorten it.")
            self.write(
                tb,
                ps,
                size,
                colors=colors,
                raw_texts=raws,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.BOTTOM if above else MSO_ANCHOR.TOP,
            )

    def _roadmap(self, slide, s, box: Box):
        periods = s["periods"]
        lanes = s["lanes"]
        np_ = len(periods)
        bold = FontSpec(self.font.family, True)
        label_w = int(
            min(
                box.w * 0.22,
                max(
                    Inches(1.2),
                    max(
                        (MEASURER.width_pt(ln["name"], bold, 16) for ln in lanes),
                        default=0,
                    )
                    * EMU_PER_PT
                    + Inches(0.3),
                ),
            )
        )
        head_h = Inches(0.45)
        gx, gw = box.x + label_w, box.w - label_w
        colw = gw / np_
        heads = []
        for j, per in enumerate(periods):
            tb = self.textbox(
                slide,
                Box(
                    int(gx + j * colw + Inches(0.05)),
                    box.y,
                    int(colw - Inches(0.1)),
                    head_h,
                ),
                f"Period {j + 1}",
            )
            heads.append(
                (tb, [Para(per, level=-1, bold=True)], [per], [self.theme.muted])
            )
        self._write_group(
            heads, 16, "period headers", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE
        )

        def idx(v):
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                return int(v)
            return periods.index(str(v).strip())

        lane_rows = []
        for ln in lanes:
            rows_end: List[int] = []
            placed = []
            for bar in ln["items"]:
                a, b = sorted((idx(bar["start"]), idx(bar["end"])))
                for r, end in enumerate(rows_end):
                    if a > end:
                        rows_end[r] = b
                        placed.append((bar, a, b, r))
                        break
                else:
                    rows_end.append(b)
                    placed.append((bar, a, b, len(rows_end) - 1))
            lane_rows.append((ln, placed, max(1, len(rows_end))))
        total_rows = sum(r for _l, _p, r in lane_rows)
        body_top = box.y + head_h + Inches(0.05)
        row_h = (box.bottom - body_top) / total_rows
        y = body_top
        lane_labels, bars = [], []
        for li, (ln, placed, nrows) in enumerate(lane_rows):
            lh = row_h * nrows
            if li % 2 == 0:
                self.shape(
                    slide,
                    MSO_SHAPE.RECTANGLE,
                    Box(box.x, int(y), box.w, int(lh)),
                    f"Background lane {li + 1}",
                    self.theme.panel,
                )
            lb = self.textbox(
                slide,
                Box(box.x, int(y), label_w - Inches(0.1), int(lh)),
                f"Lane {li + 1}",
            )
            lane_labels.append(
                (
                    lb,
                    [Para(ln["name"], level=-1, bold=True)],
                    [ln["name"]],
                    [self.theme.text],
                )
            )
            for k, (bar, a, b, r) in enumerate(placed):
                bx = gx + a * colw + Inches(0.05)
                bw = (b - a + 1) * colw - Inches(0.1)
                by = y + r * row_h + row_h * 0.15
                status = bar.get("status", "").lower()
                paint = self.theme.palette[li % min(4, len(self.theme.palette))]
                if status in ("done", "complete", "completed"):
                    paint = paint.tinted(0.55)
                shp = self.shape(
                    slide,
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Box(int(bx), int(by), int(bw), int(row_h * 0.7)),
                    f"Bar {li + 1}.{k + 1}",
                    paint,
                    round_adj=0.25,
                )
                label = bar["label"] + (
                    f" ({bar['status']})" if bar.get("status") else ""
                )
                bars.append(
                    (shp, [Para(label, level=-1)], [label], [self.theme.text_on(paint)])
                )
            y += lh
        self._write_group(lane_labels, 16, "lane labels", anchor=MSO_ANCHOR.MIDDLE)
        self._write_group(
            bars,
            16,
            "roadmap bars",
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            inset_x=Inches(0.06),
            inset_y=Inches(0.02),
        )


# ----------------------------------------------------------------------
# Entry points
# ----------------------------------------------------------------------


def save_deck(prs, output: Path) -> None:
    """Write the deck. The caller has already checked the path may be written."""
    if not output.parent.is_dir():
        raise OutlineError(
            [
                f"the folder {output.parent} does not exist. Save the deck in an existing folder (or ask the user to create it)."
            ]
        )
    prs.save(str(output))


def _renumber(prs) -> None:
    """Refresh the cached text of slide-number fields after slides move."""
    for i, slide in enumerate(prs.slides, 1):
        for fld in slide.shapes._spTree.iter(qn("a:fld")):  # noqa: SLF001
            if fld.get("type") == "slidenum":
                t = fld.find(qn("a:t"))
                if t is not None:
                    t.text = str(i)


def _verify_summary(path: str, normalized, builder, rules_prefs) -> Dict[str, Any]:
    from .verify import verify_deck

    report = verify_deck(
        path,
        expected_slides=normalized.get("expected_slides"),
        outline=normalized,
        slide_map=builder.slide_map,
        require_notes=bool(
            normalized.get("require_notes")
            or builder.tokens.get("speaker_notes") is True
        ),
        required_items=normalized.get("required_items") or None,
        render=False,
        prefs=rules_prefs,
        tokens=builder.tokens,
    )
    return {
        k: report[k]
        for k in ("status", "style_score", "summary", "issues", "required_items")
        if k in report
    }


def build_deck(
    outline: Any,
    output_path: str,
    theme: Optional[str] = None,
    template_path: Optional[str] = None,
    prefs: Optional[Dict[str, Any]] = None,
    *,
    verify: bool = True,
    base_dir: Optional[str] = None,
) -> Dict[str, Any]:
    """Build *outline* into *output_path*. Raises OutlineError on a bad outline."""
    started = time.perf_counter()
    out = Path(output_path)
    if out.suffix.lower() != ".pptx":
        raise OutlineError([f"output_path must end in .pptx, got {out.name!r}."])
    if template_path:
        tp = Path(template_path)
        if not tp.is_file():
            raise OutlineError([f"template {template_path} does not exist."])
        if tp.suffix.lower() not in (".pptx", ".potx", ".pptm", ".potm"):
            raise OutlineError(
                [f"template must be a .pptx or .potx file, got {tp.suffix!r}."]
            )
        if tp.resolve() == out.resolve():
            raise OutlineError(
                [
                    "output_path is the template itself; save the deck under a new name so the template is kept."
                ]
            )
    normalized = normalize_outline(outline, Path(base_dir) if base_dir else None)
    builder = DeckBuilder(
        normalized, theme=theme, template_path=template_path, prefs=prefs
    )
    prs = builder.build()
    save_deck(prs, out)
    result: Dict[str, Any] = {
        "status": "ok",
        "path": str(out),
        "slides": len(prs.slides),
        "outline_slides": len(normalized["slides"]),
        "slide_map": {str(i): nums for i, nums in builder.slide_map},
        "theme": builder.tokens["theme"],
        "look_from": builder.template_source,
        "text_metrics": MEASURER.method(builder.font),
        "warnings": builder.warnings,
        "notes": builder.notes,
    }
    if verify:
        merged = merge_prefs(normalized.get("style") or {}, prefs or {})
        result["verify"] = _verify_summary(str(out), normalized, builder, merged)
        if result["verify"]["status"] != "pass":
            result["status"] = "built_with_issues"
    result["build_seconds"] = round(time.perf_counter() - started, 3)
    return result


def _edit(
    deck_path: str,
    slide_spec: Any,
    *,
    position: Optional[int],
    replace_index: Optional[int],
    prefs: Optional[Dict[str, Any]],
    base_dir: Optional[str],
) -> Dict[str, Any]:
    started = time.perf_counter()
    deck = Path(deck_path)
    if not deck.is_file():
        raise OutlineError(
            [f"deck {deck_path} does not exist; build it first with build_deck."]
        )
    if isinstance(slide_spec, str):
        try:
            slide_spec = json.loads(slide_spec)
        except json.JSONDecodeError as exc:
            raise OutlineError([f"slide_spec is not valid JSON: {exc.msg}."]) from exc
    specs = slide_spec if isinstance(slide_spec, list) else [slide_spec]
    count = len(Presentation(str(deck)).slides)
    if replace_index is not None and not 1 <= replace_index <= count:
        raise OutlineError(
            [f"index {replace_index} is outside the deck's {count} slides (1-based)."]
        )
    if position is not None and not 1 <= position <= count + 1:
        raise OutlineError([f"position {position} is outside 1..{count + 1}."])
    problems: List[str] = []
    first_index = replace_index or position or count + 1
    normalized = [
        normalize_slide(
            sp, first_index + k, problems, Path(base_dir) if base_dir else None
        )
        for k, sp in enumerate(specs)
    ]
    if problems:
        raise OutlineError(problems)
    builder = DeckBuilder(
        {"slides": normalized, "allow_split": False},
        prefs=prefs,
        existing_deck=str(deck),
    )
    first_new = len(builder.prs.slides) + 1
    for s in normalized:
        builder.build_one(s)
    new_count = len(builder.prs.slides) - first_new + 1
    sld_ids = builder.prs.slides._sldIdLst  # noqa: SLF001
    new_ids = list(sld_ids)[-new_count:]
    target = replace_index if replace_index is not None else position
    if target is not None:
        for el in new_ids:
            sld_ids.remove(el)
        for k, el in enumerate(new_ids):
            sld_ids.insert(target - 1 + k, el)
        if replace_index is not None:
            old = list(sld_ids)[replace_index - 1 + new_count]
            builder.prs.part.drop_rel(old.rId)
            sld_ids.remove(old)
    _renumber(builder.prs)
    save_deck(builder.prs, deck)
    at = target or first_new
    return {
        "status": "ok",
        "path": str(deck),
        "slides": len(builder.prs.slides),
        "changed_slides": list(range(at, at + new_count)),
        "warnings": builder.warnings,
        "notes": builder.notes,
        "build_seconds": round(time.perf_counter() - started, 3),
    }


def add_slide(
    deck_path: str,
    slide_spec: Any,
    position: Optional[int] = None,
    prefs: Optional[Dict[str, Any]] = None,
    *,
    base_dir: Optional[str] = None,
) -> Dict[str, Any]:
    """Append (or insert at 1-based *position*) one slide spec, or a list of them."""
    return _edit(
        deck_path,
        slide_spec,
        position=position,
        replace_index=None,
        prefs=prefs,
        base_dir=base_dir,
    )


def update_slide(
    deck_path: str,
    index: int,
    slide_spec: Any,
    prefs: Optional[Dict[str, Any]] = None,
    *,
    base_dir: Optional[str] = None,
) -> Dict[str, Any]:
    """Replace slide *index* (1-based) with the slide(s) built from *slide_spec*."""
    return _edit(
        deck_path,
        slide_spec,
        position=None,
        replace_index=int(index),
        prefs=prefs,
        base_dir=base_dir,
    )
