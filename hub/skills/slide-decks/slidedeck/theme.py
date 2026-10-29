# Copyright(C) 2025-2026 Advanced Micro Devices, Inc. All rights reserved.
# SPDX-License-Identifier: MIT
"""Generate the skill's theme templates from the design tokens.

Each theme becomes a real .pptx template: a slide master whose theme colours
and fonts ARE the tokens, a title zone, a footer zone with a source line and a
slide number, and the layouts the builder maps slide types onto — Title
Slide, Section Header, Title and Content, Two Content, Comparison, Title Only
and Blank. The builder generates a theme in memory for every deck (so brand
colours and fonts from preferences flow into the master); the files under
``templates/`` are the same output, committed so a person can open them in
PowerPoint; ``util/gen_slide_deck_assets.py`` regenerates them.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Optional

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .style import ASPECTS, LAYOUT, SIZES, THEMES, resolve_tokens

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

#: Layouts kept from python-pptx's default master, by name, in this order.
KEEP_LAYOUTS = (
    "Title Slide",
    "Title and Content",
    "Section Header",
    "Two Content",
    "Comparison",
    "Title Only",
    "Blank",
)


def _srgb(parent, hex_rgb: str) -> None:
    for child in list(parent):
        parent.remove(child)
    el = etree.SubElement(parent, f"{{{A_NS}}}srgbClr")
    el.set("val", hex_rgb)


def _set_theme_xml(prs, tokens: Dict[str, Any]) -> None:
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = parse_xml(part.blob)
    root.set("name", f"slide-decks:{tokens['theme']}")
    ns = {"a": A_NS}
    scheme = root.find(".//a:clrScheme", ns)
    scheme.set("name", f"slide-decks {tokens['theme']}")
    c = tokens["colors"]
    pal = list(tokens["palette"]) + [c["primary"]] * 6
    slots = {
        "dk1": c["text"],
        "lt1": c["background"],
        "dk2": c["muted"],
        "lt2": c["panel"],
        "hlink": c["highlight"],
        "folHlink": c.get("context", "9CA3AF"),
    }
    for i in range(6):
        slots[f"accent{i + 1}"] = pal[i]
    for name, value in slots.items():
        _srgb(scheme.find(f"a:{name}", ns), value)
    fonts = root.find(".//a:fontScheme", ns)
    fonts.set(
        "name", f"slide-decks {tokens['fonts']['title']}/{tokens['fonts']['body']}"
    )
    fonts.find("a:majorFont/a:latin", ns).set("typeface", tokens["fonts"]["title"])
    fonts.find("a:minorFont/a:latin", ns).set("typeface", tokens["fonts"]["body"])
    part._blob = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )  # noqa: SLF001


def _drop_layouts(prs) -> None:
    master = prs.slide_master
    id_lst = master._element.find(qn("p:sldLayoutIdLst"))  # noqa: SLF001
    for layout in list(master.slide_layouts):
        if layout.name not in KEEP_LAYOUTS:
            for el in list(id_lst):
                if master.part.related_part(el.get(qn("r:id"))) is layout.part:
                    master.part.drop_rel(el.get(qn("r:id")))
                    id_lst.remove(el)
    # Reorder to KEEP_LAYOUTS order.
    by_name = {}
    for el in list(id_lst):
        part = master.part.related_part(el.get(qn("r:id")))
        by_name[part.slide_layout.name] = el
        id_lst.remove(el)
    for name in KEEP_LAYOUTS:
        if name in by_name:
            id_lst.append(by_name[name])


def _defrpr(
    ppr,
    size_pt: float,
    *,
    bold: Optional[bool] = None,
    color: Optional[str] = None,
    font: Optional[str] = None,
):
    rpr = ppr.find(qn("a:defRPr"))
    if rpr is None:
        rpr = etree.SubElement(ppr, qn("a:defRPr"))
    rpr.set("sz", str(int(size_pt * 100)))
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if color:
        for el in rpr.findall(qn("a:solidFill")):
            rpr.remove(el)
        sf = etree.Element(qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        rpr.insert(0, sf)
    if font:
        for el in rpr.findall(qn("a:latin")):
            rpr.remove(el)
        etree.SubElement(rpr, qn("a:latin")).set("typeface", font)
    return rpr


def _master_text_styles(prs, tokens: Dict[str, Any]) -> None:
    styles = prs.slide_master._element.find(qn("p:txStyles"))  # noqa: SLF001
    c = tokens["colors"]
    title = styles.find(qn("p:titleStyle")).find(qn("a:lvl1pPr"))
    title.set("algn", "l")
    _defrpr(title, SIZES["title_max"], bold=True, color=c["title"])
    body = styles.find(qn("p:bodyStyle"))
    sizes = [
        SIZES["body_max"],
        SIZES["body_max"] - 4,
        SIZES["body_min"],
        SIZES["body_min"],
        SIZES["body_min"],
    ]
    for lvl in range(1, 6):
        ppr = body.find(qn(f"a:lvl{lvl}pPr"))
        if ppr is None:
            continue
        ppr.set("marL", str(int(Pt(20 + 23 * (lvl - 1)))))
        ppr.set("indent", str(int(Pt(-14))))
        _defrpr(ppr, sizes[lvl - 1], color=c["text"] if lvl == 1 else c["muted"])
    other = styles.find(qn("p:otherStyle"))
    if other is not None:
        ppr = other.find(qn("a:lvl1pPr"))
        if ppr is not None:
            _defrpr(ppr, SIZES["body_min"], color=c["text"])


def _place(ph, x, y, w, h) -> None:
    ph.left, ph.top, ph.width, ph.height = int(x), int(y), int(w), int(h)


def _ph_text_size(
    ph,
    size_pt: float,
    color: Optional[str] = None,
    align: Optional[str] = None,
    anchor: Optional[str] = None,
) -> None:
    """Set a placeholder's own list style (lvl1) size/colour on a master/layout."""
    txbody = ph._element.find(qn("p:txBody"))  # noqa: SLF001
    if txbody is None:
        return
    body_pr = txbody.find(qn("a:bodyPr"))
    if anchor and body_pr is not None:
        body_pr.set("anchor", anchor)
    lst = txbody.find(qn("a:lstStyle"))
    if lst is None:
        lst = etree.Element(qn("a:lstStyle"))
        body_pr.addnext(lst)
    ppr = lst.find(qn("a:lvl1pPr"))
    if ppr is None:
        ppr = etree.SubElement(lst, qn("a:lvl1pPr"))
    if align:
        ppr.set("algn", align)
    _defrpr(ppr, size_pt, color=color)


def _add_rect(owner, name: str, x, y, w, h, hex_rgb: str) -> None:
    """Append a filled, borderless rectangle to a master's or layout's shape tree.

    python-pptx only adds shapes to slides, so the element is written directly.
    """
    tree = owner.shapes._spTree  # noqa: SLF001
    ids = [int(v) for v in tree.xpath(".//p:cNvPr/@id")] or [1]
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    xml = (
        f'<p:sp xmlns:p="{P}" xmlns:a="{A_NS}"><p:nvSpPr><p:cNvPr id="{max(ids) + 1}" name="{name}"/>'
        f'<p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr><a:xfrm><a:off x="{int(x)}" y="{int(y)}"/>'
        f'<a:ext cx="{int(w)}" cy="{int(h)}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f'<a:solidFill><a:srgbClr val="{hex_rgb}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>'
    )
    tree.append(parse_xml(xml))


def make_theme_presentation(tokens: Dict[str, Any]):
    """A Presentation whose master and layouts are built from *tokens* (no slides)."""
    prs = Presentation()
    w_in, h_in = ASPECTS[tokens.get("aspect", "16:9")]
    W, H = Inches(w_in), Inches(h_in)
    prs.slide_width, prs.slide_height = W, H
    L = LAYOUT
    mx = Inches(L["margin_x"])
    cw = W - 2 * mx
    title_y, title_h = Inches(L["title_top"]), Inches(
        L["title_bottom"] - L["title_top"]
    )
    body_y, body_h = Inches(L["content_top"]), Inches(
        L["content_bottom"] - L["content_top"]
    )
    foot_y, foot_h = Inches(L["footer_top"]), Inches(
        L["footer_bottom"] - L["footer_top"]
    )
    gutter = Inches(L["gutter"])
    c = tokens["colors"]

    _set_theme_xml(prs, tokens)
    _drop_layouts(prs)
    master = prs.slide_master
    master.background.fill.solid()
    from pptx.dml.color import RGBColor

    master.background.fill.fore_color.rgb = RGBColor.from_string(c["background"])
    _master_text_styles(prs, tokens)
    num_w = Inches(1.0)
    for ph in master.placeholders:
        t = ph.placeholder_format.type
        if t == PP_PLACEHOLDER.TITLE:
            _place(ph, mx, title_y, cw, title_h)
            ph._element.find(qn("p:txBody")).find(qn("a:bodyPr")).set(
                "anchor", "b"
            )  # noqa: SLF001
        elif t == PP_PLACEHOLDER.BODY:
            _place(ph, mx, body_y, cw, body_h)
        elif t == PP_PLACEHOLDER.FOOTER:
            _place(ph, mx, foot_y, cw - num_w - Inches(2.2), foot_h)
            _ph_text_size(ph, SIZES["footnote"], c["muted"], "l")
        elif t == PP_PLACEHOLDER.DATE:
            _place(ph, W - mx - num_w - Inches(2.0), foot_y, Inches(1.9), foot_h)
            _ph_text_size(ph, SIZES["footnote"], c["muted"], "r")
        elif t == PP_PLACEHOLDER.SLIDE_NUMBER:
            _place(ph, W - mx - num_w, foot_y, num_w, foot_h)
            _ph_text_size(ph, SIZES["footnote"], c["muted"], "r")
    if tokens.get("title_bar", True):
        # A short accent rule under the title zone: the deck's one signature mark.
        _add_rect(
            master,
            "Title bar",
            mx,
            title_y + title_h + Inches(0.06),
            Inches(0.9),
            Inches(0.05),
            c["primary"],
        )

    for layout in master.slide_layouts:
        name = layout.name
        phs = {}
        for ph in layout.placeholders:
            phs.setdefault(ph.placeholder_format.type, []).append(ph)
        titles = phs.get(PP_PLACEHOLDER.TITLE, []) + phs.get(
            PP_PLACEHOLDER.CENTER_TITLE, []
        )
        bodies = sorted(
            phs.get(PP_PLACEHOLDER.BODY, []) + phs.get(PP_PLACEHOLDER.OBJECT, []),
            key=lambda p: (p.placeholder_format.idx),
        )
        for t in (
            PP_PLACEHOLDER.FOOTER,
            PP_PLACEHOLDER.DATE,
            PP_PLACEHOLDER.SLIDE_NUMBER,
        ):
            for ph in phs.get(t, []):
                src = next(
                    (m for m in master.placeholders if m.placeholder_format.type == t),
                    None,
                )
                if src is not None:
                    _place(ph, src.left, src.top, src.width, src.height)
        if name == "Title Slide":
            layout._element.set("showMasterSp", "0")  # noqa: SLF001
            for ph in titles:
                _place(ph, mx, int(H * 0.30), cw, int(H * 0.22))
                _ph_text_size(ph, SIZES["cover_title_max"], c["title"], "l", "b")
            for ph in phs.get(PP_PLACEHOLDER.SUBTITLE, []):
                _place(
                    ph,
                    mx,
                    int(H * 0.30) + int(H * 0.22) + Inches(0.15),
                    cw,
                    int(H * 0.2),
                )
                _ph_text_size(ph, SIZES["subtitle"], c["muted"], "l", "t")
            _add_rect(
                layout,
                "Cover bar",
                mx,
                int(H * 0.30) - Inches(0.25),
                Inches(1.2),
                Inches(0.08),
                c["primary"],
            )
        elif name == "Section Header":
            layout._element.set("showMasterSp", "0")  # noqa: SLF001
            for ph in titles:
                _place(ph, mx, int(H * 0.36), cw, int(H * 0.18))
                _ph_text_size(ph, SIZES["title_max"] + 4, c["title"], "l", "b")
                for rpr in ph._element.iter(qn("a:defRPr")):  # noqa: SLF001
                    rpr.attrib.pop("cap", None)
            for ph in bodies:
                _place(
                    ph,
                    mx,
                    int(H * 0.36) + int(H * 0.18) + Inches(0.1),
                    cw,
                    int(H * 0.15),
                )
                _ph_text_size(ph, SIZES["subtitle"] - 2, c["muted"], "l", "t")
                ppr = (
                    ph._element.find(qn("p:txBody"))
                    .find(qn("a:lstStyle"))
                    .find(qn("a:lvl1pPr"))
                )  # noqa: SLF001
                ppr.set("marL", "0")
                ppr.set("indent", "0")
                if ppr.find(qn("a:buNone")) is None:
                    ppr.insert(0, etree.Element(qn("a:buNone")))
            _add_rect(
                layout,
                "Section bar",
                mx,
                int(H * 0.36) - Inches(0.2),
                Inches(1.2),
                Inches(0.08),
                c["primary"],
            )
        else:
            for ph in titles:
                _place(ph, mx, title_y, cw, title_h)
            if name == "Title and Content":
                for ph in bodies[:1]:
                    _place(ph, mx, body_y, cw, body_h)
            elif name == "Two Content":
                half = (cw - gutter) // 2
                for k, ph in enumerate(bodies[:2]):
                    _place(ph, mx + k * (half + gutter), body_y, half, body_h)
            elif name == "Comparison":
                half = (cw - gutter) // 2
                head_h = Inches(0.6)
                heads = [
                    p
                    for p in bodies
                    if p.placeholder_format.type == PP_PLACEHOLDER.BODY
                ]
                contents = [
                    p
                    for p in bodies
                    if p.placeholder_format.type == PP_PLACEHOLDER.OBJECT
                ]
                for k, ph in enumerate(heads[:2]):
                    _place(ph, mx + k * (half + gutter), body_y, half, head_h)
                    _ph_text_size(ph, SIZES["body_max"], c["primary"], "l", "b")
                for k, ph in enumerate(contents[:2]):
                    _place(
                        ph,
                        mx + k * (half + gutter),
                        body_y + head_h + Inches(0.1),
                        half,
                        body_h - head_h - Inches(0.1),
                    )
    master.name = f"slide-decks {tokens['theme']}"
    return prs


def write_templates(out_dir: Path) -> Dict[str, Path]:
    """Save one template per built-in theme into the existing *out_dir*."""
    written = {}
    for name in THEMES:
        tokens = resolve_tokens({"theme": name})
        prs = make_theme_presentation(tokens)
        # Core properties carry fixed values so regenerating is reproducible.
        import datetime as _dt

        cp = prs.core_properties
        cp.title = f"slide-decks {THEMES[name]['label']} theme"
        cp.author = "slide-decks skill"
        cp.last_modified_by = "slide-decks skill"
        fixed = _dt.datetime(2026, 1, 1)
        cp.created = cp.modified = fixed
        cp.revision = 1
        path = out_dir / f"{name}.pptx"
        prs.save(str(path))
        written[name] = path
    return written
